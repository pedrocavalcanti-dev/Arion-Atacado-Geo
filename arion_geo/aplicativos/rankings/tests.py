from io import BytesIO
from datetime import datetime, timedelta, timezone as datetime_timezone
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from aplicativos.base_links.models import Operadora
from aplicativos.configuracoes.models import GroupDataScope
from aplicativos.integracoes.models import (
    CMDBLinkMirror, CMDBOperatorMirror, OTRSTicket, OTRSTicketConfigItem,
)
from aplicativos.painel.services.analytics import AnalyticsQueryService
from .services import RANKING_WEIGHTS, RankingService, _apply_scores


class RankingsMirrorTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="rank-mirror", password="safe-password-1")
        group = Group.objects.get(name="Analista NOC")
        self.user.groups.set([group])
        scope, _ = GroupDataScope.objects.get_or_create(group=group)
        scope.mode = GroupDataScope.Mode.ALL
        scope.save()
        self.operator = CMDBOperatorMirror.objects.create(source_config_item_id="op-1", name="Operadora Real")
        self.link = CMDBLinkMirror.objects.create(
            source_config_item_id="ci-1", source_number="N-1", name="Link Real", state="SP",
            contracting_client="Claro", business_project="Bradesco", operator=self.operator,
        )
        self.client.force_login(self.user)
        self._next_id = 910000

    def make_ticket(self, *, day=5, status="OPEN", links=None, hours_ago=1, terminal_hours=None,
                    terminal_missing=False, priority="Médio", created_at=None):
        self._next_id += 1
        created = created_at or datetime(2026, 10, day, 12, tzinfo=datetime_timezone.utc)
        terminal = created + timedelta(hours=terminal_hours) if terminal_hours is not None else None
        if terminal_missing:
            terminal = None
        ticket = OTRSTicket.objects.create(
            external_id=self._next_id, ticket_number=f"T-{self._next_id}", title="Ticket de teste",
            queue_external_id=107, queue_name="NOC - ATACADO", external_state_id=self._next_id,
            external_state_name=status, business_status=status,
            is_business_open=status in {"OPEN", "PENDING", "SCHEDULED"},
            external_priority_id=2, priority_name=priority, source_created_at=created,
            terminal_state_at=terminal,
        )
        for link in (links or [self.link]):
            OTRSTicketConfigItem.objects.create(
                ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link,
                relation_type="DependsOn", relation_state="Valid",
            )
        return ticket

    def test_valid_ticket_count_matches_cockpit_and_excludes_canceled_unknown(self):
        self.make_ticket(status="OPEN")
        self.make_ticket(status="RESOLVED", terminal_hours=30)
        self.make_ticket(status="CANCELED")
        self.make_ticket(status="UNKNOWN")
        start, end = "2026-10-01", "2026-10-24"
        response = self.client.get(reverse("rankings:index"), {"view": "custom", "date_from": start, "date_to": end})
        cockpit = AnalyticsQueryService().query({
            "metric": "tickets", "start": start, "end": end, "group_by": "provider", "filters": {},
        }, self.user)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["tickets_valid"], 2)
        self.assertEqual(response.context["tickets_valid"], cockpit["value"])
        self.assertEqual(response.context["providers"][0]["tickets"], 2)
        self.assertContains(response, "Índice ARION")

    def test_canceled_and_unknown_do_not_change_ranking_drawer_history_or_exports(self):
        now = datetime(2026, 10, 25, 12, tzinfo=datetime_timezone.utc)
        date_from = datetime(2026, 10, 1).date()
        date_to = datetime(2026, 10, 24).date()
        valid = self.make_ticket(status="OPEN", created_at=now - timedelta(hours=30))

        def build_snapshot():
            return RankingService.build(
                self.user, date_from=date_from, date_to=date_to, now=now,
                detail_provider_id=self.operator.pk,
            )

        baseline = build_snapshot()
        canceled = self.make_ticket(
            status="CANCELED", created_at=now - timedelta(hours=72), terminal_hours=48,
        )
        unknown = self.make_ticket(
            status="UNKNOWN", created_at=now - timedelta(hours=72), terminal_hours=48,
        )
        snapshot = build_snapshot()

        for key in ("tickets_valid", "links_affected", "links_recurrent", "tickets_over_24h", "backlog_current"):
            self.assertEqual(snapshot[key], baseline[key], key)
        provider = snapshot["providers"][0]
        baseline_provider = baseline["providers"][0]
        for key in (
            "tickets", "links_affected", "recurrent_links", "tickets_over_24h", "backlog_current",
            "risk_score", "impact_noc_score", "trend", "trend_delta", "history",
        ):
            self.assertEqual(provider[key], baseline_provider[key], key)
        self.assertEqual(provider["tickets"], 1)
        self.assertEqual(provider["recurrent_links"], 0)
        self.assertEqual(provider["tickets_over_24h"], 1)
        self.assertEqual(provider["backlog_current"], 1)

        link = snapshot["ranked_links_all"][0]
        self.assertEqual((link["tickets"], link["recurrent"], link["over_24h"]), (1, False, 1))
        self.assertEqual([row["id"] for row in snapshot["ticket_export_rows"]], [valid.pk])
        detail = snapshot["detail"]
        self.assertEqual(detail["ticket_records_total"], 1)
        self.assertEqual([row["id"] for row in detail["ticket_records"]], [valid.pk])
        self.assertEqual([row["id"] for row in detail["over_24h_tickets"]], [valid.pk])
        self.assertEqual([row["id"] for row in detail["backlog_tickets"]], [valid.pk])
        self.assertEqual(detail["problematic_links"][0]["recurrent"], False)

        params = {"view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24"}
        detail_response = self.client.get(reverse("rankings:provider-detail", args=[self.operator.pk]), params)
        self.assertEqual(detail_response.status_code, 200)
        drawer_ids = {row["id"] for row in detail_response.json()["ticket_records"]}
        self.assertEqual(drawer_ids, {valid.pk})
        self.assertNotIn(canceled.pk, drawer_ids)
        self.assertNotIn(unknown.pk, drawer_ids)

        for file_format, builder in (("pdf", "build_pdf"), ("xlsx", "build_xlsx")):
            with patch(f"aplicativos.rankings.exports.{builder}", return_value=b"export") as build_export:
                response = self.client.get(reverse("rankings:export", args=[file_format]), params)
            self.assertEqual(response.status_code, 200)
            export_data = build_export.call_args.args[0]
            export_ids = {row["id"] for row in export_data["ticket_export_rows"]}
            self.assertEqual(export_ids, {valid.pk})
            self.assertNotIn(canceled.pk, export_ids)
            self.assertNotIn(unknown.pk, export_ids)

    def test_reason_filter_and_top_reasons_canonicalize_confirmed_aliases(self):
        raws = ("FalhaMassiva", "Falha Massiva", "Defeito Massivo")
        for index, raw in enumerate(raws, 1):
            ticket = self.make_ticket(day=10 + index, status="RESOLVED", terminal_hours=2)
            ticket.closing_code_raw = raw
            ticket.save(update_fields=("closing_code_raw",))
        other = self.make_ticket(day=18, status="RESOLVED", terminal_hours=2)
        other.closing_code_raw = "Falha Lógica na Operadora"
        other.save(update_fields=("closing_code_raw",))
        data = RankingService.build(
            self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date(),
            filters={"reason": "falha_massiva"}, detail_provider_id=self.operator.pk,
        )
        self.assertEqual(data["tickets_valid"], 3)
        self.assertEqual(data["detail"]["reasons_top"][0]["name"], "Falha Massiva")
        self.assertEqual(data["detail"]["reasons_top"][0]["tickets"], 3)

    def test_same_provider_multiple_links_deduplicate_ticket_and_count_affected_links(self):
        other = CMDBLinkMirror.objects.create(
            source_config_item_id="ci-2", source_number="N-2", state="SP", operator=self.operator,
        )
        first = self.make_ticket(links=[self.link, other])
        self.make_ticket(day=6, status="RESOLVED", links=[self.link], terminal_hours=30)
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date(), now=datetime(2026, 10, 24, tzinfo=datetime_timezone.utc))
        provider = data["providers"][0]
        self.assertIn(first.pk, provider["period_ticket_ids"])
        self.assertEqual(provider["tickets"], 2)
        self.assertEqual(provider["links_affected"], 2)
        self.assertEqual(provider["recurrent_links"], 1)
        # The first ticket is currently open and was created on Oct 5, so it is aged too.
        self.assertEqual(provider["tickets_over_24h"], 2)

    def test_cross_provider_ticket_counts_once_per_provider_but_once_nationally(self):
        other_provider = CMDBOperatorMirror.objects.create(source_config_item_id="op-2", name="Provider B")
        other_link = CMDBLinkMirror.objects.create(source_config_item_id="ci-2", operator=other_provider, state="RJ")
        self.make_ticket(links=[self.link, other_link])
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date())
        self.assertEqual(data["tickets_valid"], 1)
        self.assertEqual(sum(provider["tickets"] for provider in data["providers"]), 2)

    def test_open_aging_uses_now_and_terminal_aging_uses_resolution_or_terminal_fallback(self):
        now = datetime(2026, 10, 24, 12, tzinfo=datetime_timezone.utc)
        self.make_ticket(status="OPEN", created_at=now - timedelta(hours=25))
        self.make_ticket(status="CLOSED", created_at=now - timedelta(hours=40), terminal_hours=30)
        missing = self.make_ticket(status="RESOLVED", created_at=now - timedelta(hours=60), terminal_missing=True)
        OTRSTicket.objects.filter(pk=missing.pk).update(resolved_at=now - timedelta(hours=58))
        no_terminal_date = self.make_ticket(status="RESOLVED", created_at=now - timedelta(hours=60), terminal_missing=True)
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date(), now=now)
        provider = data["providers"][0]
        self.assertEqual(provider["aged_open"], 1)
        self.assertEqual(provider["terminal_over_24h"], 1)
        self.assertEqual(provider["tickets_over_24h"], 2)
        self.assertEqual(provider["eligible_tickets"], 3)
        self.assertEqual(provider["over_24h_rate"], 2 / 3)
        self.assertEqual(OTRSTicket.objects.get(pk=no_terminal_date.pk).resolved_at, None)

    def test_ranking_duration_and_xlsx_use_resolved_at_not_later_closed_date(self):
        created = datetime(2026, 10, 10, 8, tzinfo=datetime_timezone.utc)
        ticket = self.make_ticket(status="CLOSED", created_at=created, terminal_hours=31)
        ticket.resolved_at = created + timedelta(hours=7)
        ticket.save(update_fields=("resolved_at",))

        data = RankingService.build(
            self.user, date_from=created.date(), date_to=created.date(),
            now=created + timedelta(days=2), detail_provider_id=self.operator.pk,
        )
        record = data["detail"]["ticket_records"][0]
        self.assertEqual(record["status"], "CLOSED")
        self.assertEqual(record["duration_hours"], 7)
        self.assertEqual(record["resolved_at"], "10/10/2026 12:00")
        self.assertEqual(record["terminal_at"], "11/10/2026 12:00")
        self.assertEqual(record["operational_end_at"], "10/10/2026 12:00")
        self.assertEqual(record["operational_end_source"], "resolved_at")

        response = self.client.get(reverse("rankings:export", args=["xlsx"]), {
            "view": "custom", "date_from": "2026-10-10", "date_to": "2026-10-10",
        })
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        sheet = workbook["Tickets"]
        headers = [cell.value for cell in sheet[1]]
        values = [cell.value for cell in sheet[2]]
        row = dict(zip(headers, values))
        self.assertEqual(row["Status"], "CLOSED")
        self.assertEqual(row["Resolvido em"], "10/10/2026 12:00")
        self.assertEqual(row["Terminal OTRS em"], "11/10/2026 12:00")
        self.assertEqual(row["Término operacional usado"], "10/10/2026 12:00")
        self.assertEqual(row["Duração (horas)"], 7)

    def test_provider_with_zero_tickets_has_zero_index_and_stays_in_the_service_universe(self):
        small = CMDBOperatorMirror.objects.create(source_config_item_id="op-small", name="Provider pequeno")
        CMDBLinkMirror.objects.create(source_config_item_id="ci-small", operator=small)
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date())
        provider = next(row for row in data["providers"] if row["name"] == "Provider pequeno")
        self.assertEqual(provider["tickets"], 0)
        self.assertTrue(provider["sample_small"])
        self.assertEqual(provider["risk_score"], 0)
        self.assertEqual(provider["risk_label"], "Sem ocorrências")
        self.assertTrue(provider["no_occurrence"])

    def test_one_link_one_incident_remains_visible_with_high_proportional_rate(self):
        small = CMDBOperatorMirror.objects.create(source_config_item_id="op-risk", name="Provider 1 link")
        link = CMDBLinkMirror.objects.create(source_config_item_id="ci-risk", operator=small, state="MG")
        self.make_ticket(links=[link])
        large = CMDBOperatorMirror.objects.create(source_config_item_id="op-large", name="Provider 10 links")
        for index in range(10):
            CMDBLinkMirror.objects.create(source_config_item_id=f"ci-large-{index}", operator=large)
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date())
        provider = next(row for row in data["providers"] if row["name"] == "Provider 1 link")
        self.assertEqual(provider["tickets"], 1)
        self.assertEqual(provider["affected_rate"], 1)
        self.assertEqual(provider["tickets_per_100_links"], 100)
        self.assertTrue(provider["sample_small"])
        large_provider = next(row for row in data["providers"] if row["name"] == "Provider 10 links")
        self.assertGreater(provider["risk_score"], large_provider["risk_score"])

    def test_canonical_client_project_filters_and_rolling_history(self):
        for month, day in ((8, 8), (9, 8), (10, 8)):
            self.make_ticket(created_at=datetime(2026, month, day, 12, tzinfo=datetime_timezone.utc))
        data = RankingService.build(
            self.user, date_from=datetime(2026, 8, 1).date(), date_to=datetime(2026, 10, 24).date(),
            filters={"client": "Claro", "project": "Bradesco"},
        )
        provider = data["providers"][0]
        self.assertEqual(provider["tickets"], 3)
        self.assertEqual(len(provider["history"]), 3)
        self.assertIn(provider["trend"], {"worse", "better", "stable"})
        self.assertEqual(provider["risk_score"], round(provider["risk_score"], 1))
        self.assertEqual(provider["risk_score"], round(sum(provider["risk_components"].values()), 1))
        self.assertNotEqual(provider["risk_score"], sum(point["score"] for point in provider["history"]))

    def test_monthly_trend_can_show_worsening_and_improvement(self):
        other = CMDBOperatorMirror.objects.create(source_config_item_id="op-trend", name="Provider tendência")
        other_link = CMDBLinkMirror.objects.create(source_config_item_id="ci-trend", operator=other)
        self.make_ticket(day=5, status="RESOLVED", terminal_hours=1)
        self.make_ticket(day=5, status="RESOLVED", terminal_hours=1,
                         created_at=datetime(2026, 9, 5, 12, tzinfo=datetime_timezone.utc), links=[other_link])
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date())
        current_only = next(row for row in data["providers"] if row["name"] == "Operadora Real")
        previous_only = next(row for row in data["providers"] if row["name"] == "Provider tendência")
        self.assertEqual(current_only["trend"], "worse")
        self.assertGreater(current_only["trend_delta"], 2)
        self.assertEqual(previous_only["trend"], "better")
        self.assertLess(previous_only["trend_delta"], -2)

    def test_provider_detail_matches_ranking_row_and_link_identity_is_canonical(self):
        self.make_ticket(status="RESOLVED", terminal_hours=30)
        response = self.client.get(reverse("rankings:index"), {"view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24"})
        row = response.context["providers"][0]
        detail = self.client.get(reverse("rankings:provider-detail", args=[self.operator.pk]), {
            "view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24",
        })
        self.assertEqual(detail.status_code, 200)
        payload = detail.json()
        self.assertEqual(payload["risk_score"], row["risk_score"])
        self.assertEqual(payload["tickets"], row["tickets"])
        self.assertEqual(payload["links_total"], row["links_total"])
        link = response.context["ranked_links"][0]
        self.assertEqual((link["client"], link["project"], link["provider"]), ("Claro", "Bradesco", "Operadora Real"))

    def test_best_view_requires_an_event_and_podium_uses_links_or_ticket_sample(self):
        small = CMDBOperatorMirror.objects.create(source_config_item_id="op-small-best", name="Amostra pequena")
        small_link = CMDBLinkMirror.objects.create(source_config_item_id="ci-small-best", operator=small)
        eligible_by_links = CMDBOperatorMirror.objects.create(source_config_item_id="op-eligible", name="Provider elegivel por links")
        eligible_links = [CMDBLinkMirror.objects.create(source_config_item_id=f"ci-eligible-{index}", operator=eligible_by_links) for index in range(5)]
        eligible_by_tickets = CMDBOperatorMirror.objects.create(source_config_item_id="op-eligible-tickets", name="Provider elegivel por chamados")
        eligible_ticket_link = CMDBLinkMirror.objects.create(source_config_item_id="ci-eligible-ticket", operator=eligible_by_tickets)
        stable = CMDBOperatorMirror.objects.create(source_config_item_id="op-stable", name="Provider estavel")
        stable_links = [CMDBLinkMirror.objects.create(source_config_item_id=f"ci-stable-{index}", operator=stable) for index in range(8)]
        self.make_ticket(links=[small_link])
        self.make_ticket(links=[eligible_links[0]])
        for day in (6, 7, 8):
            self.make_ticket(day=day, links=[eligible_ticket_link])

        response = self.client.get(reverse("rankings:index"), {
            "view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24", "ranking_view": "best",
        })
        self.assertEqual(response.status_code, 200)
        rows = response.context["providers"]
        self.assertTrue(all(row["tickets"] >= 1 for row in rows))
        self.assertEqual(rows, sorted(rows, key=lambda row: (row["risk_score"], -row["links_total"], -row["eligible_tickets"], row["name"].casefold())))
        self.assertEqual(response.context["best_eligible_count"], 2)
        self.assertTrue(all(row["best_eligible"] for row in response.context["top_providers"]))
        self.assertNotIn(small.pk, [row["id"] for row in response.context["top_providers"]])
        self.assertNotIn(stable.pk, [row["id"] for row in rows])

        risk_response = self.client.get(reverse("rankings:index"), {
            "view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24", "ranking_view": "risk",
        })
        risk_rows = risk_response.context["providers"]
        self.assertEqual(risk_rows, sorted(risk_rows, key=lambda row: (-row["risk_score"], -row["risk_score"], -row["impact_noc_score"], -row["tickets"], -row["links_affected"], row["name"].casefold())))
        self.assertIn(small.pk, [row["id"] for row in risk_response.context["top_providers"]])

        stable_response = self.client.get(reverse("rankings:index"), {
            "view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24", "ranking_view": "stable",
        })
        self.assertEqual(stable_response.status_code, 200)
        self.assertEqual(stable_response.context["providers"][0]["id"], stable.pk)
        self.assertEqual(stable_response.context["providers"][0]["links_total"], len(stable_links))
        self.assertContains(stable_response, "Sem ocorr&ecirc;ncias no per&iacute;odo")

    def test_detail_responsibility_and_ticket_rows_deduplicate_same_provider_and_sanitize_title(self):
        second_link = CMDBLinkMirror.objects.create(
            source_config_item_id="ci-1b", source_number="N-1B", name="Segundo link", state="RJ",
            contracting_client="Claro", business_project="Bradesco", operator=self.operator,
        )
        procedente = self.make_ticket(links=[self.link, second_link])
        procedente.title = "Contato pessoa@example.com 11 99999-1234"
        procedente.responsibility_raw = "Procedente"
        procedente.responsibility_normalized = "PROCEDENTE"
        procedente.save(update_fields=("title", "responsibility_raw", "responsibility_normalized"))
        improcedente = self.make_ticket(day=6, status="RESOLVED", terminal_hours=2)
        improcedente.responsibility_raw = "Improcedente"
        improcedente.responsibility_normalized = "IMPROCEDENTE"
        improcedente.save(update_fields=("responsibility_raw", "responsibility_normalized"))
        for day in (7, 8, 9):
            ticket = self.make_ticket(day=day, status="RESOLVED", terminal_hours=2)
            ticket.closing_code_raw = "SemIntervencaoTecnica"
            ticket.responsibility_raw = "Improcedente"
            ticket.responsibility_normalized = "IMPROCEDENTE"
            ticket.save(update_fields=("closing_code_raw", "responsibility_raw", "responsibility_normalized"))
        second_procedente = self.make_ticket(day=10, status="RESOLVED", terminal_hours=2)
        second_procedente.responsibility_raw = "Procedente"
        second_procedente.responsibility_normalized = "PROCEDENTE"
        second_procedente.save(update_fields=("responsibility_raw", "responsibility_normalized"))
        self.make_ticket(day=11)

        response = self.client.get(reverse("rankings:provider-detail", args=[self.operator.pk]), {
            "view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24",
        })
        self.assertEqual(response.status_code, 200)
        detail = response.json()
        self.assertEqual(detail["tickets"], 7)
        self.assertEqual(detail["responsibility"]["procedentes"], 2)
        self.assertEqual(detail["responsibility"]["improcedentes"], 4)
        self.assertEqual(detail["responsibility"]["indeterminados"], 1)
        self.assertEqual(detail["responsibility"]["procedure_rate"], 2 / 6)
        self.assertTrue(detail["responsibility"]["review"])
        self.assertEqual(detail["ticket_records_total"], 7)
        record = next(item for item in detail["ticket_records"] if item["id"] == procedente.pk)
        self.assertEqual(len(record["links"]), 2)
        self.assertIn("Criado no per", record["why"][0])
        self.assertNotIn("pessoa@example.com", record["title"])
        self.assertNotIn("99999-1234", record["title"])

    def test_responsibility_classification_does_not_change_arion_v2_score(self):
        ticket = self.make_ticket()
        before = RankingService.build(
            self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date(),
        )["providers"][0]["risk_score"]
        OTRSTicket.objects.filter(pk=ticket.pk).update(
            responsibility_raw="Improcedente", responsibility_normalized="IMPROCEDENTE",
        )
        after = RankingService.build(
            self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date(),
        )["providers"][0]["risk_score"]
        self.assertEqual(before, after)

    def test_pdf_and_xlsx_exports_preserve_data_scope_and_include_full_filtered_rows(self):
        from aplicativos.base_links.models import Operadora, Projeto

        allowed_operator = Operadora.objects.create(name="Claro", normalized_name="claro")
        allowed_project = Projeto.objects.create(name="Bradesco", normalized_name="bradesco")
        group = self.user.groups.get(name="Analista NOC")
        scope = group.data_scope
        scope.mode = GroupDataScope.Mode.CUSTOM
        scope.save()
        scope.allowed_operators.set([allowed_operator])
        scope.allowed_projects.set([allowed_project])
        outside_operator = CMDBOperatorMirror.objects.create(source_config_item_id="op-export-out", name="Provider fora do escopo")
        outside_link = CMDBLinkMirror.objects.create(
            source_config_item_id="ci-export-out", name="Link fora do escopo", operator=outside_operator,
            business_project="Bradesco", state="RJ",
        )
        wrong_project_link = CMDBLinkMirror.objects.create(
            source_config_item_id="ci-export-project-out", name="Projeto fora do escopo", operator=self.operator,
            business_project="Projeto externo", state="RJ",
        )
        self.make_ticket(links=[self.link])
        self.make_ticket(links=[outside_link])
        self.make_ticket(links=[wrong_project_link])
        params = {"view": "custom", "date_from": "2026-10-01", "date_to": "2026-10-24", "ranking_view": "best"}

        pdf = self.client.get(reverse("rankings:export", args=["pdf"]), params)
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        self.assertGreater(len(pdf.content), 1000)

        xlsx = self.client.get(reverse("rankings:export", args=["xlsx"]), params)
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn("spreadsheetml", xlsx["Content-Type"])
        workbook = load_workbook(BytesIO(xlsx.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Resumo", "Ranking", "Providers", "Links", "Tickets"])
        self.assertEqual(workbook["Links"].max_row, 2)
        self.assertEqual(workbook["Tickets"].max_row, 2)
        ranking_headers = next(workbook["Ranking"].iter_rows(values_only=True))
        self.assertIn("Índice ARION V2", ranking_headers)
        self.assertIn("Impacto no NOC /30", ranking_headers)
        self.assertIn("Chamados / link", ranking_headers)
        self.assertNotIn("Tickets / 100 links", ranking_headers)
        values = [cell for row in workbook["Links"].iter_rows(values_only=True) for cell in row]
        self.assertNotIn("ci-export-out", values)
        self.assertNotIn("ci-export-project-out", values)

    def test_custom_data_scope_excludes_links_outside_authorized_intersection(self):
        CMDBLinkMirror.objects.create(source_config_item_id="ci-out", operator=self.operator, business_project="Outside")
        group = self.user.groups.get(name="Analista NOC")
        scope = group.data_scope
        scope.mode = GroupDataScope.Mode.CUSTOM
        scope.save()
        scope.allowed_operators.set([Operadora.objects.create(name=self.operator.name, normalized_name=self.operator.name.casefold())])
        data = RankingService.build(self.user, date_from=datetime(2026, 10, 1).date(), date_to=datetime(2026, 10, 24).date())
        self.assertEqual(data["links_total"], 0)
        self.assertEqual(data["providers"], [])


class RankingScoreTests(SimpleTestCase):
    def test_weight_sets_are_complete_and_scores_are_bounded(self):
        self.assertEqual(sum(RANKING_WEIGHTS.values()), 1)
        rows = [
            {"links_total": 50, "tickets": 0, "links_affected": 0, "recurrent_links": 0,
             "eligible_tickets": 0, "tickets_over_24h": 0, "backlog_current": 0},
            {"links_total": 1, "tickets": 4, "links_affected": 1, "recurrent_links": 1,
             "eligible_tickets": 4, "tickets_over_24h": 3, "backlog_current": 1},
            {"links_total": 267, "tickets": 28, "links_affected": 24, "recurrent_links": 4,
             "eligible_tickets": 28, "tickets_over_24h": 6, "backlog_current": 1},
            {"links_total": 267, "tickets": 2, "links_affected": 2, "recurrent_links": 0,
             "eligible_tickets": 2, "tickets_over_24h": 0, "backlog_current": 0},
        ]
        _apply_scores(rows)
        self.assertEqual(rows[0]["risk_score"], 0)
        self.assertEqual(rows[0]["risk_label"], "Sem ocorrências")
        self.assertGreaterEqual(rows[1]["risk_score"], 60)
        self.assertGreater(rows[2]["impact_noc_score"], rows[3]["impact_noc_score"])
        self.assertLess(rows[3]["risk_score"], 20)
        self.assertEqual(rows[2]["risk_components"]["impact_noc"], rows[2]["impact_noc_score"])

    def test_recurrence_aging_and_backlog_each_increase_their_component(self):
        recurrent = {"links_total": 1, "tickets": 5, "links_affected": 1, "recurrent_links": 1,
                     "eligible_tickets": 5, "tickets_over_24h": 4, "backlog_current": 1}
        _apply_scores([recurrent])
        self.assertEqual(recurrent["recurrent_score"], 20)
        self.assertEqual(recurrent["over_24h_score"], 16)

        without_backlog = {"links_total": 10, "tickets": 2, "links_affected": 1, "recurrent_links": 0,
                           "eligible_tickets": 2, "tickets_over_24h": 0, "backlog_current": 0}
        with_backlog = {**without_backlog, "backlog_current": 1}
        _apply_scores([without_backlog, with_backlog])
        self.assertGreater(with_backlog["backlog_score"], without_backlog["backlog_score"])
        self.assertGreater(with_backlog["risk_score"], without_backlog["risk_score"])

    def test_zero_ticket_provider_is_excluded_from_best_but_available_in_stable_view(self):
        rows = [
            {"links_total": 50, "tickets": 0, "links_affected": 0, "recurrent_links": 0,
             "eligible_tickets": 0, "tickets_over_24h": 0, "backlog_current": 0},
            {"links_total": 2, "tickets": 1, "links_affected": 1, "recurrent_links": 0,
             "eligible_tickets": 1, "tickets_over_24h": 0, "backlog_current": 0},
        ]
        _apply_scores(rows)
        self.assertFalse(rows[0]["risk_eligible"])
        self.assertFalse(rows[0]["best_eligible"])
        self.assertTrue(rows[0]["no_occurrence"])
        self.assertTrue(rows[1]["risk_eligible"])
        self.assertFalse(rows[1]["best_eligible"])

    def test_drawer_and_table_use_v2_portuguese_labels_without_percentile_rows(self):
        root = Path(__file__).resolve().parents[2]
        script = (root.parent / "frontend_src/js/rankings/rankings.js").read_text(encoding="utf-8")
        template = (root / "templates/rankings/index.html").read_text(encoding="utf-8")
        for label in ("Impacto no NOC", "Base afetada", "Reincid", "Backlog atual", "Chamados por link"):
            self.assertIn(label, script + template)
        self.assertNotIn("percentil'", script)
        self.assertNotIn("Tickets / 100", template)
        self.assertNotIn("Tickets ARION", script)
        for legacy in ("Affected Rate", "Recurrent Rate", "Over 24h Rate", "Criticality Rate"):
            self.assertNotIn(legacy, script + template)


class CockpitRankingPresentationTests(SimpleTestCase):
    def test_collapse_toggles_invert_state_and_uf_flags_are_local_assets(self):
        root = Path(__file__).resolve().parents[2]
        script = (root.parent / "frontend_src/js/painel/cockpit.js").read_text(encoding="utf-8")
        self.assertIn("setDetailCollapsed(key, !detailCollapseState[key])", script)
        self.assertIn("data-details-global", script)
        self.assertIn("aria-expanded", script)
        self.assertIn("/static/imagens/bandeiras/estados/${code.toLowerCase()}.svg", script)
        self.assertTrue((root / "static/imagens/bandeiras/estados/sp.svg").exists())
