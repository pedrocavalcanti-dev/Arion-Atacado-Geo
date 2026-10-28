"""Executive PDF and workbook exports for the scoped provider ranking."""
from io import BytesIO
from xml.sax.saxutils import escape


def _safe_cell(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _view_label(filters):
    return "Consolidado 3 meses" if filters["view"] == "3months" else "Período personalizado" if filters["view"] == "custom" else "Visão do mês"


def _ranking_label(filters):
    return {"best": "Melhor desempenho observado", "stable": "Sem ocorrências"}.get(filters["ranking_view"], "Maior risco")


def _filter_labels(filters):
    labels = {
        "client": "Cliente", "project": "Projeto", "operator": "Provedor / ISP", "uf": "UF",
        "priority": "Prioridade", "reason": "Motivo", "status": "Status",
    }
    result = [(labels[key], filters.get(key)) for key in labels if filters.get(key)]
    if filters.get("only_aged"):
        result.append(("Apenas ARION >24h", "Sim"))
    if filters.get("only_recurrent"):
        result.append(("Apenas reincidentes", "Sim"))
    if filters.get("q"):
        result.append(("Pesquisa", filters["q"]))
    return result or [("Filtros", "Sem filtros adicionais")]


def _provider_values(row):
    return [
        row.get("name"), row.get("risk_score"), row.get("risk_label"), row.get("impact_noc_score"),
        row.get("affected_base_score"), row.get("recurrent_score"), row.get("over_24h_score"),
        row.get("backlog_score"), row.get("links_total"), row.get("tickets"), row.get("tickets_per_link"),
        row.get("links_affected"), row.get("affected_rate_pct"), row.get("recurrent_links"),
        row.get("recurrent_rate_pct"), row.get("tickets_over_24h"), row.get("over_24h_rate_pct"),
        row.get("backlog_current"), row.get("high_critical_tickets"), row.get("procedentes"),
        row.get("improcedentes"), row.get("indeterminados"), row.get("procedure_rate", 0) * 100,
        row.get("improcedent_rate", 0) * 100,
    ]


def _ticket_values(item):
    links = item.get("links", [])
    return [
        item.get("number"), item.get("provider"), item.get("client"), item.get("project"), item.get("state"),
        " | ".join(link.get("config_item_id", "") for link in links),
        " | ".join(link.get("name", "") for link in links), item.get("status"), item.get("priority"),
        "Período" if item.get("in_period") else "Backlog atual",
        item.get("created_at"), item.get("resolved_at"), item.get("terminal_at"),
        item.get("operational_end_at"), item.get("operational_end_source"), item.get("duration_hours"),
        "Sim" if item.get("over_24h") else "Não", item.get("reason"),
        item.get("responsibility_raw"), item.get("responsibility_class"),
    ]


def build_xlsx(data, filters, providers, podium, generated_at, generated_by):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append(["ARION ATACADO GEO · NOC ATACADO"])
    summary.append(["Ranking operacional de provedores"])
    summary.append(["Período", filters["date_from"], "até", filters["date_to"]])
    summary.append(["Visão", _view_label(filters)])
    summary.append(["Ordenação", _ranking_label(filters)])
    summary.append(["Gerado em", generated_at.strftime("%d/%m/%Y %H:%M:%S %Z")])
    summary.append(["Gerado por", _safe_cell(generated_by)])
    summary.append([])
    summary.append(["Filtros ativos", "Valor"])
    for key, value in _filter_labels(filters):
        summary.append([_safe_cell(key), _safe_cell(value)])
    summary.append([])
    summary.append(["Indicadores", "Valor"])
    for label, value in (
        ("Provedores analisados", len(providers)), ("Links CMDB", data["links_total"]),
        ("Chamados válidos", data["tickets_valid"]), ("Links afetados", data["links_affected"]),
        ("Links reincidentes", data["links_recurrent"]), ("Chamados ARION >24h", data["tickets_over_24h"]),
        ("Backlog atual", data["backlog_current"]),
    ):
        summary.append([label, value])
    summary.append([])
    summary.append(["Índice ARION V2", "Impacto no NOC 30%; base afetada 20%; reincidência 20%; chamados ARION acima de 24h 20%; backlog atual 10%."])
    summary.append(["Impacto no NOC", "Componente absoluto usa 60% chamados e 40% links afetados, normalizados pelo P95 do recorte filtrado."])
    summary.append(["Responsabilidade", "Procedência é informação do fechamento OTRS e não altera o Índice ARION V2."])
    summary.append(["ARION >24h", "Indicador operacional provisório; não representa necessariamente SLA contratual."])
    summary.append(["Exposição", "Chamados por provedor podem se sobrepor; abas Links e Tickets respeitam o Data Scope."])
    summary.append([])
    summary.append(["Pódio · " + _ranking_label(filters)])
    summary.append(["Posição", "Provedor / ISP", "Índice ARION V2", "Nível", "Links", "Chamados"])
    for position, row in enumerate(podium, 1):
        summary.append([position, _safe_cell(row["name"]), row["risk_score"], row["risk_label"], row["links_total"], row["tickets"]])

    ranking = workbook.create_sheet("Ranking")
    ranking_headers = ["Pos.", "Provedor / ISP", "Índice ARION V2", "Nível", "Impacto no NOC /30", "Base afetada /20", "Reincidência /20", "ARION >24h /20", "Backlog atual /10", "Links", "Chamados", "Chamados / link", "Links afetados", "% base afetada", "Reincidentes", "Taxa de reincidência %", "Chamados >24h", "Taxa >24h %", "Backlog", "Alta / crítica", "Procedentes", "Improcedentes", "Indeterminados", "Procedência %", "Improcedentes %", "Tendência"]
    ranking.append(ranking_headers)
    for position, row in enumerate(providers, 1):
        ranking.append([position, *[_safe_cell(value) for value in _provider_values(row)], row.get("trend_label", "")])

    provider_sheet = workbook.create_sheet("Providers")
    provider_sheet.append(ranking_headers)
    for position, row in enumerate(providers, 1):
        provider_sheet.append([position, *[_safe_cell(value) for value in _provider_values(row)], row.get("trend_label", "")])

    links_sheet = workbook.create_sheet("Links")
    links_sheet.append(["Link", "ConfigItemID", "Number", "Provedor / ISP", "Cliente", "Projeto", "UF", "Chamados", "Reincidente", "ARION >24h"])
    for row in data["ranked_links_all"]:
        links_sheet.append([_safe_cell(row.get(key)) for key in ("name", "config_item_id", "number", "provider", "client", "project", "state", "tickets")] + ["Sim" if row.get("recurrent") else "Não", row.get("over_24h", 0)])

    tickets_sheet = workbook.create_sheet("Tickets")
    tickets_sheet.append([
        "Chamado", "Provedor / ISP", "Cliente", "Projeto", "UF", "ConfigItemID", "Link", "Status", "Prioridade",
        "Origem no recorte", "Criado em", "Resolvido em", "Terminal OTRS em", "Término operacional usado",
        "Origem do término operacional", "Duração (horas)", "ARION >24h", "Motivo", "Procedência raw", "Classificação",
    ])
    for item in data["ticket_export_rows"]:
        tickets_sheet.append([_safe_cell(value) for value in _ticket_values(item)])

    navy = "183445"
    pale = "EEF3F5"
    line = Side(style="thin", color="D8E1E5")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Resumo" else "A1"
        sheet.sheet_view.showGridLines = False
        if sheet.title != "Resumo" and sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
            table = Table(displayName=f"Arion{sheet.title}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        if sheet.title == "Resumo":
            sheet.column_dimensions["A"].width = 31
            sheet.column_dimensions["B"].width = 55
            sheet.column_dimensions["C"].width = 14
            sheet.column_dimensions["D"].width = 18
        else:
            for column in sheet.columns:
                letter = column[0].column_letter
                width = max(len(str(cell.value or "")) for cell in column) + 2
                sheet.column_dimensions[letter].width = min(42, max(12, width))
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.row == 1 or cell.column == 1)
                if cell.row == 1:
                    cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor=navy)
                    cell.border = Border(bottom=line)
                elif cell.column == 1 and sheet.title == "Resumo":
                    cell.font = Font(name="Aptos", bold=True, color=navy)
        sheet.row_dimensions[1].height = 28

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_pdf(data, filters, providers, podium, generated_at, generated_by):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(output, pagesize=page_size, rightMargin=12 * mm, leftMargin=12 * mm, topMargin=13 * mm, bottomMargin=14 * mm)
    palette = {"ink": colors.HexColor("#20333d"), "muted": colors.HexColor("#647782"), "line": colors.HexColor("#d9e1e5"), "pale": colors.HexColor("#eff3f5"), "accent": colors.HexColor("#287c83")}
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ArionReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=palette["ink"], alignment=TA_LEFT, spaceAfter=2))
    styles.add(ParagraphStyle(name="ArionReportSub", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=palette["muted"], spaceAfter=10))
    styles.add(ParagraphStyle(name="ArionReportSection", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=palette["ink"], spaceBefore=9, spaceAfter=5))
    styles.add(ParagraphStyle(name="ArionReportCell", parent=styles["Normal"], fontName="Helvetica", fontSize=7, leading=9, textColor=palette["ink"], wordWrap="CJK"))
    styles.add(ParagraphStyle(name="ArionReportSmall", parent=styles["Normal"], fontName="Helvetica", fontSize=7.5, leading=10, textColor=palette["muted"]))

    def p(value, style="ArionReportCell"):
        return Paragraph(escape(str(value if value is not None else "")), styles[style])

    def table(rows, widths, *, header=True):
        block = Table(rows, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT", splitByRow=1)
        commands = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LINEBELOW", (0, 0), (-1, -1), .3, palette["line"]),
        ]
        if header:
            commands += [("BACKGROUND", (0, 0), (-1, 0), palette["pale"]), ("TEXTCOLOR", (0, 0), (-1, 0), palette["ink"]), ("LINEBELOW", (0, 0), (-1, 0), .8, palette["accent"])]
        block.setStyle(TableStyle(commands))
        return block

    mode = _ranking_label(filters)
    story = [p("ARION ATACADO GEO", "ArionReportTitle"), p("NOC ATACADO · Ranking operacional de provedores", "ArionReportSub")]
    metadata = [
        [p("Período"), p(f"{filters['date_from']} a {filters['date_to']}")],
        [p("Visão"), p(_view_label(filters))], [p("Ordenação"), p(mode)],
        [p("Filtros"), p("; ".join(f"{key}: {value}" for key, value in _filter_labels(filters)))],
        [p("Gerado em"), p(generated_at.strftime("%d/%m/%Y %H:%M:%S %Z"))], [p("Gerado por"), p(generated_by)],
    ]
    story += [table(metadata, [30 * mm, 220 * mm], header=False), Spacer(1, 5 * mm), p("Indicadores do recorte", "ArionReportSection")]
    kpis = [
        [p("Provedores"), p("Links CMDB"), p("Chamados válidos"), p("Links afetados"), p("Reincidentes"), p("ARION >24h"), p("Backlog")],
        [p(len(providers)), p(data["links_total"]), p(data["tickets_valid"]), p(data["links_affected"]), p(data["links_recurrent"]), p(data["tickets_over_24h"]), p(data["backlog_current"])],
    ]
    story.append(table(kpis, [35 * mm] * 7))
    story += [p(f"Top {10 if filters['ranking_view'] == 'stable' else 3} · {mode}", "ArionReportSection")]
    top_rows = [[p("Pos."), p("Provedor / ISP"), p("Índice ARION V2"), p("Nível"), p("Links"), p("Chamados"), p("Impacto / 30")]]
    for position, row in enumerate(podium, 1):
        top_rows.append([p(position), p(row["name"]), p(f"{row['risk_score']:.1f} / 100"), p(row["risk_label"]), p(row["links_total"]), p(row["tickets"]), p(f"{row['impact_noc_score']:.1f} / 30")])
    story.append(table(top_rows, [13 * mm, 73 * mm, 28 * mm, 34 * mm, 19 * mm, 20 * mm, 25 * mm]))
    story += [p("Ranking completo · primeiros 20", "ArionReportSection")]
    headers = ["Pos.", "Provedor / ISP", "Índice V2", "Nível", "Links", "Chamados", "Chamados/link", "Impacto/30", "Base/20", "Reincid./20", ">24h/20", "Backlog/10"]
    rows = [[p(value) for value in headers]]
    for position, row in enumerate(providers[:20], 1):
        rows.append([p(position), p(row["name"]), p(f"{row['risk_score']:.1f}"), p(row["risk_label"]), p(row["links_total"]), p(row["tickets"]), p(f"{row['tickets_per_link']:.2f}"), p(f"{row['impact_noc_score']:.1f}"), p(f"{row['affected_base_score']:.1f}"), p(f"{row['recurrent_score']:.1f}"), p(f"{row['over_24h_score']:.1f}"), p(f"{row['backlog_score']:.1f}")])
    story.append(table(rows, [11 * mm, 50 * mm, 16 * mm, 24 * mm, 13 * mm, 15 * mm, 19 * mm, 17 * mm, 15 * mm, 17 * mm, 15 * mm, 18 * mm]))
    unique_tickets = {}
    for item in data["ticket_export_rows"]:
        unique_tickets.setdefault(item["external_id"], item)
    proc = sum(item.get("responsibility_class") == "PROCEDENTE" for item in unique_tickets.values())
    improc = sum(item.get("responsibility_class") == "IMPROCEDENTE" for item in unique_tickets.values())
    indeterminate = sum(item.get("responsibility_class") == "INDETERMINADO" for item in unique_tickets.values())
    story += [p("Responsabilidade observada", "ArionReportSection"), p(f"Chamados relacionados no ranking (deduplicados por ID): {len(unique_tickets)} · Procedentes: {proc} · Improcedentes: {improc} · Indeterminados: {indeterminate}. O campo vem do fechamento OTRS, não prova culpa do Provider e não altera o Índice ARION V2.", "ArionReportSmall")]
    story += [Spacer(1, 4 * mm), p("Índice ARION V2 é uma métrica operacional interna; não representa SLA contratual nem penalização financeira. O componente de impacto usa 60% de chamados e 40% de links afetados normalizados pelo P95 do recorte. ARION >24h é provisório. Procedência/Improcedência é exibida para análise e não altera o índice.", "ArionReportSmall")]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(palette["line"])
        canvas.line(12 * mm, 10 * mm, page_size[0] - 12 * mm, 10 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(palette["muted"])
        canvas.drawString(12 * mm, 6 * mm, "ARION ATACADO GEO · NOC ATACADO · Espelhos locais OTRS/CMDB · Somente leitura")
        canvas.drawRightString(page_size[0] - 12 * mm, 6 * mm, f"Página {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
