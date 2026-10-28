from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .treatment import competence_financial_summary, treatment_export_dataset


EXPORT_SCOPES = {
    "eligible": "Para tratar · Indisp. > 0",
    "any": "Qualquer multa",
    "all": "Todas as linhas",
    "treated": "Somente tratadas",
    "divergences": "Somente divergências",
    "invalid": "Somente indevidas",
}

MONEY_COLUMNS = {
    "Vlr Mensal", "Inter Claro", "Inter ARION", "Indisp Claro", "Indisp ARION",
    "Penal Claro", "Penal ARION calculado", "Penal ARION efetivo", "Valor cobrado Claro",
    "Valor calculado ARION", "Valor efetivo ARION", "Valor devido ARION",
    "Valor não cobrado", "Total Geral ARION", "Valor contestado",
}
DATE_COLUMNS = {"Abertura Claro", "Abertura ARION", "Fetec Claro", "Fetec ARION"}
EXPORT_HEADERS = (
    "REC", "Designação", "Cliente", "UF", "LP", "Fornecedor", "Contrato", "Ticket / PU Fornecedor", "Banda",
    "Abertura Claro", "Abertura ARION", "Fetec Claro", "Fetec ARION", "Tempo Claro", "TMP ARION",
    "EV Claro", "EV ARION calculado", "Vlr Mensal", "Inter Claro", "Inter ARION", "Indisp Claro",
    "Indisp ARION", "Penal Claro", "Penal ARION calculado", "Penal ARION efetivo", "Qtd/Horas", "ISP",
    "Causa Raiz", "Observação", "Valor cobrado Claro", "Valor calculado ARION", "Valor efetivo ARION",
    "Valor devido ARION", "Valor não cobrado", "Total Geral ARION", "Valor contestado", "Redução %", "Decisão", "Status",
)


def _neutralize_formula(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _decimal_value(value):
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _excel_value(column, value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, date):
        return value
    if isinstance(value, Decimal) or isinstance(value, (int, float, bool)):
        return value
    if column in DATE_COLUMNS and isinstance(value, str):
        parsed_datetime = parse_datetime(value)
        if parsed_datetime:
            return timezone.localtime(parsed_datetime).replace(tzinfo=None) if timezone.is_aware(parsed_datetime) else parsed_datetime
        parsed_date = parse_date(value)
        if parsed_date:
            return parsed_date
    if column in MONEY_COLUMNS:
        parsed_decimal = _decimal_value(value)
        if parsed_decimal is not None:
            return parsed_decimal
    return _neutralize_formula(str(value))


def _style_sheet(sheet, *, freeze="A2"):
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="26343A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color="263238")
            cell.alignment = Alignment(vertical="top")
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(cell.value, Decimal):
                cell.number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max((len(str(cell.value or "")) for cell in column_cells[:200]), default=8)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 32)


def build_treatment_workbook(competence, *, scope="eligible"):
    if scope not in EXPORT_SCOPES:
        raise ValueError("Escopo de exportação inválido.")

    workbook = Workbook()
    treatment_sheet = workbook.active
    treatment_sheet.title = "Tratativa ARION"
    rows = treatment_export_dataset(competence, scope=scope)
    headers = None
    for record in rows:
        if headers is None:
            headers = list(record.keys())
            treatment_sheet.append(headers)
        treatment_sheet.append([_excel_value(key, record.get(key)) for key in headers])
    if headers is None:
        headers = list(EXPORT_HEADERS)
        treatment_sheet.append(headers)
    _style_sheet(treatment_sheet, freeze="C2")

    summary_sheet = workbook.create_sheet("Resumo")
    summary_sheet.append(["Indicador", "Valor"])
    summary = competence_financial_summary(competence)
    values = (
        ("Competência", competence.period_label),
        ("Escopo exportado", EXPORT_SCOPES[scope]),
        ("Total Original", summary["original_imported"]),
        ("Com multa", summary["with_penalty"]),
        ("Tratadas", summary["due_count"] + summary["reduced_count"] + summary["invalid_count"]),
        ("Pendentes", summary["pending"]),
        ("Devidas", summary["due_count"]),
        ("Reduzidas", summary["reduced_count"]),
        ("Indevidas", summary["invalid_count"]),
        ("Total Claro", summary["total_claro"]),
        ("Total ARION", summary["total_arion"]),
        ("Valor contestado", summary["contested"]),
        ("Redução %", summary["reduction_percent"]),
    )
    for label, value in values:
        summary_sheet.append([label, value])
    _style_sheet(summary_sheet, freeze="A2")
    for cell in summary_sheet["B"]:
        if cell.row in (11, 12, 13):
            cell.number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
        elif cell.row == 14:
            cell.number_format = '0.00"%"'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
