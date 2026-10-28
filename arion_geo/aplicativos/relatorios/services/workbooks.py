"""Safe inspection and row normalization for Claro Multas Excel workbooks."""

from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata
import zipfile

from django.core.exceptions import ValidationError

from aplicativos.relatorios.validators import (
    ALLOWED_REPORT_EXTENSIONS,
    MAX_REPORT_FILE_SIZE,
    safe_original_filename,
)


HEADER_SCAN_ROWS = 30
PREVIEW_ROW_LIMIT = 25
MAX_WORKBOOK_ROWS = 100_000
MAX_WORKBOOK_COLUMNS = 256
MAX_WORKBOOK_SHEETS = 50
MAX_UNCOMPRESSED_XLSX_BYTES = 100 * 1024 * 1024
MIN_HEADER_MATCHES = 8
REQUIRED_IMPORT_FIELDS = frozenset(("rec", "designation"))

FIELD_LABELS = {
    "rec": "REC",
    "designation": "Designação",
    "client": "Cliente",
    "uf": "UF",
    "lp": "LP",
    "provider": "Fornecedor",
    "contract": "Contrato",
    "provider_unit_price": "PU Fornecedor",
    "bandwidth": "Banda",
    "opening": "Abertura",
    "fetec": "Fetec",
    "duration": "Tempo",
    "events": "Ev.",
    "monthly_value": "Vlr Mensal",
    "interconnection_fine": "Multas/Inter.",
    "penalty_fine": "Multas/Penal",
    "unavailability_fine": "Multas/Indisp.",
    "hours_quantity": "Qtd/Horas",
}

TEXT_FIELDS = frozenset((
    "rec", "designation", "client", "uf", "lp", "provider", "contract", "bandwidth",
))
REC_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/-]{0,79}$", re.ASCII)
MONEY_FIELDS = frozenset((
    "monthly_value", "interconnection_fine", "penalty_fine", "unavailability_fine",
))
DATE_FIELDS = frozenset(("opening", "fetec"))
DECIMAL_FIELDS = frozenset((
    "provider_unit_price", "events", "monthly_value", "interconnection_fine",
    "penalty_fine", "unavailability_fine", "hours_quantity",
))
PREVIEW_DECIMAL_FIELDS = DECIMAL_FIELDS - {"provider_unit_price"}
SEARCH_FIELDS = tuple(FIELD_LABELS)


class WorkbookImportError(Exception):
    """An expected upload or workbook validation failure with a safe message."""


def normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


_HEADER_ALIASES = {
    normalize_header(alias): field
    for field, aliases in {
        "rec": ("REC", "Nº REC", "Número REC"),
        "designation": ("Designação", "Designacao"),
        "client": ("Cliente",),
        "uf": ("UF", "Estado"),
        "lp": ("LP",),
        "provider": ("Fornecedor",),
        "contract": ("Contrato",),
        "provider_unit_price": ("PU Fornecedor", "PU do Fornecedor"),
        "bandwidth": ("Banda",),
        "opening": ("Abertura", "Data de Abertura"),
        "fetec": ("Fetec", "FETEC"),
        "duration": ("Tempo",),
        "events": ("Ev", "Ev."),
        "monthly_value": ("Vlr Mensal", "Valor Mensal"),
        "interconnection_fine": ("Multas/Inter.", "Multas Inter", "Multas Interconexão"),
        "penalty_fine": ("Multas/Penal", "Multas Penal"),
        "unavailability_fine": ("Multas/Indisp.", "Multas Indisp", "Multas Indisponibilidade"),
        "hours_quantity": ("Qtd/Horas", "Qtd Horas", "Quantidade de Horas"),
    }.items()
    for alias in aliases
}

_CANDIDATE_FIELDS = frozenset((
    "rec", "designation", "client", "uf", "provider", "opening", "fetec", "duration",
    "monthly_value", "interconnection_fine", "unavailability_fine",
))


def _safe_display_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return str(value)


def _json_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, (str, bool)) or value is None:
        return value
    return str(value)


def _cell_record(value, *, formula=None, cached_value=None, number_format=""):
    effective_value = cached_value if formula and cached_value is not None else (None if formula else value)
    if formula:
        kind = "formula"
    elif effective_value is None:
        kind = "empty"
    elif isinstance(effective_value, bool):
        kind = "boolean"
    elif isinstance(effective_value, datetime):
        kind = "datetime"
    elif isinstance(effective_value, date):
        kind = "date"
    elif isinstance(effective_value, time):
        kind = "time"
    elif isinstance(effective_value, timedelta):
        kind = "duration"
    elif isinstance(effective_value, (int, float, Decimal)):
        kind = "number"
    else:
        kind = "text"
    record = {
        "type": kind,
        "value": _json_value(effective_value),
        "display": _safe_display_value(effective_value) if effective_value is not None else (str(formula) if formula else ""),
    }
    if formula:
        record["formula"] = str(formula)
        record["cached_value"] = _json_value(cached_value)
    if number_format:
        record["number_format"] = str(number_format)[:120]
    return record, effective_value


def _serialize_rows_for_hash(source_data):
    return json.dumps(source_data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _parse_decimal(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    if isinstance(value, (int, float)):
        try:
            parsed = Decimal(str(value))
            return parsed if parsed.is_finite() else None
        except InvalidOperation:
            return None
    text = str(value).replace("\u00a0", " ").strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("() ").replace("R$", "").replace("r$", "").replace(" ", "")
    if not text or any(character.isalpha() for character in text) or any(character in text for character in "=*%@"):
        return None
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif text.count(".") > 1:
        last = text.rsplit(".", 1)[1]
        text = text.replace(".", "") if len(last) == 3 else text[:text.rfind(".")].replace(".", "") + "." + last
    elif "." in text and len(text.rsplit(".", 1)[1]) == 3:
        text = text.replace(".", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        return None
    try:
        result = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return -result if negative else result


def canonical_ticket_number(value):
    """Return an exact decimal ticket identifier as digits, never scientific notation."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
        except (InvalidOperation, ValueError):
            return ""
    else:
        text = str(value).replace("\u00a0", " ").strip().replace(" ", "")
        if not text:
            return ""
        # Excel exports scientific notation using either decimal separator.
        if "e" in text.casefold():
            if "," in text and "." not in text:
                text = text.replace(",", ".")
            elif "," in text:
                text = text.replace(".", "").replace(",", ".")
            try:
                number = Decimal(text)
            except (InvalidOperation, ValueError):
                return ""
        elif re.fullmatch(r"[0-9]+", text):
            return text if 8 <= len(text) <= 80 else ""
        else:
            # Accept an integral value serialized by a numeric workbook cell,
            # including a locale-formatted value, without ever passing through float.
            number = _parse_decimal(text)
            if number is None:
                return ""
    if not number.is_finite() or number != number.to_integral_value():
        return ""
    ticket_number = format(number.quantize(Decimal("1")), "f")
    return ticket_number if re.fullmatch(r"[0-9]{8,80}", ticket_number) else ""


def _canonical_rec(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            numeric = Decimal(str(value))
        except InvalidOperation:
            return None
        if not numeric.is_finite() or numeric != numeric.to_integral_value():
            return None
        text = format(numeric.quantize(Decimal("1")), "f")
    else:
        text = _safe_display_value(value).strip()
    return text if REC_PATTERN.fullmatch(text) else None


def _is_real_data_row(cells, mapping):
    rec_index = next((int(index) for index, field in mapping.items() if field == "rec"), None)
    if rec_index is None:
        # Keep an inspectable preview for files that need manual column mapping.
        return any(record.get("value") not in (None, "") or record.get("formula") for record, _value in cells)
    if rec_index >= len(cells):
        return False
    record, value = cells[rec_index]
    if record.get("type") == "formula" and record.get("cached_value") is None:
        return False
    return _canonical_rec(value) is not None


def _format_preview_decimal(value, *, fixed_places=False):
    if value is None or value == "":
        return None
    try:
        parsed = value if isinstance(value, Decimal) else Decimal(str(value))
        if not parsed.is_finite():
            return None
    except (InvalidOperation, ValueError):
        parsed = _parse_decimal(value)
        if parsed is None:
            return None
    rounded = parsed.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if not fixed_places and rounded == rounded.to_integral_value():
        return format(rounded.quantize(Decimal("1")), "f")
    rendered = format(rounded, ",.2f")
    return rendered.replace(",", "\u0000").replace(".", ",").replace("\u0000", ".")


def _preview_cells(source_data, canonical_data, mapping):
    preview = [dict(record) for record in source_data]
    for index, field in mapping.items():
        if field not in PREVIEW_DECIMAL_FIELDS:
            continue
        cell_index = int(index)
        if cell_index >= len(preview):
            continue
        formatted = _format_preview_decimal(canonical_data.get(field), fixed_places=field in MONEY_FIELDS)
        if formatted is not None:
            preview[cell_index]["display"] = formatted
    return preview


def _parse_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, time):
        return value
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        pass
    for pattern in (
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%d/%m/%y %H:%M:%S", "%d/%m/%y %H:%M", "%d/%m/%y",
    ):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


class WorkbookReader:
    def __init__(self, data, extension):
        self.data = data
        self.extension = str(extension or "").casefold().lstrip(".")
        self.workbook = None
        self.cached_workbook = None
        self._xls_module = None
        if self.extension not in ALLOWED_REPORT_EXTENSIONS:
            raise WorkbookImportError("Formato não permitido. Envie um arquivo XLS ou XLSX.")
        if not data:
            raise WorkbookImportError("O arquivo está vazio.")
        if len(data) > MAX_REPORT_FILE_SIZE:
            raise WorkbookImportError("O arquivo ultrapassa o limite de 25 MB.")
        try:
            if data.startswith(b"PK"):
                self._validate_xlsx_container()
                # Alguns arquivos operacionais mantêm .xls no nome apesar do
                # conteúdo OOXML. A assinatura confiável escolhe o parser.
                self.extension = "xlsx"
                from openpyxl import load_workbook
                self.workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
                self.cached_workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
            elif data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                if self.extension == "xlsx":
                    raise WorkbookImportError("A extensão XLSX não corresponde ao conteúdo XLS enviado.")
                self.extension = "xls"
                import xlrd
                self._xls_module = xlrd
                self.workbook = xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True)
            else:
                raise WorkbookImportError("O conteúdo não corresponde a uma planilha XLS ou XLSX válida.")
        except WorkbookImportError:
            self.close()
            raise
        except ImportError as exc:
            self.close()
            raise WorkbookImportError("O leitor necessário para este formato não está instalado.") from exc
        except Exception as exc:
            self.close()
            raise WorkbookImportError("Não foi possível ler a planilha. Verifique se o arquivo não está corrompido ou protegido por senha.") from exc
        self.sheet_names = self.workbook.sheetnames if self.extension == "xlsx" else list(self.workbook.sheet_names())
        if not self.sheet_names:
            self.close()
            raise WorkbookImportError("O arquivo não contém abas para leitura.")
        if len(self.sheet_names) > MAX_WORKBOOK_SHEETS:
            self.close()
            raise WorkbookImportError("O arquivo contém abas demais para inspeção segura.")

    def _validate_xlsx_container(self):
        if not self.data.startswith(b"PK"):
            raise WorkbookImportError("O conteúdo não corresponde a um arquivo XLSX válido.")
        try:
            with zipfile.ZipFile(BytesIO(self.data)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise WorkbookImportError("O arquivo não contém uma estrutura XLSX válida.")
                if any(name.casefold().endswith("vbaproject.bin") for name in names):
                    raise WorkbookImportError("Arquivos com macros não são aceitos. Envie uma planilha XLSX sem macros.")
                if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_XLSX_BYTES:
                    raise WorkbookImportError("O conteúdo compactado excede o limite seguro para leitura.")
                if archive.testzip() is not None:
                    raise WorkbookImportError("O arquivo XLSX está corrompido.")
        except WorkbookImportError:
            raise
        except (OSError, zipfile.BadZipFile, RuntimeError):
            raise WorkbookImportError("O arquivo XLSX está corrompido ou incompleto.") from None

    def dimensions(self, sheet_name):
        if self.extension == "xlsx":
            sheet = self.workbook[sheet_name]
            rows, columns = sheet.max_row or 0, sheet.max_column or 0
        else:
            sheet = self.workbook.sheet_by_name(sheet_name)
            rows, columns = sheet.nrows, sheet.ncols
        if rows > MAX_WORKBOOK_ROWS:
            raise WorkbookImportError("A aba selecionada excede o limite de linhas para importação.")
        if columns > MAX_WORKBOOK_COLUMNS:
            raise WorkbookImportError("A aba selecionada excede o limite de 256 colunas.")
        return rows, columns

    def iter_cell_rows(self, sheet_name):
        max_row, max_column = self.dimensions(sheet_name)
        if self.extension == "xls":
            sheet = self.workbook.sheet_by_name(sheet_name)
            for row_index in range(max_row):
                cells = []
                for column_index in range(max_column):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == self._xls_module.XL_CELL_DATE:
                        try:
                            value = self._xls_module.xldate_as_datetime(value, self.workbook.datemode)
                        except (ValueError, self._xls_module.XLDateError):
                            value = cell.value
                    if cell.ctype == self._xls_module.XL_CELL_ERROR:
                        value = f"#ERRO {cell.value}"
                    record, effective = _cell_record(value)
                    cells.append((record, effective))
                yield row_index + 1, cells
            return

        formula_sheet = self.workbook[sheet_name]
        cached_sheet = self.cached_workbook[sheet_name]
        formula_rows = formula_sheet.iter_rows(min_row=1, max_col=max_column, values_only=False)
        cached_rows = cached_sheet.iter_rows(min_row=1, max_col=max_column, values_only=False)
        for row_number, (formula_cells, cached_cells) in enumerate(zip(formula_rows, cached_rows), start=1):
            cells = []
            for formula_cell, cached_cell in zip(formula_cells, cached_cells):
                is_formula = formula_cell.data_type == "f"
                formula = formula_cell.value if is_formula else None
                cached_value = cached_cell.value if is_formula else None
                value = None if is_formula else formula_cell.value
                record, effective = _cell_record(
                    value,
                    formula=formula,
                    cached_value=cached_value,
                    number_format=getattr(formula_cell, "number_format", ""),
                )
                cells.append((record, effective))
            yield row_number, cells

    def excel_serial_to_datetime(self, value):
        try:
            if self.extension == "xlsx":
                from openpyxl.utils.datetime import from_excel
                return from_excel(value, self.workbook.epoch)
            return self._xls_module.xldate_as_datetime(value, self.workbook.datemode)
        except (ValueError, TypeError, self._xls_module.XLDateError if self._xls_module else ValueError):
            return None

    def close(self):
        for workbook in (self.workbook, self.cached_workbook):
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass


def _row_header_info(cells):
    headers = [record.get("display", "") for record, _value in cells]
    mapping = {}
    matched = set()
    for index, header in enumerate(headers):
        field = _HEADER_ALIASES.get(normalize_header(header))
        if field and field not in matched:
            mapping[str(index)] = field
            matched.add(field)
    candidate = (
        len(matched.intersection(_CANDIDATE_FIELDS)) >= MIN_HEADER_MATCHES
        and REQUIRED_IMPORT_FIELDS.issubset(matched)
    )
    return headers, mapping, matched, candidate


def _combine_header_rows(upper_cells, lower_cells):
    """Resolve vertically merged labels and the Multas/Qtd grouped headings."""
    width = max(len(upper_cells), len(lower_cells))
    combined = []
    active_group = ""
    group_columns_left = 0
    for index in range(width):
        upper_record = upper_cells[index][0] if index < len(upper_cells) else {"display": ""}
        lower_record = lower_cells[index][0] if index < len(lower_cells) else {"display": ""}
        upper = str(upper_record.get("display", "") or "").strip()
        lower = str(lower_record.get("display", "") or "").strip()
        upper_normalized = normalize_header(upper)
        if upper_normalized in {"multas", "multa"}:
            active_group, group_columns_left = upper, 2
        elif upper_normalized in {"qtd", "quantidade"}:
            active_group, group_columns_left = upper, 0
        elif upper_normalized in _HEADER_ALIASES:
            active_group, group_columns_left = "", 0

        if active_group and lower and (upper_normalized in {"multas", "multa"} or not upper) and group_columns_left >= 0:
            label = f"{active_group}/{lower}"
        elif upper_normalized in {"multas", "multa", "qtd", "quantidade"} and lower:
            label = f"{upper}/{lower}"
        elif lower:
            label = lower
        elif upper_normalized in _HEADER_ALIASES:
            label = upper
        else:
            label = ""
        if group_columns_left > 0 and upper_normalized not in {"multas", "multa"}:
            group_columns_left -= 1
            if group_columns_left == 0:
                active_group = ""
        if upper_normalized in {"qtd", "quantidade"}:
            active_group = ""
        record, effective = _cell_record(label)
        combined.append((record, effective))
    return combined


def _read_header_cells(reader, sheet_name, header_row):
    upper_cells = []
    lower_cells = []
    for row_number, cells in reader.iter_cell_rows(sheet_name):
        if row_number == header_row - 1:
            upper_cells = cells
        elif row_number == header_row:
            lower_cells = cells
            break
    if header_row > 1 and upper_cells:
        return _combine_header_rows(upper_cells, lower_cells)
    return lower_cells


def _detect_sheet_header(reader, sheet_name):
    best = {"row": None, "headers": [], "mapping": {}, "matched_fields": [], "match_count": 0, "candidate": False}
    previous_cells = []
    for row_number, cells in reader.iter_cell_rows(sheet_name):
        if row_number > HEADER_SCAN_ROWS:
            break
        choices = [(cells, row_number)]
        if previous_cells:
            choices.append((_combine_header_rows(previous_cells, cells), row_number))
        for candidate_cells, logical_header_row in choices:
            headers, mapping, matched, candidate = _row_header_info(candidate_cells)
            count = len(matched.intersection(_CANDIDATE_FIELDS))
            if count > best["match_count"]:
                best = {
                    "row": logical_header_row,
                    "headers": headers,
                    "mapping": mapping,
                    "matched_fields": sorted(matched),
                    "match_count": count,
                    "candidate": candidate,
                }
        previous_cells = cells
    return best


def _canonical_value(reader, value, field):
    if value is None:
        return None
    if field in TEXT_FIELDS:
        if field == "rec":
            return _canonical_rec(value)
        text = _safe_display_value(value)
        return text.strip() if text else None
    if field in DATE_FIELDS:
        parsed = _parse_datetime(value)
        if parsed is None and isinstance(value, (int, float, Decimal)):
            parsed = reader.excel_serial_to_datetime(value)
        if parsed is None:
            return _safe_display_value(value).strip() or None
        return parsed.isoformat()
    if field in DECIMAL_FIELDS:
        return _parse_decimal(value)
    if isinstance(value, str):
        return value.strip() or None
    return value


def _build_source_row(reader, row_number, cells, mapping):
    source_data = [record for record, _value in cells]
    canonical_data = {}
    for index, field in mapping.items():
        column_index = int(index)
        if column_index >= len(cells):
            continue
        record, value = cells[column_index]
        if record.get("type") == "formula" and record.get("cached_value") is None:
            canonical_data[field] = None
        else:
            canonical_value = _canonical_value(reader, value, field)
            canonical_data[field] = format(canonical_value, "f") if isinstance(canonical_value, Decimal) else canonical_value
            if field == "provider_unit_price":
                canonical_data["pu_fornecedor_ticket"] = canonical_ticket_number(value)
    row_hash = sha256(_serialize_rows_for_hash(source_data)).hexdigest()
    return {
        "source_row_number": row_number,
        "source_data": source_data,
        "canonical_data": canonical_data,
        "row_hash": row_hash,
    }


def inspect_workbook(data, extension, *, selected_sheet=None, header_row=None, preview_limit=PREVIEW_ROW_LIMIT):
    reader = WorkbookReader(data, extension)
    try:
        sheets = []
        header_by_sheet = {}
        for sheet_name in reader.sheet_names:
            rows, columns = reader.dimensions(sheet_name)
            header = _detect_sheet_header(reader, sheet_name)
            header_by_sheet[sheet_name] = header
            sheets.append({
                "name": sheet_name,
                "row_count": rows,
                "column_count": columns,
                "header_row": header["row"],
                "match_count": header["match_count"],
                "matched_fields": header["matched_fields"],
                "candidate": header["candidate"],
            })
        candidates = [sheet["name"] for sheet in sheets if sheet["candidate"]]
        auto_selected = len(candidates) == 1
        if selected_sheet:
            if selected_sheet not in reader.sheet_names:
                raise WorkbookImportError("A aba escolhida não existe no arquivo.")
            selected = selected_sheet
        elif auto_selected:
            selected = candidates[0]
        elif not candidates and len(reader.sheet_names) == 1:
            selected = reader.sheet_names[0]
        else:
            selected = None

        result = {
            "sheets": sheets,
            "candidate_sheets": candidates,
            "selected_sheet": selected,
            "auto_selected": auto_selected and selected == candidates[0],
            "header_detected": False,
            "header_row": None,
            "headers": [],
            "mapping": {},
            "source_row_count": 0,
            "physical_row_count": 0,
            "preview_rows": [],
            "preview_limit": preview_limit,
        }
        if selected is None:
            result["warning"] = "Selecione uma aba para inspecionar a estrutura do arquivo."
            return result

        detected = header_by_sheet[selected]
        if header_row is None:
            selected_header_row = detected["row"] or 1
        else:
            try:
                selected_header_row = int(header_row)
            except (TypeError, ValueError):
                raise WorkbookImportError("Informe uma linha de cabeçalho válida.") from None
        rows, _columns = reader.dimensions(selected)
        if selected_header_row < 1 or selected_header_row > min(rows or 1, HEADER_SCAN_ROWS):
            raise WorkbookImportError(f"A linha de cabeçalho deve estar entre 1 e {min(rows or 1, HEADER_SCAN_ROWS)}.")

        header_cells = _read_header_cells(reader, selected, selected_header_row)
        headers, default_mapping, matched, candidate = _row_header_info(header_cells)
        result.update(
            header_row=selected_header_row,
            header_detected=bool(candidate and selected_header_row == detected["row"]),
            headers=headers,
            mapping=default_mapping,
            source_row_count=0,
            physical_row_count=max(0, rows - selected_header_row),
        )
        if header_row is not None:
            result["mapping"] = default_mapping
        if result["header_detected"]:
            result["warning"] = "Cabeçalho compatível identificado automaticamente. Confira o mapeamento antes de importar."
        else:
            result["warning"] = "Nenhum cabeçalho compatível foi confirmado nesta linha. Selecione a linha correta e ajuste o mapeamento."

        source_count = 0
        previews = []
        for row_number, cells in reader.iter_cell_rows(selected):
            if row_number <= selected_header_row:
                continue
            if not _is_real_data_row(cells, default_mapping):
                continue
            source_count += 1
            if len(previews) < preview_limit:
                row_data = _build_source_row(reader, row_number, cells, default_mapping)
                previews.append({
                    "row_number": row_number,
                    "cells": row_data["source_data"],
                    "preview_cells": _preview_cells(row_data["source_data"], row_data["canonical_data"], default_mapping),
                    "canonical_data": row_data["canonical_data"],
                })
        result["source_row_count"] = source_count
        result["preview_rows"] = previews
        if not source_count:
            result["warning"] += " A aba não contém linhas preenchidas abaixo do cabeçalho."
        return result
    finally:
        reader.close()


@contextmanager
def iter_import_rows(data, extension, *, sheet_name, header_row, mapping):
    reader = WorkbookReader(data, extension)
    try:
        if sheet_name not in reader.sheet_names:
            raise WorkbookImportError("A aba escolhida não existe no arquivo.")
        try:
            header_row = int(header_row)
        except (TypeError, ValueError):
            raise WorkbookImportError("Informe uma linha de cabeçalho válida.") from None
        rows, _columns = reader.dimensions(sheet_name)
        if header_row < 1 or header_row > min(rows or 1, HEADER_SCAN_ROWS):
            raise WorkbookImportError("A linha de cabeçalho selecionada está fora do intervalo permitido.")

        normalized_mapping = {}
        seen_fields = set()
        for source_index, field in mapping.items():
            try:
                index = int(source_index)
            except (TypeError, ValueError):
                raise WorkbookImportError("O mapeamento de colunas é inválido.") from None
            if field not in FIELD_LABELS or index < 0 or index >= MAX_WORKBOOK_COLUMNS:
                raise WorkbookImportError("O mapeamento de colunas é inválido.")
            if field in seen_fields:
                raise WorkbookImportError("Cada campo interno pode ser associado a apenas uma coluna.")
            normalized_mapping[str(index)] = field
            seen_fields.add(field)
        if not REQUIRED_IMPORT_FIELDS.issubset(seen_fields):
            raise WorkbookImportError("Associe ao menos as colunas REC e Designação antes de confirmar.")

        header_cells = _read_header_cells(reader, sheet_name, header_row)
        if not header_cells:
            raise WorkbookImportError("A linha de cabeçalho não pôde ser lida.")
        headers, _defaults, _matched, _candidate = _row_header_info(header_cells)
        if any(int(index) >= len(headers) for index in normalized_mapping):
            raise WorkbookImportError("O mapeamento contém uma coluna que não existe na aba escolhida.")

        def source_rows():
            for row_number, cells in reader.iter_cell_rows(sheet_name):
                if row_number <= header_row:
                    continue
                if not _is_real_data_row(cells, normalized_mapping):
                    continue
                yield _build_source_row(reader, row_number, cells, normalized_mapping)

        yield headers, source_rows()
    finally:
        reader.close()


def prepare_uploaded_workbook(uploaded_file):
    original_name = str(getattr(uploaded_file, "name", "") or "")
    extension = Path(original_name.replace("\\", "/")).suffix.casefold().lstrip(".")
    if extension not in ALLOWED_REPORT_EXTENSIONS:
        raise WorkbookImportError("Formato não permitido. Envie somente arquivos XLS ou XLSX.")
    if getattr(uploaded_file, "size", 0) > MAX_REPORT_FILE_SIZE:
        raise WorkbookImportError("O arquivo ultrapassa o limite de 25 MB.")
    chunks = []
    total = 0
    try:
        uploaded_file.seek(0)
    except (AttributeError, OSError):
        pass
    for chunk in uploaded_file.chunks():
        total += len(chunk)
        if total > MAX_REPORT_FILE_SIZE:
            raise WorkbookImportError("O arquivo ultrapassa o limite de 25 MB.")
        chunks.append(chunk)
    data = b"".join(chunks)
    if not data:
        raise WorkbookImportError("O arquivo está vazio.")
    inspection = inspect_workbook(data, extension)
    return {
        "data": data,
        "extension": extension,
        "filename": safe_original_filename(original_name),
        "file_size": total,
        "sha256": sha256(data).hexdigest(),
        "inspection": inspection,
    }
