import base64
from datetime import datetime
from decimal import Decimal
import tempfile
import xml.etree.ElementTree as ET
import zlib
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from aplicativos.auditoria.models import AuditEvent

from .models import PenaltyCompetence, PenaltySourceRow, PenaltyTreatmentRow
from .services.imports import confirm_competence_import
from .services.treatment import claro_treatment_candidates
from .services.workbooks import (
    WorkbookImportError,
    _parse_decimal,
    canonical_ticket_number,
    inspect_workbook,
    iter_import_rows,
    prepare_uploaded_workbook,
)
from .validators import MAX_REPORT_FILE_SIZE, safe_original_filename


HEADERS = [
    "REC", "Designação", "Cliente", "UF", "LP", "Fornecedor", "Contrato", "PU Fornecedor",
    "Banda", "Abertura", "Fetec", "Tempo", "Ev.", "Vlr Mensal", "Multas/Inter.",
    "Multas/Penal", "Multas/Indisp.", "Qtd/Horas",
]
VALUES = [
    "REC-DEMO-01", "Link demonstrativo", "Cliente de teste", "SP", "LP-1", "Fornecedor de teste",
    "Contrato-1", "R$ 10,00", "100 Mbps", datetime(2026, 8, 3, 10, 15), datetime(2026, 8, 4, 11, 30),
    "01:15:00", 2, "R$ 1.234,56", 12.5, 0, 56.78, None,
]

# XLS legado mínimo (duas colunas e uma linha), gerado apenas para os testes.
XLS_FIXTURE_ZLIB_B64 = (
    "eNrtWE1oE0EU/maTND80aRJToRVKKFi1rdDiRZB2qzb2oqXWglIE3TZBQuquJEHQi9WaoyB4UrwUevFS9eIPetCbB6GiB0EQ"
    "Ez16EhQ8tFnfvOxKqj00oPWH+Zb55u2bNztvd9683dkXS7Hy/N32Cn7AIDyo2kE01ekElaB7EgW127YU3TpAxVb4pxAM0EQ2"
    "+fAo/Nwv51DOdwUa7nifEgPvqRzHGYxaZia5gdjHPhhC+jBALHCTNBG0sVdx5mnmTcy32fIx8xBrrjAPkG1ZTGJJH+3e7UTx"
    "Ma2T2yKQ173Pfd6wph+teCaj+MJVUbP1YW8+a8z8nQ0d3mYsgOZtJGNm8sZMGQmawAV8sZPAZ3elPkkq/cbqBUj/dbXev4b+"
    "muYFZmGf5AAvIYx3Htniw7CRtgrL2AMy4EL68dT+EDCcKWRPmYauW5RuJ1JHJnb29fXHgYNZM5dMZ05bZqGYN4rZs1ZIpmte"
    "3tFVyzvMYd9MnEYLyzEO/igl8OVbn14emhrTT7BmllN6LfFvlb7CxkXZgzpHuEX7zt3MPcyX+KpbWG5nTtA9Ud011uoIB+bY5"
    "jK3dtE4uxiv9G118naSSx8PP+gofdB3kLw4UjmfWHytz6OTXkRp6i+POfSKXnHjusRD3a2FkyTeMrf9lDACWtTx3Xbebi1YQY"
    "jFGHPtTD4d4diLNewF22uOvWB7OYmDWhz3+PEM1b1CQ1BQUFBQUFBQUFD4gxDO57vH2Wf4nE2C3/mvs0Klqn6T/LcYh0VHkTam"
    "KZhU53GuofjZDJ9wryXW2cf9XyhxlEbPI4cp9iPXcPzSJk/U38+6O0Z/3RJqdPxqI37+5vG/AYsMyec="
)


def make_xlsx(*, header_row=2, headers=None, values=None, sheet_names=("Penalidades",), row_count=1, formula=False):
    headers = list(headers or HEADERS)
    values = list(values or VALUES)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in sheet_names:
        sheet = workbook.create_sheet(sheet_name)
        if header_row > 1:
            sheet.cell(1, 1, "Relatório demonstrativo")
        for column, header in enumerate(headers, start=1):
            sheet.cell(header_row, column, header)
        for row_offset in range(row_count):
            row_number = header_row + row_offset + 1
            current_values = list(values)
            if row_offset:
                current_values[0] = f"REC-DEMO-{row_offset + 1:04d}"
            for column, value in enumerate(current_values, start=1):
                if value is not None:
                    sheet.cell(row_number, column, value)
        if formula:
            formula_column = len(headers) + 1
            sheet.cell(header_row, formula_column, "Coluna adicional")
            sheet.cell(header_row + 1, formula_column, "=1+1")
    from io import BytesIO
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def xls_fixture():
    return zlib.decompress(base64.b64decode(XLS_FIXTURE_ZLIB_B64))


def make_grouped_xlsx():
    from io import BytesIO
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Penalidades"
    for column, header in enumerate(HEADERS[:14], start=1):
        sheet.cell(1, column, header)
        sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
    sheet.cell(1, 15, "Multas")
    sheet.merge_cells("O1:Q1")
    sheet.cell(1, 18, "Qtd")
    for column, header in enumerate(("Inter.", "Penal", "Indisp.", "Horas"), start=15):
        sheet.cell(2, column, header)
    for column, value in enumerate(VALUES, start=1):
        if value is not None:
            sheet.cell(3, column, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_cached_formula_cut_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Penalidades"
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:B2")
    sheet["A1"] = "REC"
    sheet["B1"] = "Designação"
    sheet.merge_cells("C1:E1")
    sheet["C1"] = "Multas"
    sheet["C2"] = "Inter."
    sheet["D2"] = "Penal"
    sheet["E2"] = "Indisp."
    rows = (
        (603204, "Link A", 52.41, 0, "=+N3*0.05"),
        (601945, "Link B", 34.30, 0, "=+N4*0.05"),
        (601000, "Link C", 0, 0, 0),
        (601001, "Link D", 0, 0, "-"),
    )
    for row_number, values in enumerate(rows, start=3):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
    sheet["A7"] = "217 Registros"
    sheet["A8"] = "Multas Interrupção R$ 4.035,06"
    output = BytesIO()
    workbook.save(output)

    with ZipFile(BytesIO(output.getvalue()), "r") as source_zip:
        members = [(member, source_zip.read(member.filename)) for member in source_zip.infolist()]
    cached_values = {"E3": "16.55", "E4": "17.15"}
    namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cached_sheet = BytesIO()
    with ZipFile(cached_sheet, "w", ZIP_DEFLATED) as result_zip:
        for member, content in members:
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                for reference, value in cached_values.items():
                    cell = root.find(f".//main:c[@r='{reference}']", namespace)
                    cached = cell.find("main:v", namespace)
                    if cached is None:
                        cached = ET.SubElement(cell, f"{{{namespace['main']}}}v")
                    cached.text = value
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            result_zip.writestr(member, content)
    return cached_sheet.getvalue()


class WorkbookInspectionTests(TestCase):
    def test_scientific_pu_fornecedor_is_canonicalized_without_rounding(self):
        expected = "2026100136000300"
        self.assertEqual(canonical_ticket_number("2.026100136000300E+15"), expected)
        self.assertEqual(canonical_ticket_number("2,026100136000300E+15"), expected)
        self.assertEqual(canonical_ticket_number(2026100136000300.0), expected)
        self.assertEqual(canonical_ticket_number("2026100136000300.0"), expected)
        self.assertEqual(canonical_ticket_number("2026100136000300.5"), "")

        values = list(VALUES)
        values[7] = 2.026100136000300E+15
        data = make_xlsx(values=values)
        result = inspect_workbook(data, "xlsx")
        from aplicativos.relatorios.services.workbooks import prepare_uploaded_workbook
        with iter_import_rows(
            data, "xlsx", sheet_name="Penalidades", header_row=result["header_row"],
            mapping=result["mapping"],
        ) as (_headers, rows):
            row = next(rows)
        self.assertEqual(row["canonical_data"]["pu_fornecedor_ticket"], expected)

    def test_xlsx_detects_header_on_second_row_and_exact_source_columns(self):
        result = inspect_workbook(make_xlsx(), "xlsx")
        self.assertEqual(result["selected_sheet"], "Penalidades")
        self.assertTrue(result["header_detected"])
        self.assertEqual(result["header_row"], 2)
        self.assertEqual(result["headers"], HEADERS)
        self.assertEqual(result["sheets"][0]["column_count"], 18)
        self.assertEqual(result["source_row_count"], 1)
        self.assertEqual(len(result["preview_rows"]), 1)

    def test_xls_legacy_is_read_with_xlrd(self):
        result = inspect_workbook(xls_fixture(), "xls")
        self.assertEqual(result["selected_sheet"], "Dados")
        self.assertEqual(result["headers"][0], "REC")
        self.assertEqual(result["preview_rows"][0]["cells"][0]["display"], "TEST-001")

    def test_xlsx_content_with_legacy_xls_filename_uses_ooxml_reader(self):
        result = inspect_workbook(make_xlsx(), "xls")
        self.assertEqual(result["selected_sheet"], "Penalidades")
        self.assertEqual(result["headers"], HEADERS)

    def test_grouped_multiline_headers_cover_a_through_r_without_importing_header_row_as_data(self):
        data = make_grouped_xlsx()
        result = inspect_workbook(data, "xlsx")
        self.assertEqual(result["header_row"], 2)
        self.assertEqual(result["headers"], HEADERS)
        self.assertEqual(result["source_row_count"], 1)
        self.assertEqual(result["preview_rows"][0]["row_number"], 3)

    def test_single_non_candidate_can_be_inspected_and_multiple_tabs_require_choice(self):
        simple = make_xlsx(headers=["REC", "Designacao", "Campo extra"], values=["R1", "Link", "x"])
        one_sheet = inspect_workbook(simple, "xlsx")
        self.assertEqual(one_sheet["selected_sheet"], "Penalidades")
        self.assertFalse(one_sheet["sheets"][0]["candidate"])

        multiple = make_xlsx(headers=["REC", "Designacao"], values=["R1", "Link"], sheet_names=("Dados 1", "Dados 2"))
        result = inspect_workbook(multiple, "xlsx")
        self.assertIsNone(result["selected_sheet"])
        self.assertEqual(result["warning"], "Selecione uma aba para inspecionar a estrutura do arquivo.")
        selected = inspect_workbook(multiple, "xlsx", selected_sheet="Dados 2")
        self.assertEqual(selected["selected_sheet"], "Dados 2")

    def test_manual_mapping_still_shows_unmapped_rows_for_inspection(self):
        data = make_xlsx(
            header_row=1,
            headers=["Identificador operacional", "Referência do link", "Campo extra"],
            values=["REC-VALID-1", "ASPA/IP/00001", "valor"],
        )
        result = inspect_workbook(data, "xlsx")
        self.assertFalse(result["mapping"])
        self.assertEqual(result["source_row_count"], 1)
        self.assertEqual(result["preview_rows"][0]["cells"][0]["display"], "REC-VALID-1")

    def test_multiple_compatible_tabs_are_not_selected_by_name(self):
        data = make_xlsx(sheet_names=("Penalidades", "Outra aba"))
        result = inspect_workbook(data, "xlsx")
        self.assertIsNone(result["selected_sheet"])
        self.assertEqual(result["candidate_sheets"], ["Penalidades", "Outra aba"])
        self.assertEqual(inspect_workbook(data, "xlsx", selected_sheet="Outra aba")["header_row"], 2)

    def test_header_normalization_accepts_case_accents_and_spacing_and_keeps_unknowns(self):
        headers = [value.upper().replace(" ", "  ") for value in HEADERS]
        headers[1] = "DESIGNACAO"
        headers[-1] = "Campo que não existe"
        result = inspect_workbook(make_xlsx(headers=headers), "xlsx")
        self.assertTrue(result["header_detected"])
        self.assertEqual(result["mapping"]["0"], "rec")
        self.assertEqual(result["mapping"]["1"], "designation")
        self.assertEqual(result["headers"][-1], "Campo que não existe")

    def test_header_can_be_on_first_row_and_user_can_select_another_header_line(self):
        row_one = inspect_workbook(make_xlsx(header_row=1), "xlsx")
        self.assertEqual(row_one["header_row"], 1)
        later_header = make_xlsx(header_row=3)
        adjusted = inspect_workbook(later_header, "xlsx", selected_sheet="Penalidades", header_row=3)
        self.assertEqual(adjusted["header_row"], 3)
        self.assertTrue(adjusted["header_detected"])

    def test_canonical_values_keep_decimal_and_datetime_and_formula_is_not_executed(self):
        result = inspect_workbook(make_xlsx(formula=True), "xlsx")
        preview = result["preview_rows"][0]
        self.assertEqual(preview["canonical_data"]["monthly_value"], "1234.56")
        self.assertEqual(preview["canonical_data"]["opening"], "2026-08-03T10:15:00")
        formula_cell = preview["cells"][-1]
        self.assertEqual(formula_cell["type"], "formula")
        self.assertEqual(formula_cell["formula"], "=1+1")
        self.assertIsNone(formula_cell["cached_value"])
        self.assertEqual(_parse_decimal("R$ 1.234,56"), Decimal("1234.56"))

    def test_preview_formats_decimal_cells_without_changing_original_display(self):
        values = list(VALUES)
        values[13] = "17.150000000000002"
        values[17] = "24.133333333333333"
        preview = inspect_workbook(make_xlsx(values=values), "xlsx")["preview_rows"][0]
        self.assertEqual(preview["preview_cells"][13]["display"], "17,15")
        self.assertEqual(preview["preview_cells"][17]["display"], "24,13")
        self.assertEqual(preview["cells"][13]["display"], "17.150000000000002")
        numeric_value = list(VALUES)
        numeric_value[13] = 1.234
        canonical_decimal = inspect_workbook(make_xlsx(values=numeric_value), "xlsx")["preview_rows"][0]
        self.assertEqual(canonical_decimal["preview_cells"][13]["display"], "1,23")

    def test_formula_cache_drives_decimal_canonical_and_operational_import_cut(self):
        data = make_cached_formula_cut_xlsx()
        inspection = inspect_workbook(data, "xlsx")
        self.assertEqual(inspection["header_row"], 2)
        self.assertEqual(inspection["mapping"]["2"], "interconnection_fine")
        self.assertEqual(inspection["mapping"]["3"], "penalty_fine")
        self.assertEqual(inspection["mapping"]["4"], "unavailability_fine")
        self.assertEqual(inspection["source_row_count"], 4)
        first = inspection["preview_rows"][0]
        self.assertEqual(first["cells"][4]["formula"], "=+N3*0.05")
        self.assertEqual(first["cells"][4]["cached_value"], "16.55")
        self.assertEqual(first["canonical_data"]["unavailability_fine"], "16.55")
        self.assertEqual(first["preview_cells"][4]["display"], "16,55")
        with iter_import_rows(data, "xlsx", sheet_name=inspection["selected_sheet"],
                             header_row=inspection["header_row"], mapping=inspection["mapping"]) as (_headers, rows):
            imported = list(rows)
        self.assertEqual(len(imported), 4)
        positive_indisp = [row for row in imported if Decimal(row["canonical_data"]["unavailability_fine"] or "0") > 0]
        self.assertEqual(len(positive_indisp), 2)
        self.assertEqual([row["source_row_number"] for row in imported], [3, 4, 5, 6])
        self.assertTrue(all(row["canonical_data"]["rec"].isdigit() for row in imported))

    def test_decimal_parser_rejects_placeholders_and_excel_formula_text(self):
        for value in (0, 0.0, "0", "0.00", "0,00", None, "", "-", "null", "=+N3*0.05"):
            with self.subTest(value=value):
                parsed = _parse_decimal(value)
                self.assertTrue(parsed is None or parsed <= Decimal("0"))
        self.assertEqual(_parse_decimal("0,01"), Decimal("0.01"))
        self.assertIsNone(_parse_decimal("=+N3*0.05"))

    def test_original_formula_like_text_is_not_changed_during_import_preview(self):
        values = list(VALUES)
        values[2] = "+1+1"
        preview = inspect_workbook(make_xlsx(values=values), "xlsx")["preview_rows"][0]
        self.assertEqual(preview["cells"][2]["display"], "+1+1")
        self.assertEqual(preview["canonical_data"]["client"], "+1+1")

    def test_macro_enabled_xlsx_is_rejected(self):
        source = make_xlsx()
        output = BytesIO()
        with ZipFile(BytesIO(source), "r") as input_zip, ZipFile(output, "w", ZIP_DEFLATED) as output_zip:
            for member in input_zip.infolist():
                output_zip.writestr(member, input_zip.read(member.filename))
            output_zip.writestr("xl/vbaProject.bin", b"macro")
        with self.assertRaisesRegex(WorkbookImportError, "macros"):
            inspect_workbook(output.getvalue(), "xlsx")

    def test_upload_validation_rejects_bad_format_signature_empty_and_oversize(self):
        with self.assertRaises(WorkbookImportError):
            prepare_uploaded_workbook(SimpleUploadedFile("origem.csv", b"a,b"))
        with self.assertRaises(WorkbookImportError):
            prepare_uploaded_workbook(SimpleUploadedFile("vazio.xlsx", b""))
        with self.assertRaises(WorkbookImportError):
            prepare_uploaded_workbook(SimpleUploadedFile("corrompido.xlsx", b"PK-not-a-zip"))
        with self.assertRaises(WorkbookImportError):
            prepare_uploaded_workbook(SimpleUploadedFile("grande.xlsx", b"x" * (MAX_REPORT_FILE_SIZE + 1)))

    def test_xls_signature_must_match_content_and_path_is_sanitized(self):
        with self.assertRaises(WorkbookImportError):
            prepare_uploaded_workbook(SimpleUploadedFile("arquivo.xls", b"not a workbook"))
        self.assertEqual(safe_original_filename(r"..\..\cliente\penalidades.xlsx"), "penalidades.xlsx")


class ReportsImportWorkflowTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_media.cleanup)
        settings_override = override_settings(MEDIA_ROOT=self.temp_media.name)
        settings_override.enable()
        self.addCleanup(settings_override.disable)
        self.user = get_user_model().objects.create_user(username="reports-import", password="secure-test-password")
        permission = Permission.objects.get(content_type__app_label="relatorios", codename="view_reports")
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)
        self.data = make_xlsx()

    def create_with_upload(self, *, data=None, filename="penalidades.xlsx", month=8):
        response = self.client.post(reverse("relatorios:competence-create"), {
            "competence_month": month,
            "competence_year": 2026,
            "original_file": SimpleUploadedFile(filename, data or self.data),
        })
        if not PenaltyCompetence.objects.filter(competence_month=month, competence_year=2026).exists():
            errors = response.context["form"].errors.as_text() if response.context else "sem detalhes de formulário"
            self.fail(f"A competência não foi criada (HTTP {response.status_code}): {errors}")
        competence = PenaltyCompetence.objects.get(competence_month=month, competence_year=2026)
        return competence, response

    def mapping_post(self, *, sheet="Penalidades", header_row=2, intent="confirm"):
        result = inspect_workbook(self.data, "xlsx", selected_sheet=sheet, header_row=header_row)
        payload = {"sheet_name": sheet, "header_row": str(header_row), "intent": intent}
        for index in range(len(result["headers"])):
            payload[f"column_{index}"] = result["mapping"].get(str(index), "")
        return payload

    def test_upload_persists_private_original_metadata_and_review_preview(self):
        competence, response = self.create_with_upload(filename=r"..\unsafe\penalidades.xlsx")
        self.assertRedirects(response, reverse("relatorios:import-review", args=(competence.pk,)))
        competence.refresh_from_db()
        self.assertEqual(competence.original_file_size, len(self.data))
        self.assertEqual(competence.original_file_extension, "xlsx")
        self.assertEqual(len(competence.original_file_sha256), 64)
        self.assertEqual(competence.original_uploaded_by, self.user)
        self.assertTrue(competence.original_file.name.startswith("relatorios/originais/2026-08/"))
        self.assertNotIn("penalidades.xlsx", competence.original_file.name)
        self.assertTrue(competence.workbook_sheets[0]["candidate"])
        review = self.client.get(reverse("relatorios:import-review", args=(competence.pk,)))
        self.assertContains(review, "Cabeçalho na linha 2")
        self.assertContains(review, "REC-DEMO-01")
        self.assertContains(review, "Confirmar importação")
        self.assertTrue(AuditEvent.objects.filter(module="relatorios", action="ORIGINAL_UPLOADED_AND_VALIDATED").exists())

    def test_valid_legacy_xls_upload_is_stored_and_reviewed(self):
        competence, response = self.create_with_upload(data=xls_fixture(), filename="legado.xls", month=5)
        self.assertRedirects(response, reverse("relatorios:import-review", args=(competence.pk,)))
        self.assertEqual(competence.original_file_extension, "xls")
        review = self.client.get(reverse("relatorios:import-review", args=(competence.pk,)))
        self.assertContains(review, "TEST-001")
        self.assertContains(review, "REC")

    def test_confirm_import_creates_immutable_source_and_linked_empty_treatment(self):
        competence, _ = self.create_with_upload()
        response = self.client.post(reverse("relatorios:import-review", args=(competence.pk,)), self.mapping_post())
        form_errors = response.context["form"].errors if response.context else response.content[:1200]
        feedback = list(response.context.get("messages", [])) if response.context else []
        self.assertEqual(response.status_code, 302, f"{form_errors} {feedback}")
        self.assertRedirects(response, reverse("relatorios:workspace", args=(competence.pk,)))
        competence.refresh_from_db()
        source = PenaltySourceRow.objects.get(competence=competence)
        treatment = PenaltyTreatmentRow.objects.get(source_row=source)
        self.assertEqual(competence.status, PenaltyCompetence.Status.IN_TREATMENT)
        self.assertEqual(competence.imported_row_count, 1)
        self.assertIsNotNone(competence.imported_at)
        self.assertEqual(source.source_row_number, 3)
        self.assertEqual(source.source_data[0]["display"], "REC-DEMO-01")
        self.assertEqual(source.canonical_data["monthly_value"], "1234.56")
        self.assertEqual(treatment.source_values["designation"], "Link demonstrativo")
        self.assertEqual(treatment.arion_values.get("otrs_tickets", []), [])
        self.assertEqual(treatment.status, PenaltyTreatmentRow.Status.PENDING)
        self.assertEqual(treatment.source_row_id, source.pk)
        self.assertTrue(AuditEvent.objects.filter(module="relatorios", action="IMPORT_CONFIRMED").exists())
        grid = self.client.get(reverse("relatorios:workspace", args=(competence.pk,)), {"aba": "original"})
        self.assertContains(grid, "REC-DEMO-01")
        treatment_page = self.client.get(reverse("relatorios:workspace", args=(competence.pk,)), {"aba": "tratativa"})
        self.assertContains(treatment_page, "Link demonstrativo")
        self.assertContains(treatment_page, "Status da tratativa")

    def test_confirm_import_skips_footer_and_keeps_formula_cache_for_claro_cut(self):
        data = make_cached_formula_cut_xlsx()
        competence, _ = self.create_with_upload(data=data, filename="09-set26.xls", month=9)
        inspection = inspect_workbook(data, "xls")
        payload = {"sheet_name": inspection["selected_sheet"], "header_row": str(inspection["header_row"]), "intent": "confirm"}
        for index in range(len(inspection["headers"])):
            payload[f"column_{index}"] = inspection["mapping"].get(str(index), "")

        response = self.client.post(reverse("relatorios:import-review", args=(competence.pk,)), payload)
        self.assertEqual(response.status_code, 302)
        competence.refresh_from_db()
        sources = list(PenaltySourceRow.objects.filter(competence=competence).order_by("source_row_number"))
        self.assertEqual(competence.imported_row_count, 4)
        self.assertEqual([row.source_row_number for row in sources], [3, 4, 5, 6])
        first_indisp = sources[0].source_data[4]
        self.assertEqual(first_indisp["formula"], "=+N3*0.05")
        self.assertEqual(first_indisp["cached_value"], "16.55")
        self.assertEqual(sources[0].canonical_data["unavailability_fine"], "16.55")
        self.assertEqual(claro_treatment_candidates(PenaltyTreatmentRow.objects.filter(source_row__competence=competence)).count(), 2)

    def test_repeated_confirmation_does_not_duplicate_and_original_cannot_be_replaced(self):
        competence, _ = self.create_with_upload()
        import_url = reverse("relatorios:import-review", args=(competence.pk,))
        self.client.post(import_url, self.mapping_post())
        duplicate = self.client.post(import_url, self.mapping_post())
        self.assertRedirects(duplicate, reverse("relatorios:workspace", args=(competence.pk,)))
        self.assertEqual(PenaltySourceRow.objects.filter(competence=competence).count(), 1)
        old_hash = competence.original_file_sha256
        old_name = competence.original_file.name
        response = self.client.post(reverse("relatorios:competence-upload", args=(competence.pk,)), {
            "original_file": SimpleUploadedFile("substituto.xlsx", make_xlsx()),
        })
        self.assertRedirects(response, reverse("relatorios:workspace", args=(competence.pk,)))
        competence.refresh_from_db()
        self.assertEqual(competence.original_file_sha256, old_hash)
        self.assertEqual(competence.original_file.name, old_name)

    def test_original_grid_search_and_page_size_are_server_side(self):
        large_data = make_xlsx(row_count=30)
        competence, _ = self.create_with_upload(data=large_data, month=4)
        self.client.post(reverse("relatorios:import-review", args=(competence.pk,)), self.mapping_post())
        page_one = self.client.get(reverse("relatorios:workspace", args=(competence.pk,)), {
            "aba": "original", "page_size": "25",
        })
        self.assertContains(page_one, "Página 1 de 2")
        filtered = self.client.get(reverse("relatorios:workspace", args=(competence.pk,)), {
            "aba": "original", "q": "REC-DEMO-0010", "page_size": "50",
        })
        self.assertContains(filtered, "REC-DEMO-0010")
        self.assertNotContains(filtered, "REC-DEMO-0020")

    def test_original_download_is_authenticated_and_attachment_is_private(self):
        competence, _ = self.create_with_upload()
        response = self.client.get(reverse("relatorios:original-download", args=(competence.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(b"".join(response.streaming_content), self.data)

    def test_invalid_upload_shows_safe_error_and_creates_no_competence(self):
        response = self.client.post(reverse("relatorios:competence-create"), {
            "competence_month": 8,
            "competence_year": 2026,
            "original_file": SimpleUploadedFile("malicioso.xlsx", b"not a zip"),
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "não corresponde a uma planilha XLS ou XLSX válida")
        self.assertFalse(PenaltyCompetence.objects.filter(competence_month=8, competence_year=2026).exists())

    def test_transaction_rolls_back_source_and_treatment_rows_when_later_batch_fails(self):
        large_data = make_xlsx(row_count=501)
        competence, _ = self.create_with_upload(data=large_data, month=7)
        import aplicativos.relatorios.services.imports as import_service
        real_flush = import_service._flush_source_batch
        calls = 0

        def fail_after_second_batch(instance, batch):
            nonlocal calls
            calls += 1
            count = real_flush(instance, batch)
            if calls == 2:
                raise RuntimeError("simulated batch failure")
            return count

        with patch("aplicativos.relatorios.services.imports._flush_source_batch", side_effect=fail_after_second_batch):
            with self.assertRaises(RuntimeError):
                confirm_competence_import(
                    competence_id=competence.pk,
                    data=large_data,
                    extension="xlsx",
                    sheet_name="Penalidades",
                    header_row=2,
                    mapping={str(index): field for index, field in enumerate((
                        "rec", "designation", "client", "uf", "lp", "provider", "contract", "provider_unit_price",
                        "bandwidth", "opening", "fetec", "duration", "events", "monthly_value",
                        "interconnection_fine", "penalty_fine", "unavailability_fine", "hours_quantity",
                    ))},
                    physical_row_count=501,
                    user=self.user,
                )
        self.assertEqual(PenaltySourceRow.objects.filter(competence=competence).count(), 0)
        self.assertEqual(PenaltyTreatmentRow.objects.filter(source_row__competence=competence).count(), 0)
        competence.refresh_from_db()
        self.assertFalse(competence.imported_at)


class ReportsImportPermissionTests(TestCase):
    def test_all_import_routes_require_reports_permission(self):
        user = get_user_model().objects.create_user(username="reports-no-access", password="secure-test-password")
        competence = PenaltyCompetence.objects.create(competence_year=2026, competence_month=3)
        self.client.force_login(user)
        for url in (
            reverse("relatorios:competence-upload", args=(competence.pk,)),
            reverse("relatorios:import-review", args=(competence.pk,)),
            reverse("relatorios:original-download", args=(competence.pk,)),
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 403)
