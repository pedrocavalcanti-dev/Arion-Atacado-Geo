from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from aplicativos.integracoes.models import CMDBLinkMirror, CMDBOperatorMirror, OTRSTicket, OTRSTicketConfigItem
from aplicativos.auditoria.models import AuditEvent
from .models import PenaltyCompetence, PenaltyFinancialChange, PenaltySourceRow, PenaltyTreatmentOverride, PenaltyTreatmentRow
from .services.calculations import calculate_treatment, row_has_penalty, source_total_claro
from .services.treatment import (
    claro_treatment_candidates,
    competence_financial_summary,
    eligible_treatments,
    enrich_competence,
    treatment_queryset,
    treatment_export_dataset,
    validate_financial_decision,
)
from .services.exports import build_treatment_workbook


class ReconciliationTests(TestCase):
    def setUp(self):
        self.competence = PenaltyCompetence.objects.create(
            competence_year=2026, competence_month=8,
            status=PenaltyCompetence.Status.IN_TREATMENT,
            imported_at=timezone.now(),
        )

    def make_row(self, number, *, inter="0", penalty="0", indisp="0", **extra):
        source = {
            "rec": f"REC-{number}", "designation": f"Link {number}",
            "client": "Cliente", "uf": "SP", "provider_unit_price": "",
            "monthly_value": "100.00", "interconnection_fine": inter,
            "penalty_fine": penalty, "unavailability_fine": indisp,
        }
        source.update(extra)
        source_row = PenaltySourceRow.objects.create(
            competence=self.competence, source_row_number=number,
            source_data=[], canonical_data=source, row_hash=f"{number:064x}"[-64:],
        )
        treatment = PenaltyTreatmentRow.objects.create(source_row=source_row, total_claro=source_total_claro(source))
        return treatment

    def test_null_blank_and_zero_fines_are_not_eligible_but_one_cent_is(self):
        no_fine = self.make_row(2, inter=None, penalty="", indisp="0")
        fine = self.make_row(3, inter="0.01")
        self.assertFalse(row_has_penalty(no_fine.source_row.canonical_data))
        self.assertTrue(row_has_penalty(fine.source_row.canonical_data))
        self.assertEqual(list(eligible_treatments(PenaltyTreatmentRow.objects.all())), [fine])
        self.assertEqual(source_total_claro(no_fine.source_row.canonical_data), Decimal("0.00"))

    def test_default_claro_cut_uses_only_positive_indisp_and_other_filters_remain_explicit(self):
        indisp = self.make_row(2, indisp="0.01")
        inter_only = self.make_row(3, inter="25")
        penalty_only = self.make_row(4, penalty="8")
        no_fine = self.make_row(5, inter=None, penalty="-", indisp="0.00")
        all_rows = PenaltyTreatmentRow.objects.filter(source_row__competence=self.competence)

        self.assertEqual(list(claro_treatment_candidates(all_rows)), [indisp])
        self.assertEqual(list(treatment_queryset(self.competence)), [indisp])
        self.assertEqual(list(treatment_queryset(self.competence, {"multa": "indisp"})), [indisp])
        self.assertEqual(list(treatment_queryset(self.competence, {"multa": "inter"})), [inter_only])
        self.assertEqual(list(treatment_queryset(self.competence, {"multa": "penal"})), [penalty_only])
        self.assertEqual(list(treatment_queryset(self.competence, {"multa": "com"})), [indisp, inter_only, penalty_only])
        self.assertEqual(list(treatment_queryset(self.competence, {"multa": "todas"})), [indisp, inter_only, penalty_only, no_fine])

    def test_claro_cut_is_numeric_and_excludes_common_zero_and_placeholder_values(self):
        values = (0, 0.0, "0", "0.00", "0,00", None, "", "-", "=+N3*0.05")
        excluded = [self.make_row(index + 2, indisp=value) for index, value in enumerate(values)]
        one_cent = self.make_row(len(values) + 2, indisp="0.01")
        candidates = list(claro_treatment_candidates(PenaltyTreatmentRow.objects.filter(source_row__competence=self.competence)))
        self.assertEqual(candidates, [one_cent])
        self.assertEqual(len(excluded), len(values))

    def test_summary_keeps_canonical_decimal_point_as_decimal_separator(self):
        self.make_row(2, inter="1.234", indisp="0.01")
        summary = competence_financial_summary(self.competence)
        self.assertEqual(summary["inter_claro"], Decimal("1.234"))

    def test_decision_rules_require_consistent_effective_totals_and_reason(self):
        row = self.make_row(2, inter="100")
        row.total_arion = Decimal("100.00")
        row.decision = PenaltyTreatmentRow.Decision.DUE
        row.decision_reason = "Validado conforme evidência operacional."
        validate_financial_decision(row)
        row.decision = PenaltyTreatmentRow.Decision.REDUCED
        row.total_arion = Decimal("40.00")
        validate_financial_decision(row)
        row.decision = PenaltyTreatmentRow.Decision.INVALID
        row.total_arion = Decimal("0.00")
        validate_financial_decision(row)
        row.decision_reason = ""
        with self.assertRaisesMessage(Exception, "justificativa"):
            validate_financial_decision(row)

    def test_august_aggregate_reference_uses_decimal_values(self):
        first = self.make_row(2, indisp="2500.00")
        first.total_arion = Decimal("1000.00")
        first.amount_contested = Decimal("1500.00")
        first.save(update_fields=("total_arion", "amount_contested"))
        second = self.make_row(3, indisp="2549.63")
        second.total_arion = Decimal("1040.44")
        second.amount_contested = Decimal("1509.19")
        second.save(update_fields=("total_arion", "amount_contested"))
        summary = competence_financial_summary(self.competence)
        self.assertEqual(summary["total_claro"], Decimal("5049.63"))
        self.assertEqual(summary["with_penalty"], 2)
        self.assertEqual(summary["any_fine_count"], 2)
        self.assertEqual(summary["total_arion"], Decimal("2040.44"))
        self.assertEqual(summary["contested"], Decimal("3009.19"))
        self.assertEqual(summary["reduction_percent"], Decimal("59.59"))

    def test_enrichment_uses_exact_local_ticket_and_safe_cmdb_isp(self):
        row = self.make_row(2, inter="10", provider_unit_price="20260001234")
        ticket = OTRSTicket.objects.create(
            external_id=1234, ticket_number="20260001234", title="Atendimento teste",
            queue_external_id=1, queue_name="NOC", external_state_id=2,
            external_state_name="Fechado", business_status="CLOSED",
            source_created_at=timezone.now() - timedelta(hours=2), resolved_at=timezone.now(),
            responsible_user_id=77,
        )
        operator = CMDBOperatorMirror.objects.create(source_config_item_id="10013", name="FIERCE TELECOM")
        link = CMDBLinkMirror.objects.create(source_config_item_id="ci-123", isp_raw="10013", operator=operator)
        OTRSTicketConfigItem.objects.create(
            ticket=ticket, config_item_id="ci-123", cmdb_link=link,
            relation_type="DependsOn", relation_state="Valid",
        )
        result = enrich_competence(self.competence.pk)
        row.refresh_from_db()
        self.assertEqual(result["found"], 1)
        self.assertEqual(row.ticket, ticket)
        self.assertEqual(row.evidence_status, PenaltyTreatmentRow.EvidenceStatus.FOUND)
        self.assertEqual(row.isp, "FIERCE TELECOM")
        self.assertEqual(row.field_origins["isp"], PenaltyTreatmentRow.Origin.CMDB)
        self.assertEqual(row.arion_values["cmdb_isp_sources"][0]["id"], "10013")
        self.assertEqual(row.arion_values["cmdb_isp_sources"][0]["confidence"], "HIGH")
        self.assertEqual(row.arion_opening_origin, PenaltyTreatmentRow.Origin.OTRS)
        self.assertEqual(row.arion_values["otrs_responsible_user_id"], 77)
        self.assertEqual(row.root_cause, "")
        financial_change = PenaltyFinancialChange.objects.get(treatment=row)
        self.assertIsNone(financial_change.previous_total)
        self.assertEqual(financial_change.new_total, row.total_arion)
        self.assertEqual(financial_change.source, "ENRICHMENT")

    def test_cmdb_isp_identity_prefers_explicit_attribute_and_structured_operator_name(self):
        from .services.treatment import cmdb_isp_identity

        explicit_link = CMDBLinkMirror.objects.create(
            source_config_item_id="CI-ISP-ATTR", isp_raw="Infonet Provedor",
        )
        self.assertEqual(cmdb_isp_identity(explicit_link)["name"], "Infonet Provedor")
        self.assertEqual(cmdb_isp_identity(explicit_link)["confidence"], "HIGH")

        operator = CMDBOperatorMirror.objects.create(
            source_config_item_id="10013", source_number="36132000705",
            name="Infonet Provedor (Infotreini Provedores e Informática Ltda)",
        )
        explicit_and_related_link = CMDBLinkMirror.objects.create(
            source_config_item_id="CI-ISP-ATTR-REL", isp_raw="Infonet Provedor", operator=operator,
        )
        corroborated = cmdb_isp_identity(explicit_and_related_link)
        self.assertEqual(corroborated["name"], "Infonet Provedor")
        self.assertEqual(corroborated["full_name"], "Infonet Provedor (Infotreini Provedores e Informática Ltda)")
        self.assertEqual(corroborated["confidence"], "HIGH")

        associated_link = CMDBLinkMirror.objects.create(
            source_config_item_id="CI-ISP-REL", isp_raw="10013", operator=operator,
        )
        identity = cmdb_isp_identity(associated_link)
        self.assertEqual(identity["name"], "Infonet Provedor")
        self.assertEqual(identity["full_name"], "Infonet Provedor (Infotreini Provedores e Informática Ltda)")
        self.assertEqual(identity["id"], "10013")
        self.assertEqual(identity["number"], "36132000705")
        self.assertEqual(identity["confidence"], "HIGH")

    def test_explicit_isp_attribute_separates_a_trailing_technical_number(self):
        from .services.treatment import cmdb_isp_identity

        link = CMDBLinkMirror.objects.create(
            source_config_item_id="CI-ISP-ATTR-ID",
            isp_raw="Infonet Provedor (Infotreini Provedores e Informática Ltda) (36132000705)",
        )
        identity = cmdb_isp_identity(link)
        self.assertEqual(identity["name"], "Infonet Provedor")
        self.assertEqual(identity["full_name"], "Infonet Provedor (Infotreini Provedores e Informática Ltda)")
        self.assertEqual(identity["id"], "36132000705")
        self.assertEqual(identity["raw_id"], "36132000705")
        self.assertEqual(identity["confidence"], "HIGH")

    def test_numeric_isp_identifier_without_resolved_provider_is_never_a_name(self):
        from .services.treatment import cmdb_isp_identity

        link = CMDBLinkMirror.objects.create(source_config_item_id="CI-ISP-UNKNOWN", isp_raw="36132000705")
        identity = cmdb_isp_identity(link)
        self.assertEqual(identity["name"], "")
        self.assertEqual(identity["id"], "36132000705")
        self.assertEqual(identity["confidence"], "NONE")

    def test_enrichment_preserves_manual_isp_and_refreshes_legacy_numeric_isp(self):
        row = self.make_row(31, inter="10", indisp="10", provider_unit_price="20260003100")
        row.isp = "ISP MANUAL"
        row.field_origins = {"isp": PenaltyTreatmentRow.Origin.MANUAL}
        row.save(update_fields=("isp", "field_origins"))
        ticket = OTRSTicket.objects.create(
            external_id=3100, ticket_number="20260003100", title="ISP refresh test",
            queue_external_id=1, queue_name="NOC", external_state_id=2,
            external_state_name="Fechado", business_status="CLOSED",
        )
        operator = CMDBOperatorMirror.objects.create(
            source_config_item_id="10013", name="Infonet Provedor (Infotreini Provedores e Informática Ltda)",
        )
        link = CMDBLinkMirror.objects.create(source_config_item_id="CI-3100", isp_raw="10013", operator=operator)
        OTRSTicketConfigItem.objects.create(
            ticket=ticket, config_item_id="CI-3100", cmdb_link=link,
            relation_type="DependsOn", relation_state="Valid",
        )
        enrich_competence(self.competence.pk)
        row.refresh_from_db()
        self.assertEqual(row.isp, "ISP MANUAL")

        row.isp = "10013"
        row.field_origins = {"isp": PenaltyTreatmentRow.Origin.CMDB}
        row.save(update_fields=("isp", "field_origins"))
        enrich_competence(self.competence.pk)
        row.refresh_from_db()
        self.assertEqual(row.isp, "Infonet Provedor")
        self.assertEqual(row.field_origins["isp"], PenaltyTreatmentRow.Origin.CMDB)
        self.assertEqual(list(treatment_queryset(self.competence, {"isp": "infonet"})), [row])
        self.assertEqual(list(treatment_queryset(self.competence, {"isp": "10013"})), [])
        exported = list(treatment_export_dataset(self.competence))
        self.assertEqual(exported[0]["ISP"], "Infonet Provedor")
        self.assertEqual(exported[0]["ISP ID interno CMDB"], "10013")

    def test_enrichment_marks_attribute_operator_divergence_for_review(self):
        row = self.make_row(34, inter="10", provider_unit_price="20260003400")
        ticket = OTRSTicket.objects.create(
            external_id=3400, ticket_number="20260003400", title="ISP divergence test",
            queue_external_id=1, queue_name="NOC", external_state_id=2,
            external_state_name="Fechado", business_status="CLOSED",
        )
        operator = CMDBOperatorMirror.objects.create(source_config_item_id="10014", name="Provider B")
        link = CMDBLinkMirror.objects.create(
            source_config_item_id="CI-3400", isp_raw="10013", operator=operator,
        )
        OTRSTicketConfigItem.objects.create(
            ticket=ticket, config_item_id="CI-3400", cmdb_link=link,
            relation_type="DependsOn", relation_state="Valid",
        )
        enrich_competence(self.competence.pk)
        row.refresh_from_db()
        self.assertEqual(row.isp, "")
        self.assertNotIn("10013", row.arion_values["cmdb_isp_candidates"])
        self.assertIn("Provider B", row.arion_values["cmdb_isp_candidates"])
        self.assertEqual(row.arion_values["cmdb_isp_sources"][0]["confidence"], "AMBIGUOUS")

    def test_enrichment_keeps_multiple_providers_ambiguous_and_empty_provider_unidentified(self):
        row = self.make_row(32, inter="10", provider_unit_price="20260003200")
        ticket = OTRSTicket.objects.create(
            external_id=3200, ticket_number="20260003200", title="Multiple ISP test",
            queue_external_id=1, queue_name="NOC", external_state_id=2,
            external_state_name="Fechado", business_status="CLOSED",
        )
        for index, provider_name in enumerate(("ISP A", "ISP B"), start=1):
            link = CMDBLinkMirror.objects.create(
                source_config_item_id=f"CI-320{index}", isp_raw=provider_name,
            )
            OTRSTicketConfigItem.objects.create(
                ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link,
                relation_type="DependsOn", relation_state="Valid",
            )
        result = enrich_competence(self.competence.pk)
        row.refresh_from_db()
        self.assertGreaterEqual(result["ambiguous"], 1)
        self.assertEqual(row.isp, "")
        self.assertEqual(set(row.arion_values["cmdb_isp_candidates"]), {"ISP A", "ISP B"})

        empty_row = self.make_row(33, inter="10", provider_unit_price="20260003300")
        empty_ticket = OTRSTicket.objects.create(
            external_id=3300, ticket_number="20260003300", title="Missing ISP test",
            queue_external_id=1, queue_name="NOC", external_state_id=2,
            external_state_name="Fechado", business_status="CLOSED",
        )
        empty_link = CMDBLinkMirror.objects.create(source_config_item_id="CI-3300", isp_raw="")
        OTRSTicketConfigItem.objects.create(
            ticket=empty_ticket, config_item_id="CI-3300", cmdb_link=empty_link,
            relation_type="DependsOn", relation_state="Valid",
        )
        enrich_competence(self.competence.pk)
        empty_row.refresh_from_db()
        self.assertEqual(empty_row.isp, "")
        self.assertNotIn("cmdb_isp_candidates", empty_row.arion_values)

    def test_enrichment_does_not_overwrite_manual_values_and_marks_missing_ticket(self):
        row = self.make_row(2, inter="10", provider_unit_price="20260009999")
        manual_date = timezone.now() - timedelta(days=3)
        row.arion_opening_at = manual_date
        row.arion_opening_origin = PenaltyTreatmentRow.Origin.MANUAL
        row.isp = "ISP MANUAL"
        row.field_origins = {"isp": PenaltyTreatmentRow.Origin.MANUAL}
        row.root_cause = "CAUSA MANUAL"
        row.save()
        result = enrich_competence(self.competence.pk)
        row.refresh_from_db()
        self.assertEqual(result["not_found"], 1)
        self.assertEqual(row.evidence_status, PenaltyTreatmentRow.EvidenceStatus.NOT_FOUND)
        self.assertEqual(row.arion_opening_at, manual_date)
        self.assertEqual(row.isp, "ISP MANUAL")
        self.assertEqual(row.root_cause, "CAUSA MANUAL")

    def test_export_dataset_defaults_to_penalty_rows_and_preserves_financial_semantics(self):
        eligible = self.make_row(2, indisp="100")
        any_only = self.make_row(3, inter="25")
        self.make_row(4)
        dataset = list(treatment_export_dataset(self.competence))
        self.assertEqual(len(dataset), 1)
        self.assertEqual(dataset[0]["Valor cobrado Claro"], eligible.total_claro)
        self.assertIn("Valor calculado ARION", dataset[0])
        self.assertIn("Valor devido ARION", dataset[0])
        self.assertIn("Valor contestado", dataset[0])
        self.assertEqual(len(list(treatment_export_dataset(self.competence, scope="any"))), 2)
        self.assertIn(any_only.total_claro, [row["Valor cobrado Claro"] for row in treatment_export_dataset(self.competence, scope="any")])

    def test_xlsx_export_has_summary_and_neutralizes_formula_text_without_secrets(self):
        source = {
            "rec": "12345678", "designation": "=HYPERLINK(\"https://example.invalid\")",
            "client": "Cliente", "monthly_value": "331.00", "interconnection_fine": "0",
            "penalty_fine": "0", "unavailability_fine": "12.50", "validation_code": "SENHA-NAO-EXPORTAR",
        }
        source_row = PenaltySourceRow.objects.create(
            competence=self.competence, source_row_number=9, source_data=[], canonical_data=source, row_hash="9" * 64,
        )
        PenaltyTreatmentRow.objects.create(source_row=source_row, total_claro=Decimal("12.50"))
        workbook_file = build_treatment_workbook(self.competence)
        workbook = load_workbook(BytesIO(workbook_file.read()), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Tratativa ARION", "Resumo"])
        treatment_sheet = workbook["Tratativa ARION"]
        self.assertEqual(treatment_sheet.freeze_panes, "C2")
        self.assertEqual(treatment_sheet["B2"].value, "'=HYPERLINK(\"https://example.invalid\")")
        self.assertNotIn("SENHA-NAO-EXPORTAR", str(list(treatment_sheet.values)))
        self.assertEqual(workbook["Resumo"]["A2"].value, "Competência")

    def test_export_scopes_keep_the_shared_dataset_rules(self):
        due = self.make_row(20, inter="100", opening="2026-08-01T10:00:00")
        due.decision = PenaltyTreatmentRow.Decision.DUE
        due.save(update_fields=("decision",))
        invalid = self.make_row(21, penalty="5")
        invalid.decision = PenaltyTreatmentRow.Decision.INVALID
        invalid.save(update_fields=("decision",))
        divergent = self.make_row(22, indisp="3", opening="2026-08-01T10:00:00")
        divergent.arion_opening_at = timezone.now()
        divergent.save(update_fields=("arion_opening_at",))
        self.make_row(23)

        expected_rows = {"eligible": 1, "any": 3, "all": 4, "treated": 2, "divergences": 1, "invalid": 1}
        for scope, count in expected_rows.items():
            with self.subTest(scope=scope):
                workbook = load_workbook(BytesIO(build_treatment_workbook(self.competence, scope=scope).read()))
                self.assertEqual(workbook["Tratativa ARION"].max_row - 1, count)


class TreatmentExportAccessTests(TestCase):
    def setUp(self):
        self.competence = PenaltyCompetence.objects.create(competence_year=2026, competence_month=9)
        source = PenaltySourceRow.objects.create(
            competence=self.competence, source_row_number=2, source_data=[],
            canonical_data={"rec": "987654321", "designation": "Link", "interconnection_fine": "10", "penalty_fine": "0", "unavailability_fine": "0"},
            row_hash="2" * 64,
        )
        PenaltyTreatmentRow.objects.create(source_row=source, total_claro=Decimal("10.00"))
        self.url = reverse("relatorios:treatment-export", args=(self.competence.pk,))
        content_type = Permission.objects.get(content_type__app_label="relatorios", codename="view_reports").content_type
        self.view_permission = Permission.objects.get(content_type=content_type, codename="view_reports")
        self.export_permission = Permission.objects.get(content_type=content_type, codename="export_treatments")
        self.user = get_user_model().objects.create_user(username="reports-export", password="test-password")

    def test_download_requires_authentication_reports_access_and_export_permission(self):
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.user.user_permissions.add(self.export_permission)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.user.user_permissions.add(self.view_permission)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("Claro_Multas_setembro_2026_ARION.xlsx", response["Content-Disposition"])
        payload = b"".join(response.streaming_content)
        workbook = load_workbook(BytesIO(payload), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Tratativa ARION", "Resumo"])


class InlineTreatmentEditTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="reports-inline", password="test-password")
        content_type = Permission.objects.get(content_type__app_label="relatorios", codename="view_reports").content_type
        self.user.user_permissions.add(
            Permission.objects.get(content_type=content_type, codename="view_reports"),
            Permission.objects.get(content_type=content_type, codename="edit_treatments"),
        )
        self.client.force_login(self.user)
        self.competence = PenaltyCompetence.objects.create(
            competence_year=2026, competence_month=8, status=PenaltyCompetence.Status.IN_TREATMENT,
        )
        source = PenaltySourceRow.objects.create(
            competence=self.competence, source_row_number=5, source_data=[],
            canonical_data={"rec": "20260001234", "designation": "Link inline", "monthly_value": "100.00",
                            "interconnection_fine": "10.00", "penalty_fine": "0", "unavailability_fine": "0"},
            row_hash="5" * 64,
        )
        opening = timezone.make_aware(datetime(2026, 8, 1, 10, 0))
        self.treatment = PenaltyTreatmentRow.objects.create(source_row=source, arion_opening_at=opening)
        self.url = reverse("relatorios:treatment-inline-edit", args=(self.treatment.pk,))

    def _post(self, **values):
        return self.client.post(self.url, values, HTTP_X_REQUESTED_WITH="XMLHttpRequest", HTTP_ACCEPT="application/json")

    def test_authorized_inline_edit_recalculates_and_audits(self):
        response = self._post(field="arion_fetec_at", value="2026-08-01T12:00", reason="Horário validado", version=1)
        self.assertEqual(response.status_code, 200, response.content)
        self.treatment.refresh_from_db()
        self.assertEqual(self.treatment.tmp_calculated_minutes, 120)
        self.assertEqual(self.treatment.ev_calculated, Decimal("4.0000"))
        self.assertEqual(self.treatment.inter_effective, Decimal("0.83"))
        self.assertEqual(self.treatment.version, 2)
        self.assertTrue(AuditEvent.objects.filter(module="relatorios", action="TREATMENT_FIELDS_CHANGED").exists())
        change = PenaltyFinancialChange.objects.get(treatment=self.treatment)
        self.assertEqual(change.source, "INLINE_EDIT")

    def test_inline_edit_rejects_missing_permission_and_stale_version(self):
        self.user.user_permissions.remove(Permission.objects.get(codename="edit_treatments", content_type__app_label="relatorios"))
        denied = self._post(field="isp", value="ISP manual", version=1)
        self.assertEqual(denied.status_code, 403)
        self.user.user_permissions.add(Permission.objects.get(codename="edit_treatments", content_type__app_label="relatorios"))
        stale = self._post(field="isp", value="ISP manual", version=0)
        self.assertEqual(stale.status_code, 409)
        self.treatment.refresh_from_db()
        self.assertEqual(self.treatment.isp, "")


class CompetenceFinalizationTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="finance-finalizer", password="test-password")
        permission = Permission.objects.get(content_type__app_label="relatorios", codename="finalize_treatments")
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)
        self.competence = PenaltyCompetence.objects.create(
            competence_year=2026, competence_month=7,
            status=PenaltyCompetence.Status.IN_TREATMENT,
            imported_at=timezone.now(), calculation_rule_set="CLARO_LEGACY_2026_V1",
        )

    def create_treatment(self, number, inter, decision):
        source = PenaltySourceRow.objects.create(
            competence=self.competence, source_row_number=number, source_data=[], row_hash=str(number) * 64,
            canonical_data={"interconnection_fine": inter, "penalty_fine": "0",
                            "unavailability_fine": "0.01" if inter != "0" else "0",
                           "monthly_value": "100.00"},
        )
        treatment = PenaltyTreatmentRow.objects.create(
            source_row=source, decision=decision, decision_reason="Conciliação validada.",
            arion_opening_at=timezone.now() - timedelta(minutes=14_400),
            arion_fetec_at=timezone.now(), amount_not_charged=Decimal("0.00"),
        )
        if inter != "0":
            PenaltyTreatmentOverride.objects.create(
                treatment=treatment, field=PenaltyTreatmentOverride.Field.EV,
                override_value=Decimal("480"), calculated_value=None, justification="Tempo validado.", created_by=self.user,
            )
            PenaltyTreatmentOverride.objects.create(
                treatment=treatment, field=PenaltyTreatmentOverride.Field.UNAVAILABILITY,
                override_value=Decimal("0"), calculated_value=None, justification="Sem indisp devida.", created_by=self.user,
            )
        return treatment

    def test_no_fine_row_does_not_block_but_pending_fine_does(self):
        self.create_treatment(2, "0", PenaltyTreatmentRow.Decision.PENDING)
        eligible = self.create_treatment(3, "99.99", PenaltyTreatmentRow.Decision.PENDING)
        url = reverse("relatorios:competence-finalize", args=(self.competence.pk,))
        self.client.post(url)
        self.competence.refresh_from_db()
        self.assertNotEqual(self.competence.status, PenaltyCompetence.Status.FINALIZED)
        eligible.decision = PenaltyTreatmentRow.Decision.DUE
        eligible.save(update_fields=("decision",))
        self.client.post(url)
        self.competence.refresh_from_db()
        self.assertEqual(self.competence.status, PenaltyCompetence.Status.FINALIZED)
