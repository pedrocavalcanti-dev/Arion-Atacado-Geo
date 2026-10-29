import json
import csv
import logging
from datetime import datetime
from io import BytesIO, StringIO

from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_POST

from aplicativos.integracoes.services.analytics import LinkQueryService
from aplicativos.integracoes.services.closing_reason import ClosingReasonCanonicalizer
from aplicativos.configuracoes.permissions import access_required
from config.errors import api_error, require_media_type
from aplicativos.auditoria.services import AuditService
from .models import SavedAnalyticsView
from .services.analytics import AnalyticsQueryService, DashboardAnalyticsService, STATE_NAMES

logger = logging.getLogger(__name__)


@login_required
@access_required("view_dashboard")
def index(request):
    links = LinkQueryService.base(request.user)
    summary = links.aggregate(
        total_links=Count("pk"),
        state_count=Count("state", filter=Q(state__in=STATE_NAMES), distinct=True),
        provider_count=Count(
            "operator_id",
            filter=Q(operator__isnull=False, operator__name__gt=""),
            distinct=True,
        ),
    )
    access = {
        "can_view_chamados": request.user.has_perm("configuracoes.view_operational_tickets"),
        "can_view_base": request.user.has_perm("configuracoes.view_base"),
        "can_view_cockpit": request.user.has_perm("configuracoes.view_cockpit"),
        "can_view_map": request.user.has_perm("configuracoes.view_map"),
        "can_view_rankings": request.user.has_perm("configuracoes.view_rankings"),
        "can_view_ari": request.user.has_perm("configuracoes.view_ari"),
        "can_view_integrations": request.user.has_perm("configuracoes.view_integrations"),
        "can_manage_users": request.user.has_perm("configuracoes.manage_users"),
        "can_manage_groups": request.user.has_perm("configuracoes.manage_groups"),
        "can_manage_scopes": request.user.has_perm("configuracoes.manage_data_scopes"),
        "can_view_audit": request.user.has_perm("configuracoes.view_audit_log"),
        "can_view_documentation": request.user.has_perm("configuracoes.manage_groups"),
        "can_view_system": request.user.has_perm("configuracoes.view_system_info"),
    }
    access["can_view_operation"] = any(access[key] for key in (
        "can_view_chamados", "can_view_base", "can_view_cockpit", "can_view_map", "can_view_rankings", "can_view_ari",
    ))
    access["can_view_admin_area"] = any(access[key] for key in (
        "can_view_integrations", "can_manage_users", "can_manage_groups", "can_manage_scopes",
        "can_view_audit", "can_view_documentation", "can_view_system",
    ))
    return render(request, "painel/index.html", {
        **summary,
        **access,
    })


@login_required
@access_required("view_dashboard")
def dashboard(request):
    from datetime import date

    from django.utils import timezone
    from aplicativos.integracoes.services.analytics import PeriodService

    now = timezone.localtime()
    try:
        start, end = PeriodService.parse(request.GET)
        period_error = False
    except ValueError:
        start, end, period_error = None, None, True
    start_year, start_month = now.year, now.month - 11
    while start_month <= 0:
        start_month += 12
        start_year -= 1
    default_start = date(start_year, start_month, 1)
    if not start:
        start = default_start
    end = end or now.date()
    if start > end:
        start, end, period_error = default_start, now.date(), True

    context = DashboardAnalyticsService.build(user=request.user, start=start, end=end)
    return render(request, "painel/dashboard.html", {
        **context,
        "data_source_mode": "cmdb",
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "period_error": period_error,
    })

@login_required
@access_required("view_cockpit")
def cockpit(request):
    from datetime import date

    from django.utils import timezone

    now = timezone.localtime()
    start_year, start_month = now.year, now.month - 11
    while start_month <= 0:
        start_month += 12
        start_year -= 1
    return render(request, "painel/cockpit.html", {
        "can_save_views": request.user.has_perm("configuracoes.save_analytics_views"),
        "default_start": date(start_year, start_month, 1).isoformat(),
        "default_end": now.date().isoformat(),
    })


@login_required
@access_required("view_cockpit")
@require_POST
@csrf_protect
def analytics_query(request):
    media_error = require_media_type(request, "application/json")
    if media_error:
        return media_error
    try:
        payload = json.loads(request.body or b"{}")
        if not isinstance(payload, dict):
            return api_error(request, "BAD_REQUEST")
        return JsonResponse(AnalyticsQueryService().query(payload, user=request.user))
    except (ValueError, TypeError, json.JSONDecodeError):
        return api_error(request, "VALIDATION_ERROR", status=400)


def _pdf_report(generated, generated_by, result, filter_rows):
    from pathlib import Path
    from html import escape
    from reportlab import __file__ as reportlab_file
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    period_text = result["description"]["period"]
    fonts = Path(reportlab_file).parent / "fonts"
    if "ArionVera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ArionVera", str(fonts / "Vera.ttf")))
        pdfmetrics.registerFont(TTFont("ArionVera-Bold", str(fonts / "VeraBd.ttf")))
    ink, muted = colors.HexColor("#263746"), colors.HexColor("#687781")
    line, pale, accent = colors.HexColor("#dce3e7"), colors.HexColor("#f2f5f6"), colors.HexColor("#287c83")
    page_width, page_height = A4
    content_width = page_width - 32 * mm
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ArionTitle", parent=styles["Title"], fontName="ArionVera-Bold", fontSize=19, leading=23, textColor=ink, alignment=TA_LEFT, spaceAfter=5))
    styles.add(ParagraphStyle(name="ArionMeta", parent=styles["Normal"], fontName="ArionVera", fontSize=8.5, leading=12, textColor=muted))
    styles.add(ParagraphStyle(name="ArionSection", parent=styles["Heading2"], fontName="ArionVera-Bold", fontSize=11, leading=14, textColor=ink, spaceBefore=9, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle(name="ArionCell", parent=styles["Normal"], fontName="ArionVera", fontSize=7.5, leading=9.5, textColor=ink, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="ArionRight", parent=styles["ArionCell"], alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="ArionSmall", parent=styles["Normal"], fontName="ArionVera", fontSize=7.5, leading=10, textColor=muted))

    def para(value, style="ArionCell"):
        return Paragraph(escape(str(value if value is not None else "")), styles[style])

    def fmt(value):
        if isinstance(value, float):
            return f"{value:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
        if isinstance(value, int):
            return f"{value:,}".replace(",", ".")
        return str(value)

    def table(rows, widths, *, repeat=1):
        result_table = Table(rows, colWidths=widths, repeatRows=repeat, hAlign="LEFT", splitByRow=1)
        result_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), pale), ("TEXTCOLOR", (0, 0), (-1, 0), ink),
            ("LINEBELOW", (0, 0), (-1, 0), .7, line), ("LINEBELOW", (0, 1), (-1, -1), .35, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return result_table

    def chart(spec):
        horizontal = spec["type"] in {"horizontal_bar", "compact_distribution"}
        chart_height = max(42 * mm, min(76 * mm, len(spec.get("values", [])) * 7 * mm + 18 * mm)) if horizontal else 61 * mm
        drawing = Drawing(content_width, chart_height)
        chart_obj = HorizontalBarChart() if horizontal else VerticalBarChart()
        chart_obj.x, chart_obj.y = (76 if horizontal else 38), 12
        chart_obj.width, chart_obj.height = content_width - (88 if horizontal else 54), chart_height - 22
        chart_obj.data = [spec.get("values", [])]
        chart_obj.categoryAxis.categoryNames = [str(label) for label in spec.get("labels", [])]
        chart_obj.categoryAxis.labels.fontName = "ArionVera"
        chart_obj.categoryAxis.labels.fontSize = 6.5 if horizontal else (5 if len(chart_obj.categoryAxis.categoryNames) > 18 else 6.5)
        if not horizontal and len(chart_obj.categoryAxis.categoryNames) > 18:
            chart_obj.categoryAxis.labels.angle = 45
            chart_obj.categoryAxis.labels.boxAnchor = "ne"
            chart_obj.categoryAxis.labels.dy = -18
        chart_obj.valueAxis.labels.fontName = "ArionVera"
        chart_obj.valueAxis.labels.fontSize = 7
        chart_obj.valueAxis.valueMin = 0
        chart_obj.valueAxis.valueMax = max([1, *chart_obj.data[0]])
        chart_obj.valueAxis.valueStep = max(1, int(chart_obj.valueAxis.valueMax / 5))
        chart_obj.bars[0].fillColor = accent
        if horizontal:
            chart_obj.categoryAxis.reverseDirection = True
            chart_obj.barSpacing = 3
        drawing.add(chart_obj)
        return drawing

    def monthly_calendar(spec):
        month_rows = [[para(day) for day in spec.get("weekdays", [])]]
        for week in spec.get("weeks", []):
            cells = []
            for item in week:
                if not item.get("inside_month"):
                    cells.append("")
                elif item.get("inside_query_period"):
                    cells.append(Paragraph(f"<b>{item.get('day', '')}</b><br/>{fmt(item.get('value') or 0)}", styles["ArionCell"]))
                else:
                    cells.append(Paragraph(f"<font color='#8a969e'>{item.get('day', '')}<br/>fora do período</font>", styles["ArionSmall"]))
            month_rows.append(cells)
        calendar_table = Table(month_rows, colWidths=[content_width / 7] * 7, rowHeights=[8 * mm] + [13 * mm] * (len(month_rows) - 1), repeatRows=1)
        calendar_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), pale), ("TEXTCOLOR", (0, 0), (-1, 0), ink),
            ("GRID", (0, 0), (-1, -1), .35, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return calendar_table

    class NumberedCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.states = []
        def showPage(self):
            self.states.append(dict(self.__dict__))
            self._startPage()
        def save(self):
            total = len(self.states)
            for state in self.states:
                self.__dict__.update(state)
                self.saveState()
                self.setFont("ArionVera", 7)
                self.setFillColor(muted)
                self.drawString(16 * mm, 10 * mm, f"ARION ATACADO GEO | NOC ATACADO | {generated} | {generated_by}")
                self.drawRightString(page_width - 16 * mm, 10 * mm, f"P\u00e1gina {self._pageNumber} de {total}")
                self.restoreState()
                super().showPage()
            super().save()

    def page_frame(canvas_obj, doc):
        canvas_obj.saveState()
        canvas_obj.setFont("ArionVera-Bold", 8)
        canvas_obj.setFillColor(ink)
        canvas_obj.drawString(doc.leftMargin, page_height - 13 * mm, "ARION ATACADO GEO")
        canvas_obj.setFont("ArionVera", 8)
        canvas_obj.setFillColor(muted)
        canvas_obj.drawRightString(page_width - doc.rightMargin, page_height - 13 * mm, "NOC ATACADO / COCKPIT")
        canvas_obj.setStrokeColor(line)
        canvas_obj.line(doc.leftMargin, page_height - 16 * mm, page_width - doc.rightMargin, page_height - 16 * mm)
        canvas_obj.restoreState()

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=23 * mm, bottomMargin=18 * mm,
                            title=f"ARION ATACADO GEO - {result['description']['title']}", author=generated_by)
    story = [para("RELAT\u00d3RIO ANAL\u00cdTICO DO COCKPIT", "ArionMeta"),
             para(result["description"]["title"], "ArionTitle"),
             para(result["description"]["subtitle"], "ArionMeta"), Spacer(1, 4)]
    meta_rows = [[para("M\u00e9trica"), para(result["metric_label"])],
                 [para("Per\u00edodo"), para(period_text)],
                 [para("Base temporal"), para(result.get("period_date_label", "data de cria\u00e7\u00e3o"))],
                 [para("Visualizar por"), para(result["group_by_label"])],
                 [para("Gerado por"), para(generated_by)], [para("Gerado em"), para(generated)],
                 [para("Data Scope"), para("Aplicado antes dos filtros e agrega\u00e7\u00f5es")]]
    story += [table(meta_rows, [43 * mm, content_width - 43 * mm], repeat=0), para("Filtros ativos", "ArionSection")]
    if filter_rows:
        story.append(table([[para("Filtro"), para("Valor")]] + [[para(key), para(value)] for key, value in filter_rows],
                           [48 * mm, content_width - 48 * mm]))
    if result.get("scope") != "support":
        story.append(para(result.get("metric_semantics", "Consulta operacional do escopo selecionado."), "ArionSmall"))
    else:
        story.append(para("A m\u00e9trica exclui tickets cancelados e estados n\u00e3o classificados.", "ArionSmall")
                     if result["metric"] != "canceled_tickets" else para("Consulta espec\u00edfica de tickets cancelados.", "ArionSmall"))
    if result.get("comparison"):
        comparison = result["comparison"]
        story.append(para(f"Per\u00edodo anterior ({comparison['previous_start']} a {comparison['previous_end']}): {fmt(comparison['previous_value'])}; varia\u00e7\u00e3o {fmt(comparison['absolute_change'])} ({fmt(comparison['percentage_change'])}%).", "ArionSmall"))
    story.append(para("Resumo executivo", "ArionSection"))
    kpi_rows = [[para("Indicador"), para("Valor", "ArionRight"), para("Varia\u00e7\u00e3o", "ArionRight")]]
    for item in result["kpis"]:
        trend = item.get("trend")
        trend_text = ""
        if trend:
            trend_text = f"{fmt(trend['absolute_change'])} {trend['unit']}" if trend["unit"] else f"{fmt(trend['percentage_change'])}%"
        kpi_rows.append([para(item["label"]), para(fmt(item["value"]), "ArionRight"), para(trend_text, "ArionRight")])
    story.append(table(kpi_rows, [content_width - 65 * mm, 30 * mm, 35 * mm]))
    quality = result.get("provider_quality", {})
    if quality.get("unresolved"):
        story.append(para(f"Qualidade dos v\u00ednculos CMDB: {fmt(quality['unresolved'])} tickets sem provedor identificado; cobertura de {fmt(quality['coverage_percent'])}%.", "ArionSmall"))
    chart_specs = result.get("chart_specs", [])
    temporal_spec = next((item for item in chart_specs if item.get("type") == "temporal_overview"), None)
    calendar_spec = next((item for item in chart_specs if item.get("type") == "monthly_calendar"), None)
    story.append(PageBreak())
    if temporal_spec:
        story.append(para("Evolu\u00e7\u00e3o geral do per\u00edodo", "ArionSection"))
        if temporal_spec.get("subtitle"):
            story.append(para(temporal_spec["subtitle"], "ArionSmall"))
        story.append(chart(temporal_spec))
    if calendar_spec:
        story.append(para(calendar_spec.get("title", "Calend\u00e1rio mensal"), "ArionSection"))
        summary = calendar_spec.get("summary", {})
        story.append(para(f"{fmt(summary.get('total', 0))} {calendar_spec.get('unit', 'tickets')} · m\u00e9dia por dia {fmt(summary.get('average_per_day', 0))} · maior dia {summary.get('peak_day') or 'sem ocorr\u00eancias'} ({fmt(summary.get('peak_value', 0))}) · {fmt(summary.get('days_with_occurrence', 0))} dias com ocorr\u00eancia", "ArionSmall"))
        story.append(monthly_calendar(calendar_spec))
    dimension_specs = [item for item in chart_specs if item.get("type") in {"ranking", "horizontal_bar", "compact_distribution"}]
    if dimension_specs:
        story.append(PageBreak())
        story.append(para("Ranking e dimens\u00f5es complementares", "ArionSection"))
    for spec in dimension_specs:
        story.append(para(spec["title"], "ArionSection"))
        if spec.get("subtitle"):
            story.append(para(spec["subtitle"], "ArionSmall"))
        if spec["type"] == "ranking":
            dimension_label = "Provedor" if spec.get("dimension") == "provider" else "UF" if spec.get("dimension") == "uf" else result["group_by_label"]
            ranking = [[para("Posi\u00e7\u00e3o"), para(dimension_label), para(spec.get("unit", "Tickets"), "ArionRight"), para("Participa\u00e7\u00e3o", "ArionRight")]]
            ranking += [[para(i + 1), para(row["label"]), para(fmt(row["value"]), "ArionRight"), para(f"{fmt(row['share'])}%", "ArionRight")]
                        for i, row in enumerate(spec.get("details", []))]
            story.append(table(ranking, [20 * mm, content_width - 78 * mm, 28 * mm, 30 * mm]))
        else:
            story.append(chart(spec))
    story.append(PageBreak())
    story.append(para("ANEXO - TICKETS DO RESULTADO", "ArionSection"))
    ticket_table = [[para("Ticket"), para("Data base"), para("Status"), para("Motivo"), para("Raw OTRS"), para("Provedor"), para("UF"), para("Prioridade")]]
    ticket_table += [[para(row["ticket"]), para(row["base_date"]), para(row["status"]), para(row["reason"]), para(row.get("closing_code_raw", "")),
                      para(row["provider"]), para(row["uf"]), para(row["priority"])] for row in result.get("tickets", [])]
    widths = [20 * mm, 20 * mm, 19 * mm, 28 * mm, 30 * mm, content_width - 20*mm - 20*mm - 19*mm - 28*mm - 30*mm - 12*mm - 17*mm, 12*mm, 17*mm]
    story.append(table(ticket_table, widths))
    doc.build(story, onFirstPage=page_frame, onLaterPages=page_frame, canvasmaker=NumberedCanvas)
    return buffer.getvalue()

def _safe_spreadsheet_text(value):
    text = str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


@login_required
@access_required("view_cockpit")
@require_POST
@csrf_protect
def analytics_export(request):
    media_error = require_media_type(request, "application/json")
    if media_error:
        return media_error
    try:
        payload = json.loads(request.body or b"{}")
        if not isinstance(payload, dict):
            raise ValueError("A consulta deve ser um objeto JSON.")
        export_format = payload.pop("format", "csv")
        if export_format not in ("csv", "xlsx", "pdf"):
            raise ValueError("Formato de exporta\u00e7\u00e3o inv\u00e1lido.")
        payload["ticket_page"] = 1
        payload["ticket_page_size"] = 100
        payload.pop("ticket_search", None)
        result = AnalyticsQueryService().query(payload, user=request.user, include_all_rows=True)
        from django.utils import timezone
        from django.utils.text import slugify
        generated_at = timezone.localtime()
        generated = generated_at.strftime("%d/%m/%Y %H:%M")
        today = generated_at.strftime("%Y-%m-%d")
        generated_by = request.user.get_username()
        filter_labels = {
            "status": "Status", "source_state": "Situa\u00e7\u00e3o OTRS", "provider": "Provedor", "uf": "UF",
            "sla": "SLA", "priority": "Prioridade", "service": "Servi\u00e7o", "termination_reason": "Motivo",
            "stage": "Etapa", "origin": "Origem", "project": "Projeto / cliente",
        }
        raw_filters = result["normalized_query"]["filters"]
        options = result["options"]
        period_text = result["description"]["period"]
        scope_label = {"support": "Suporte", "cancellation": "Cancelamentos", "general": "Geral"}.get(result["scope"], result["scope"])
        filter_rows = [("Contexto", scope_label), ("Per\u00edodo", period_text), ("Base temporal", result["period_date_label"]),
                       ("M\u00e9trica", result["metric_label"])]
        for key, label in filter_labels.items():
            selected = raw_filters.get(key)
            if selected in (None, "", "all"):
                continue
            if key == "priority" and selected == "high-critical":
                display = "Alto / Cr\u00edtico"
            else:
                values = options.get(key, [])
                choice = next((item for item in values if isinstance(item, dict) and str(item.get("value")) == str(selected)), None)
                display = choice.get("label", selected) if choice else selected
            filter_rows.append((label, display))
        filter_rows += [("Visualizar por", result["group_by_label"]), ("Escopo de dados", "Aplicado pelo servidor")]
        if result.get("comparison"):
            comp = result["comparison"]
            filter_rows.append(("Per\u00edodo comparado", f"{comp['previous_start']} a {comp['previous_end']}"))
        if result.get("scope") == "support":
            if result["metric"] != "canceled_tickets" and raw_filters.get("status") not in {"CANCELED", "UNKNOWN"}:
                filter_rows.append(("Regra operacional", "A m\u00e9trica exclui tickets cancelados e estados n\u00e3o classificados"))
        else:
            filter_rows.append(("Regra operacional", result.get("metric_semantics", "Escopo operacional do Cockpit")))
        group_slug = slugify(result["group_by"]) or "resultado"
        ticket_headers = ["ticket", "data_base", "created_at", "status", "otrs_state", "closing_reason", "closing_reason_raw", "provider", "uf",
                          "priority", "sla", "source", "config_item", "number", "designation", "circuit"]
        def ticket_values(row):
            links = row.get("links", [])
            fields = ("config_item_id", "number", "designation", "circuit")
            flattened = [" | ".join(str(link.get(field) or "") for link in links) for field in fields]
            base = [row.get("ticket"), row.get("base_date"), row.get("created_at"), row.get("status"), row.get("source_state"),
                    row.get("reason"), row.get("closing_code_raw"), row.get("provider"), row.get("uf"), row.get("priority"), row.get("sla"), row.get("source")]
            return base + flattened
        if export_format == "csv":
            buffer = StringIO(newline="")
            writer = csv.writer(buffer, delimiter=";")
            writer.writerow(ticket_headers)
            for row in result["tickets"]:
                writer.writerow([_safe_spreadsheet_text(value or "") for value in ticket_values(row)])
            response = HttpResponse("\ufeff" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
            filename = f"arion-cockpit-tickets-{today}.csv"
        elif export_format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font
            workbook = Workbook()
            summary = workbook.active
            summary.title = "Resumo"
            summary.append(["ARION ATACADO GEO - NOC ATACADO - Cockpit"])
            summary["A1"].font = Font(bold=True, size=14)
            summary.append(["An\u00e1lise", _safe_spreadsheet_text(result["description"]["title"])])
            summary.append(["Gerado por", _safe_spreadsheet_text(generated_by)])
            summary.append(["Gerado em", generated])
            summary.append(["M\u00e9trica", result["metric_label"]])
            summary.append(["Per\u00edodo", _safe_spreadsheet_text(period_text)])
            summary.append(["Base temporal", result["period_date_label"]])
            summary.append(["Resultado", result["value"]])
            summary.append(["Tickets do resultado", result["ticket_result_total"]])
            if result.get("comparison"):
                comp = result["comparison"]
                summary.append(["Per\u00edodo anterior", _safe_spreadsheet_text(f"{comp['previous_start']} a {comp['previous_end']}")])
                summary.append(["Valor anterior", comp["previous_value"]])
                summary.append(["Varia\u00e7\u00e3o percentual", comp["percentage_change"]])
            summary.append([])
            summary.append(["Filtros ativos", "Valor"])
            summary.append(["Escopo de dados", "Aplicado pelo servidor"])
            for label, value in filter_rows:
                summary.append([_safe_spreadsheet_text(label), _safe_spreadsheet_text(value)])
            summary.append([])
            summary.append(["Indicador", "Valor", "Varia\u00e7\u00e3o"])
            for item in result["kpis"]:
                trend = item.get("trend")
                text = ""
                if trend:
                    text = f"{trend['absolute_change']} {trend['unit']}" if trend["unit"] else f"{trend['percentage_change']}%"
                summary.append([item["label"], item["value"], text])
            quality = result.get("provider_quality", {})
            summary.append(["Com provedor identificado", quality.get("identified", 0)])
            summary.append(["Sem provedor identificado", quality.get("unresolved", 0)])
            summary.append(["Cobertura de provedor (%)", quality.get("coverage_percent", 0)])
            distribution = workbook.create_sheet("Distribui\u00e7\u00e3o")
            distribution.append([result["group_by_label"], result["metric_label"], "Participa\u00e7\u00e3o (%)"])
            for row in result["detail"]:
                distribution.append([_safe_spreadsheet_text(row["dimension"]), row["value"], row["share"]])
            tickets_sheet = workbook.create_sheet("Tickets")
            tickets_sheet.append(ticket_headers)
            for row in result["tickets"]:
                tickets_sheet.append([_safe_spreadsheet_text(value or "") for value in ticket_values(row)])
            for sheet in (distribution, tickets_sheet):
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
            for sheet in workbook:
                for column in sheet.columns:
                    letter = column[0].column_letter
                    sheet.column_dimensions[letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            output = BytesIO()
            workbook.save(output)
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            filename = f"arion-cockpit-completo-{today}.xlsx"
        else:
            response = HttpResponse(_pdf_report(generated, generated_by, result, filter_rows), content_type="application/pdf")
            filename = f"arion-cockpit-analitico-{today}.pdf"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        AuditService.record(request=request, category="EXPORT", action=f"EXPORT_{export_format.upper()}", module="Cockpit",
                            description=f"Cockpit export {export_format.upper()}", result="SUCCESS",
                            metadata={"format": export_format, "filters": raw_filters, "start": result["start_date"],
                                      "end": result["end_date"], "group_by": result["group_by_label"],
                                      "ticket_rows": len(result["tickets"])})
        return response
    except (ValueError, TypeError, json.JSONDecodeError):
        return api_error(request, "VALIDATION_ERROR", status=400)
    except Exception:
        logger.exception("Cockpit export failed (format=%s)", locals().get("export_format", "unknown"))
        return api_error(request, "EXPORT_FAILED", status=500)

@login_required
@access_required("view_cockpit")
@require_POST
@csrf_protect
def saved_views(request):
    media_error = require_media_type(request, "application/json")
    if media_error:
        return media_error
    try:
        payload = json.loads(request.body or b"{}")
        if not isinstance(payload, dict):
            return api_error(request, "BAD_REQUEST")
        action = payload.get("action", "list")
        if action in ("save", "rename", "delete") and not request.user.has_perm("configuracoes.save_analytics_views"):
            return api_error(request, "ACCESS_DENIED")
        if action == "list":
            rows = SavedAnalyticsView.objects.filter(user=request.user)
            views = []
            for row in rows:
                configuration = dict(row.configuration or {})
                raw_filters = configuration.get("filters")
                if isinstance(raw_filters, dict):
                    filters = dict(raw_filters)
                    for legacy_key in ("closing_reason", "closing_code_raw"):
                        if legacy_key in filters and "termination_reason" not in filters:
                            filters["termination_reason"] = filters[legacy_key]
                        filters.pop(legacy_key, None)
                    reason = filters.get("termination_reason")
                    if reason not in (None, "", "all"):
                        filters["termination_reason"] = ClosingReasonCanonicalizer.key(reason)
                    configuration["filters"] = filters
                views.append({"id": row.pk, "name": row.name, "configuration": configuration, "favorite": row.favorite})
            return JsonResponse({"views": views})
        if action == "save":
            name, config = str(payload.get("name", "")).strip(), payload.get("configuration")
            if not name or len(name) > 80 or not isinstance(config, dict):
                raise ValueError("Informe um nome e uma configuração válida.")
            allowed = {"metric", "group_by", "period", "period_from", "period_to", "start", "end", "compare_to", "filters", "detail_page"}
            if set(config) - allowed:
                raise ValueError("Configuração inválida.")
            row, created = SavedAnalyticsView.objects.update_or_create(user=request.user, name=name, defaults={"configuration": config})
            if created:
                AuditService.record(request=request, user=request.user, category="COCKPIT", action="SAVED_VIEW_CREATED", description="Visão salva criada", obj=row, metadata={"name": row.name})
            return JsonResponse({"id": row.pk, "name": row.name})
        if action == "delete":
            row = SavedAnalyticsView.objects.filter(user=request.user, pk=payload.get("id")).first()
            if row:
                AuditService.record(request=request, user=request.user, category="COCKPIT", action="SAVED_VIEW_DELETED", description="Visão salva removida", obj=row, metadata={"name": row.name})
                row.delete()
            return JsonResponse({"ok": True})
        if action == "rename":
            name = str(payload.get("name", "")).strip()
            if not name or len(name) > 80:
                raise ValueError("Informe um nome de até 80 caracteres.")
            row = SavedAnalyticsView.objects.filter(user=request.user, pk=payload.get("id")).first()
            if row is None:
                return api_error(request, "NOT_FOUND")
            if SavedAnalyticsView.objects.filter(user=request.user, name=name).exclude(pk=row.pk).exists():
                raise ValueError("Você já tem uma consulta com esse nome.")
            old_name = row.name
            row.name = name
            row.save(update_fields=("name", "updated_at"))
            AuditService.record(request=request, user=request.user, category="COCKPIT", action="SAVED_VIEW_UPDATED", description="Consulta salva renomeada", obj=row, metadata={"old_name": old_name, "name": row.name})
            return JsonResponse({"id": row.pk, "name": row.name})
        raise ValueError("Ação inválida.")
    except (ValueError, TypeError, json.JSONDecodeError):
        return api_error(request, "VALIDATION_ERROR", status=400)
