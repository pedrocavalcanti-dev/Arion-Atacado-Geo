from datetime import date, datetime, timezone as datetime_timezone
from io import BytesIO
import json
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.conf import settings
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from unittest.mock import patch
from aplicativos.configuracoes.models import GroupDataScope
from aplicativos.base_links.models import Operadora, Projeto
from aplicativos.integracoes.models import CMDBLinkMirror, CMDBOperatorMirror, OTRSTicket, OTRSTicketConfigItem
from aplicativos.integracoes.services.analytics import TicketQueryService
from aplicativos.integracoes.services.semantics import LinkSemantics
from .services.analytics import AnalyticsQueryService, DashboardAnalyticsService


class MirrorAnalyticsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="analytics-user", password="safe-password-1")
        group = Group.objects.get(name="Analista NOC")
        self.user.groups.set([group])
        scope, _ = GroupDataScope.objects.get_or_create(group=group)
        scope.mode = GroupDataScope.Mode.ALL
        scope.save()
        self.operator = CMDBOperatorMirror.objects.create(source_config_item_id="analytics-op", name="Operadora Analytics")
        self.link = CMDBLinkMirror.objects.create(source_config_item_id="analytics-ci-1", source_number="AN-1", state="SP", city="Sao Paulo", operator=self.operator, legacy_project="Projeto A", isp_raw="ISP real")
        self.other_link = CMDBLinkMirror.objects.create(source_config_item_id="analytics-ci-2", source_number="AN-2", state="RJ", operator=self.operator, legacy_project="Projeto A", isp_raw="ISP real")
        self.client.force_login(self.user)

    def test_home_summary_counts_only_scoped_links_named_providers_and_valid_states(self):
        scope = self.user.groups.get(name="Analista NOC").data_scope
        scope.mode = GroupDataScope.Mode.CUSTOM
        scope.save()
        allowed_operator = Operadora.objects.create(name=self.operator.name, normalized_name="operadora-analytics")
        allowed_project = Projeto.objects.create(name="Projeto A", normalized_name="projeto-a")
        scope.allowed_operators.set([allowed_operator])
        scope.allowed_projects.set([allowed_project])

        CMDBLinkMirror.objects.create(
            source_config_item_id="analytics-invalid-state", state="N/A",
            operator=self.operator, legacy_project="Projeto A",
        )
        CMDBLinkMirror.objects.create(
            source_config_item_id="analytics-invalid-state-name", state="Sao Paulo",
            operator=self.operator, legacy_project="Projeto A",
        )
        hidden_provider = CMDBOperatorMirror.objects.create(
            source_config_item_id="analytics-hidden-provider", name="Fora do escopo",
        )
        CMDBLinkMirror.objects.create(
            source_config_item_id="analytics-hidden-link", state="AM",
            operator=hidden_provider, legacy_project="Projeto B",
        )

        response = self.client.get(reverse("painel:index"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_links"], 4)
        self.assertEqual(response.context["state_count"], 2)
        self.assertEqual(response.context["provider_count"], 1)
        self.assertContains(response, "Estados")
        self.assertNotContains(response, "Nenhuma importa&ccedil;&atilde;o registrada")

    def test_home_shows_only_permitted_accesses_and_preserves_route_guards(self):
        response = self.client.get(reverse("painel:index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ACESSO OPERACIONAL")
        self.assertContains(response, 'class="overview-module-card" href="/chamados/"')
        self.assertContains(response, 'class="overview-module-card" href="/ari/"')
        self.assertContains(response, 'class="overview-module-card"', count=6)
        self.assertContains(response, "Conta")
        self.assertNotContains(response, "ACESSO RESTRITO")
        self.assertNotContains(response, "Documenta&ccedil;&atilde;o")
        self.assertEqual(self.client.get(reverse("configuracoes:documentation")).status_code, 403)

        admin_user = get_user_model().objects.create_user(
            username="home-admin", email="home-admin@example.test", password="safe-password-2",
        )
        admin_user.groups.set([Group.objects.get(name="Administrador do Sistema")])
        self.assertFalse(admin_user.is_superuser)
        self.client.force_login(admin_user)
        response = self.client.get(reverse("painel:index"))
        self.assertEqual(response.status_code, 200)
        for label in (
            "Chamados", "Base de Links", "Cockpit", "Mapa", "Rankings", "ARI I.A", "Conta",
            "Administra&ccedil;&atilde;o", "Integra&ccedil;&otilde;es", "Usu&aacute;rios", "Grupos",
            "Escopos", "Auditoria", "Documenta&ccedil;&atilde;o", "Sistema", "Minha conta",
        ):
            self.assertContains(response, label)

        limited = get_user_model().objects.create_user(
            username="home-limited", password="safe-password-3",
        )
        dashboard_permission = Permission.objects.get(
            content_type__app_label="configuracoes", codename="view_dashboard",
        )
        limited.user_permissions.add(dashboard_permission)
        self.client.force_login(limited)
        response = self.client.get(reverse("painel:index"))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'href="/chamados/"')
        self.assertNotContains(response, 'href="/ari/"')
        self.assertNotContains(response, reverse("configuracoes:documentation"))
        self.assertNotContains(response, "ACESSO RESTRITO")
        self.assertContains(response, "Minha conta")
        self.assertEqual(self.client.get(reverse("ari:index")).status_code, 403)
        self.assertEqual(self.client.get(reverse("configuracoes:documentation")).status_code, 403)

    def _ticket(self, ext, status, open_state, day):
        created_at = datetime(2026, 10, day, 3, tzinfo=datetime_timezone.utc)
        terminal_at = created_at if status in {"RESOLVED", "CLOSED", "CANCELED"} else None
        return OTRSTicket.objects.create(external_id=ext, ticket_number=f"A-{ext}", queue_external_id=107, queue_name="NOC - ATACADO", external_state_id=ext, external_state_name=status, business_status=status, is_business_open=open_state, priority_name="3 normal", service_name="Internet", sla_name="Resposta NOC", source_created_at=created_at, terminal_state_at=terminal_at, canceled_at=terminal_at if status == "CANCELED" else None)

    def test_dashboard_renders_real_snapshot_and_period(self):
        ticket = self._ticket(101, "OPEN", True, 5)
        OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id="analytics-ci-1", cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        response = self.client.get(reverse("dashboard"), {"date_from": "2026-10-01", "date_to": "2026-10-31"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_links"], 2)
        self.assertEqual(response.context["business_open"], 1)
        self.assertEqual(response.context["links_affected"], 1)
        self.assertContains(response, "Situação atual dos chamados")
        self.assertContains(response, "Em atendimento")
        self.assertContains(response, "Agendados")
        queue_url = reverse("chamados:index")
        for status in ("OPEN", "PENDING", "SCHEDULED"):
            self.assertContains(response, f'href="{queue_url}?tab={status}"')
        self.assertNotContains(response, "otrs_status=")
        self.assertEqual(response.context["open_count"], 1)
        self.assertEqual(response.context["pending_count"], 0)
        self.assertEqual(response.context["scheduled_count"], 0)
        self.assertEqual(response.context["resolved_count"], 0)
        self.assertEqual(response.context["closed_count"], 0)
        self.assertEqual(response.context["canceled_count"], 0)
        self.assertNotContains(response, ">OPEN<")
        self.assertEqual(response.context["top_providers"], [{"name": "Operadora Analytics", "value": 2}])
        self.assertNotContains(response, "CLASSIFICAÇÃO LEGADA")
        self.assertNotContains(response, "VIVO")
        self.assertNotContains(response, "CLARO")
        self.assertNotContains(response, "ALGAR")
        self.assertContains(response, "Top 10 Estados com mais Links")
        self.assertContains(response, "Provedores com mais Links")

    def test_dashboard_period_and_snapshot_metrics_are_independent_and_noc_only(self):
        october_open = self._ticket(201, "OPEN", True, 5)
        october_pending = self._ticket(202, "PENDING", True, 8)
        october_scheduled = self._ticket(208, "SCHEDULED", True, 9)
        september_resolved = self._ticket(203, "RESOLVED", False, 5)
        september_resolved.source_created_at = datetime(2026, 9, 5, tzinfo=datetime_timezone.utc)
        september_resolved.save()
        self._ticket(204, "CLOSED", False, 10)
        self._ticket(205, "CANCELED", False, 11)
        self._ticket(206, "UNKNOWN", False, 12)
        other_queue = self._ticket(207, "OPEN", True, 13)
        other_queue.queue_name = "OUTRA FILA"
        other_queue.save()
        invalid_state_link = CMDBLinkMirror.objects.create(
            source_config_item_id="analytics-invalid-state", state="N/A",
        )

        def relate(ticket, link, suffix):
            OTRSTicketConfigItem.objects.create(
                ticket=ticket, config_item_id=link.source_config_item_id,
                cmdb_link=link, relation_type=f"DependsOn-{suffix}", relation_state="Valid",
            )

        relate(october_open, self.link, "201")
        relate(october_pending, self.link, "202")
        relate(october_scheduled, self.link, "208")
        relate(october_pending, self.other_link, "202")
        relate(september_resolved, self.other_link, "203")
        relate(other_queue, self.link, "207")

        with patch("aplicativos.integracoes.clients.cmdb.CMDBClient._post", side_effect=AssertionError("Dashboard não deve chamar ServiceUp")), patch(
            "aplicativos.integracoes.clients.otrs.OTRSReadOnlyClient._connection", side_effect=AssertionError("Dashboard não deve conectar ao MySQL OTRS")
        ), CaptureQueriesContext(connection) as queries:
            october = DashboardAnalyticsService.build(
                user=self.user, start=date(2026, 10, 1), end=date(2026, 10, 31),
            )
        self.assertLessEqual(len(queries), 20)
        self.assertEqual(october["calls"], 6)  # Tickets NOC sem relação também entram no volume.
        self.assertEqual(october["repeat"], 1)  # Somente o Link com dois Tickets distintos.
        self.assertEqual(october["monthly"][0]["calls"], 6)
        self.assertEqual(october["monthly"][0]["repeat"], 1)
        self.assertEqual(october["business_open"], 3)
        self.assertEqual(october["open_count"], 1)
        self.assertEqual(october["pending_count"], 1)
        self.assertEqual(october["scheduled_count"], 1)
        self.assertEqual(october["resolved_count"], 1)
        self.assertEqual(october["closed_count"], 1)
        self.assertEqual(october["canceled_count"], 1)
        self.assertEqual(october["unknown_count"], 1)
        self.assertEqual(october["total_links"], 3)
        self.assertEqual(october["state_count"], 2)
        self.assertEqual([item["code"] for item in october["top_states"]], ["RJ", "SP"])

        september = DashboardAnalyticsService.build(
            user=self.user, start=date(2026, 9, 1), end=date(2026, 9, 30),
        )
        self.assertEqual(september["calls"], 1)
        self.assertEqual(september["repeat"], 0)
        self.assertEqual(september["monthly"][0]["calls"], 1)
        self.assertEqual(september["business_open"], october["business_open"])
        self.assertEqual(september["scheduled_count"], october["scheduled_count"])
        self.assertEqual(september["open_count"], october["open_count"])
        self.assertEqual(september["pending_count"], october["pending_count"])
        self.assertEqual(september["resolved_count"], october["resolved_count"])
        self.assertEqual(september["closed_count"], october["closed_count"])
        self.assertEqual(september["canceled_count"], october["canceled_count"])
        self.assertEqual(september["total_links"], october["total_links"])
        self.assertEqual(september["state_count"], october["state_count"])
        rendered = self.client.get(reverse("dashboard"), {
            "date_from": "2026-09-01", "date_to": "2026-09-30",
        })
        self.assertContains(rendered, "Não classificados")
        self.assertNotContains(rendered, "UNKNOWN")

    def test_unknown_ticket_status_is_hidden_when_zero(self):
        response = self.client.get(reverse("dashboard"), {
            "date_from": "2026-10-01", "date_to": "2026-10-31",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["unknown_count"], 0)
        self.assertContains(response, "Chamados abertos agora")
        self.assertContains(response, "Em atendimento + pendentes")
        self.assertNotContains(response, "Não classificados")
        self.assertNotContains(response, "Prioridade, serviço e SLA")
        self.assertNotContains(response, "Tickets por data de criação")
        self.assertNotContains(response, "Distribuição por provedor")

    def test_dashboard_provider_ranking_uses_resolved_names_and_current_links_only(self):
        self.link.legacy_operator_classification = CMDBLinkMirror.LegacyOperator.VIVO
        self.link.save(update_fields=("legacy_operator_classification",))
        self.other_link.legacy_operator_classification = CMDBLinkMirror.LegacyOperator.CLARO
        self.other_link.save(update_fields=("legacy_operator_classification",))
        self.operator.name = "FIERCE TELECOM"
        self.operator.save()
        absent = CMDBLinkMirror.objects.create(
            source_config_item_id="analytics-absent", source_number="AN-ABSENT",
            state="SP", operator=self.operator, is_present_in_source=False,
        )
        unresolved = CMDBLinkMirror.objects.create(
            source_config_item_id="analytics-unresolved", source_number="AN-UNRESOLVED",
            state="SP", isp_raw="10097",
        )

        result = DashboardAnalyticsService.build(
            user=self.user, start=date(2026, 10, 1), end=date(2026, 10, 31),
        )
        self.assertEqual(result["total_links"], 3)
        self.assertEqual(result["top_providers"], [{"name": "FIERCE TELECOM", "value": 2}])
        self.assertNotIn(str(absent.source_config_item_id), [row["name"] for row in result["top_providers"]])
        self.assertNotIn("10097", [row["name"] for row in result["top_providers"]])
        self.assertNotIn("macro_operators", result)

    def test_dashboard_rankings_are_limited_to_top_ten(self):
        for index in range(11):
            operator = CMDBOperatorMirror.objects.create(
                source_config_item_id=f"provider-{index}", name=f"Provider {index:02d}",
            )
            CMDBLinkMirror.objects.create(
                source_config_item_id=f"link-{index}", source_number=f"LINK-{index}",
                state="AC", operator=operator,
            )
        result = DashboardAnalyticsService.build(
            user=self.user, start=date(2026, 10, 1), end=date(2026, 10, 31),
        )
        self.assertEqual(len(result["top_states"]), 3)
        self.assertEqual(len(result["top_providers"]), 10)
        self.assertEqual(result["top_providers"][0]["name"], "Operadora Analytics")

    def test_link_semantics_resolve_provider_name_and_keep_ids_technical(self):
        self.operator.source_config_item_id = "6804"
        self.operator.name = "CLARO"
        self.operator.save()
        self.link.isp_raw = "6804"
        self.link.save()
        semantic = LinkSemantics.from_link(self.link)
        self.assertEqual(semantic.provider_name, "CLARO")
        self.assertNotEqual(semantic.provider_name, "6804")
        self.assertEqual(semantic.provider_config_item_id, "6804")
        self.assertIsNone(semantic.project_name)
        self.assertIsNone(semantic.contract_operator_name)

    def test_unresolved_numeric_isp_is_never_a_provider_label(self):
        self.link.operator = None
        self.link.isp_raw = "10097"
        self.link.save()
        semantic = LinkSemantics.from_link(self.link)
        self.assertEqual(semantic.provider_name, "Não classificado")
        self.assertEqual(semantic.provider_config_item_id, "10097")
        self.assertEqual(semantic.isp_raw, "10097")

    def test_cockpit_provider_options_use_names_not_raw_config_item_ids(self):
        self.operator.source_config_item_id = "6804"
        self.operator.name = "CLARO"
        self.operator.save()
        self.link.isp_raw = "6804"
        self.link.save()
        ticket = self._ticket(120, "OPEN", True, 5)
        OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=self.link.source_config_item_id, cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "provider", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {},
        }, user=self.user)
        self.assertEqual(result["options"]["provider"], ["CLARO"])
        self.assertEqual(result["breakdown"][0]["key"], "CLARO")
        self.assertNotIn("6804", result["options"]["provider"])

    def test_cockpit_options_and_status_labels_are_noc_only(self):
        noc = self._ticket(121, "OPEN", True, 5)
        noc.service_name = "Internet"
        noc.sla_name = "Resposta NOC"
        noc.external_state_name = "em atendimento"
        noc.save()
        OTRSTicketConfigItem.objects.create(ticket=noc, config_item_id=self.link.source_config_item_id, cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        outside = self._ticket(122, "OPEN", True, 6)
        outside.queue_name = "OUTRA FILA"
        outside.service_name = "Firewall"
        outside.sla_name = "SLA global"
        outside.external_state_name = "global state"
        outside.save()

        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "source_state", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {},
        }, user=self.user)
        self.assertEqual(result["options"]["service"], ["Internet"])
        self.assertEqual(result["options"]["sla"], ["Resposta NOC"])
        self.assertEqual(result["options"]["source_state"], [{"value": "em atendimento", "label": "Em atendimento"}])
        self.assertEqual(result["options"]["status"], [{"value": "OPEN", "label": "Em atendimento"}])
        self.assertEqual(result["breakdown"][0], {"key": "em atendimento", "dimension": "Em atendimento", "value": 1, "share": 100.0})
        self.assertEqual(result["value"], 1)

    def test_cockpit_page_starts_with_noc_filters_and_paged_results(self):
        response = self.client.get(reverse("cockpit"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "NOC ATACADO")
        self.assertContains(response, "Situa&ccedil;&atilde;o OTRS")
        self.assertContains(response, "Provedor")
        self.assertContains(response, "data-detail-pagination")
        self.assertNotContains(response, "Firewall")
        self.assertContains(response, 'name="termination_reason"')
        self.assertNotContains(response, "CodigoFechamento ainda")
        self.assertNotContains(response, "Projeto legado")
        self.assertNotContains(response, "Tickets por servi")
        self.assertContains(response, "Mais filtros")
        self.assertContains(response, 'name="priority"')
        self.assertContains(response, 'data-more-filters')
        self.assertNotContains(response, 'data-snapshot-kpis')
        self.assertContains(response, 'Visualizar por')
        self.assertContains(response, 'data-dimension-grid')
        self.assertContains(response, 'data-ticket-drawer')
        self.assertContains(response, 'data-ticket-search')
        self.assertContains(response, 'data-initial-loading="true"')
        self.assertContains(response, 'PDF anal&iacute;tico')
        self.assertContains(response, 'Excel completo')
        self.assertContains(response, 'CSV de tickets')
        self.assertNotContains(response, '<details class="cockpit-more-filters" data-more-filters open')
        self.assertNotContains(response, "Operadora contratada")
        self.assertNotContains(response, "name=\"operator\"")

    def test_cockpit_top_ten_and_detail_are_database_paged(self):
        for index in range(27):
            provider = CMDBOperatorMirror.objects.create(source_config_item_id=f"cockpit-provider-{index}", name=f"NOC Provider {index:02d}")
            link = CMDBLinkMirror.objects.create(source_config_item_id=f"cockpit-link-{index}", source_number=f"NOC-{index}", operator=provider, state="SP", legacy_project="Projeto A")
            ticket = self._ticket(300 + index, "OPEN", True, 5)
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link, relation_type="DependsOn", relation_state="Valid")
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "provider", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "detail_page": 2, "filters": {},
        }, user=self.user)
        self.assertEqual(len(result["breakdown"]), 10)
        self.assertEqual(result["detail_total"], 27)
        self.assertEqual(result["detail_page"], 2)
        self.assertEqual(result["detail_pages"], 2)
        self.assertEqual(len(result["detail"]), 2)
        self.assertEqual(result["ticket_result_total"], 27)
        self.assertEqual(result["ticket_total"], 27)
        self.assertEqual(len(result["tickets"]), 25)
        exported = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "provider", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "detail_page": 2, "filters": {},
        }, user=self.user, include_all_rows=True)
        self.assertEqual(len(exported["detail"]), 27)
        self.assertEqual(len(exported["tickets"]), 27)
        self.assertEqual(exported["value"], 27)

    def test_cockpit_created_metric_excludes_canceled_and_resolution_excludes_it(self):
        self._ticket(810, "RESOLVED", False, 3)
        self._ticket(811, "CLOSED", False, 4)
        self._ticket(812, "CANCELED", False, 5)
        payload = {"group_by": "status", "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {}}
        service = AnalyticsQueryService()
        created = service.query({**payload, "metric": "tickets"}, user=self.user)
        resolved = service.query({**payload, "metric": "resolved_closed_tickets"}, user=self.user)
        canceled = service.query({**payload, "metric": "canceled_tickets"}, user=self.user)
        self.assertEqual(created["value"], 2)
        self.assertEqual(resolved["value"], 2)
        self.assertEqual({row["dimension"] for row in resolved["detail"]}, {"Resolvido", "Encerrado"})
        self.assertEqual(canceled["value"], 1)

    def test_operational_state_universe_and_explicit_cancel_unknown_queries(self):
        states = ("OPEN", "PENDING", "SCHEDULED", "RESOLVED", "CLOSED", "CANCELED", "UNKNOWN")
        for index, status in enumerate(states):
            self._ticket(1800 + index, status, status in {"OPEN", "PENDING", "SCHEDULED"}, index + 1)
        service = AnalyticsQueryService()
        base = {"start": "2026-10-01", "end": "2026-10-31", "group_by": "status", "filters": {}}
        created = service.query({**base, "metric": "tickets"}, user=self.user, include_all_rows=True)
        canceled = service.query({**base, "metric": "canceled_tickets"}, user=self.user)
        unknown = service.query({**base, "metric": "tickets", "filters": {"status": "UNKNOWN"}}, user=self.user)
        explicit_canceled = service.query({**base, "metric": "tickets", "filters": {"status": "CANCELED"}}, user=self.user)
        backlog = service.query({**base, "metric": "backlog"}, user=self.user)
        resolved = service.query({**base, "metric": "resolved_tickets"}, user=self.user)
        closed = service.query({**base, "metric": "closed_tickets"}, user=self.user)
        both = service.query({**base, "metric": "resolved_closed_tickets"}, user=self.user)
        self.assertEqual(created["value"], 5)
        self.assertEqual({row["key"] for row in created["detail"]}, {"OPEN", "PENDING", "SCHEDULED", "RESOLVED", "CLOSED"})
        self.assertEqual((canceled["value"], unknown["value"], explicit_canceled["value"]), (1, 0, 0))
        self.assertEqual([row["value"] for row in canceled["options"]["status"]], ["CANCELED"])
        self.assertFalse({"CANCELED", "UNKNOWN"} & {row["value"] for row in created["options"]["status"]})
        self.assertEqual(backlog["value"], 3)
        self.assertEqual((resolved["value"], closed["value"], both["value"]), (1, 1, 2))

    def test_terminal_metrics_use_terminal_date_and_reason_aliases_merge(self):
        resolved = self._ticket(813, "RESOLVED", False, 2)
        resolved.source_created_at = datetime(2026, 9, 2, tzinfo=datetime_timezone.utc)
        resolved.terminal_state_at = datetime(2026, 10, 8, tzinfo=datetime_timezone.utc)
        resolved.closing_code_raw = "PendenciaFinanceira"
        resolved.save()
        closed = self._ticket(814, "CLOSED", False, 3)
        closed.closing_code_raw = "Pendência Financeira"
        closed.save()
        canceled = self._ticket(815, "CANCELED", False, 4)
        canceled.closing_code_raw = "ChamadoCancelado"
        canceled.save()
        payload = {"group_by": "termination_reason", "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {}}
        service = AnalyticsQueryService()
        closed_result = service.query({**payload, "metric": "closed_tickets", "filters": {"termination_reason": "Pendência Financeira"}}, user=self.user)
        success = service.query({**payload, "metric": "resolved_closed_tickets"}, user=self.user)
        canceled_result = service.query({**payload, "metric": "canceled_tickets"}, user=self.user)
        created = service.query({**payload, "metric": "tickets"}, user=self.user)
        self.assertEqual(closed_result["value"], 1)
        self.assertEqual(success["value"], 2)
        self.assertEqual(canceled_result["value"], 1)
        self.assertEqual(created["value"], 1)
        self.assertEqual(success["detail"][0]["dimension"], "Pendência Financeira")
        self.assertEqual(success["period_date_label"], "data de encerramento")
        self.assertEqual(closed_result["tickets"][0]["base_date"], "03/10/2026 00:00")
        self.assertTrue(any(value.startswith("Data de encerramento entre 01/10/2026 e 31/10/2026") for value in closed_result["tickets"][0]["evidence"]))

    def test_massive_failure_reason_filter_and_monthly_series_use_all_raw_aliases(self):
        raws = ("FalhaMassiva", "Defeito Massivo", "Falha Massiva")
        for index, (month, raw) in enumerate(zip((7, 8, 9), raws), 1):
            ticket = self._ticket(2600 + index, "CLOSED", False, 5)
            ticket.source_created_at = datetime(2026, month, 5, tzinfo=datetime_timezone.utc)
            ticket.terminal_state_at = datetime(2026, month, 6, tzinfo=datetime_timezone.utc)
            ticket.closing_code_raw = raw
            ticket.save()

        service = AnalyticsQueryService()
        grouped = service.query({
            "metric": "closed_tickets", "group_by": "termination_reason",
            "start": "2026-07-01", "end": "2026-09-30", "filters": {},
        }, user=self.user)
        filtered = service.query({
            "metric": "tickets", "group_by": "month",
            "start": "2026-07-01", "end": "2026-09-30",
            "filters": {"termination_reason": "falha_massiva"},
        }, user=self.user, include_all_rows=True)
        self.assertEqual([(row["dimension"], row["value"]) for row in grouped["detail"]], [("Falha Massiva", 3)])
        self.assertIn({"value": "falha_massiva", "label": "Falha Massiva"}, grouped["options"]["termination_reason"])
        self.assertEqual(filtered["value"], 3)
        self.assertEqual([row["value"] for row in filtered["series"]], [1, 1, 1])
        self.assertEqual({row["reason"] for row in filtered["tickets"]}, {"Falha Massiva"})
        self.assertEqual({row["closing_code_raw"] for row in filtered["tickets"]}, set(raws))

    def test_transition_reason_filters_group_options_and_saved_views_are_canonical(self):
        raws = {
            "equipamento_travado": (
                "Equipamento travado/reinicilizado", "EquipamentoTravado",
            ),
            "rompimento_fibra": ("Rompimento/manut Fibra", "RompimentoFibra"),
            "chamado_cancelado": ("Cancelado", "ChamadoCancelado"),
            "conexao_incorreta_cabo_interface": (
                "Conexão Incorreta - Cabeamento/porta", "ConexaoIncorreta",
            ),
        }
        ext = 1800
        for aliases in raws.values():
            for month, raw in zip((7, 9), aliases):
                ticket = self._ticket(ext, "CLOSED", False, 5)
                ticket.source_created_at = datetime(2026, month, 5, tzinfo=datetime_timezone.utc)
                ticket.terminal_state_at = ticket.source_created_at
                ticket.closing_code_raw = raw
                ticket.save(update_fields=("closing_code_raw", "source_created_at", "terminal_state_at"))
                ext += 1

        payload = {
            "metric": "closed_tickets", "start": "2026-07-01", "end": "2026-09-30",
            "group_by": "termination_reason", "filters": {},
        }
        result = AnalyticsQueryService().query(payload, user=self.user)
        reason_options = result["options"]["termination_reason"]
        self.assertEqual(len(reason_options), len({item["value"] for item in reason_options}))
        self.assertEqual(len(reason_options), len({item["label"].casefold() for item in reason_options}))
        for key, label in (
            ("equipamento_travado", "Equipamento Travado"),
            ("rompimento_fibra", "Rompimento de Fibra"),
            ("chamado_cancelado", "Chamado Cancelado"),
            ("conexao_incorreta_cabo_interface", "Conexão Incorreta - Cabo/Interface"),
        ):
            self.assertEqual(sum(item["value"] == key and item["label"] == label for item in reason_options), 1)

        service = AnalyticsQueryService()
        for key, aliases in raws.items():
            with self.subTest(key=key):
                filtered = service.query(
                    {**payload, "filters": {"termination_reason": aliases[0]}},
                    user=self.user, include_all_rows=True,
                )
                self.assertEqual(filtered["value"], 2)
                self.assertEqual(filtered["normalized_query"]["filters"]["termination_reason"], key)
                self.assertEqual({row["closing_code_raw"] for row in filtered["tickets"]}, set(aliases))

        monthly = service.query(
            {**payload, "metric": "closed_tickets", "group_by": "month",
             "filters": {"termination_reason": "Equipamento travado/reinicilizado"}},
            user=self.user,
        )
        self.assertEqual([row["value"] for row in monthly["series"]], [1, 0, 1])

        from .models import SavedAnalyticsView
        saved = SavedAnalyticsView.objects.create(
            user=self.user, name="Legado", configuration={"filters": {"termination_reason": "EquipamentoTravado"}},
        )
        listed = self.client.post(
            reverse("analytics-saved-views"), data=json.dumps({"action": "list"}),
            content_type="application/json",
        )
        self.assertEqual(listed.status_code, 200)
        saved_view = next(item for item in listed.json()["views"] if item["id"] == saved.pk)
        self.assertEqual(saved_view["configuration"]["filters"]["termination_reason"], "equipamento_travado")

        legacy_url = self.client.post(
            reverse("analytics-query"), data=json.dumps({**payload, "filters": {"termination_reason": "EquipamentoTravado"}}),
            content_type="application/json",
        )
        self.assertEqual(legacy_url.status_code, 200)
        legacy_payload = legacy_url.json()
        self.assertEqual(legacy_payload["normalized_query"]["filters"]["termination_reason"], "equipamento_travado")
        self.assertEqual(legacy_payload["value"], 2)
        self.assertEqual(
            [item for item in legacy_payload["options"]["termination_reason"] if item["value"] == "equipamento_travado"],
            [{"value": "equipamento_travado", "label": "Equipamento Travado"}],
        )
        self.assertEqual(len(legacy_payload["detail"]), 1)
        self.assertEqual(legacy_payload["detail"][0]["dimension"], "Equipamento Travado")
        self.assertEqual(legacy_payload["detail"][0]["value"], 2)

    def test_kpis_are_scoped_to_the_selected_query_without_dashboard_snapshot_cards(self):
        self._ticket(816, "OPEN", True, 3)
        self._ticket(817, "CLOSED", False, 4)
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "status", "period_from": "2026-10", "period_to": "2026-10",
            "period": "2026-10", "filters": {"status": "CLOSED"},
        }, user=self.user)
        self.assertEqual(result["value"], 1)
        self.assertEqual(result["kpis"][0]["value"], 1)
        self.assertNotIn("snapshot_kpis", result)
        self.assertEqual(result["provider_quality"]["unresolved"], 1)

    def test_cockpit_priority_filter_and_grouping_support_high_critical_shortcut(self):
        high = self._ticket(820, "OPEN", True, 5)
        high.priority_name = "Alto"
        high.save()
        critical = self._ticket(821, "CLOSED", False, 5)
        critical.priority_name = "Crítico"
        critical.save()
        normal = self._ticket(822, "CANCELED", False, 5)
        normal.priority_name = "Médio"
        normal.save()
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "priority", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {"priority": "high-critical"},
        }, user=self.user)
        self.assertEqual(result["value"], 2)
        self.assertEqual({row["dimension"] for row in result["detail"]}, {"Alto", "Crítico"})

    def test_cockpit_ticket_totals_are_distinct_and_recurrences_count_distinct_links(self):
        second_provider = CMDBOperatorMirror.objects.create(source_config_item_id="cockpit-second-provider", name="Provider B")
        second_link = CMDBLinkMirror.objects.create(source_config_item_id="cockpit-second-link", source_number="NOC-B", operator=second_provider, state="RJ", legacy_project="Projeto A")
        first = self._ticket(530, "OPEN", True, 5)
        second = self._ticket(531, "RESOLVED", False, 6)
        for ticket in (first, second):
            for index, link in enumerate((self.link, second_link)):
                OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link, relation_type=f"DependsOn-{ticket.pk}-{index}", relation_state="Valid")
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "provider", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {},
        }, user=self.user)
        self.assertEqual(result["value"], 2)
        self.assertEqual({row["dimension"]: row["value"] for row in result["detail"]}, {"Operadora Analytics": 2, "Provider B": 2})
        recurrences = AnalyticsQueryService().query({
            "metric": "recurrences", "group_by": "provider", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {},
        }, user=self.user)
        self.assertEqual(recurrences["value"], 2)
        self.assertEqual({row["dimension"]: row["value"] for row in recurrences["detail"]}, {"Operadora Analytics": 1, "Provider B": 1})

    def test_cockpit_otr_status_values_keep_raw_filter_keys_and_friendly_labels(self):
        states = (
            ("OPEN", "em atendimento", "Em atendimento"),
            ("PENDING", "Pendente", "Pendente"),
            ("SCHEDULED", "agendado", "Agendado"),
            ("RESOLVED", "resolvido", "Resolvido"),
            ("CLOSED", "closed successful", "Encerrado com sucesso"),
            ("CLOSED", "Encerrado com viabilidade", "Encerrado com viabilidade"),
            ("CANCELED", "cancelado", "Cancelado"),
            ("UNKNOWN", "estado novo", "estado novo"),
        )
        for index, (status, source_state, label) in enumerate(states):
            ticket = self._ticket(400 + index, status, status in {"OPEN", "PENDING", "SCHEDULED"}, index + 5)
            ticket.external_state_name = source_state
            ticket.save()
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=self.link.source_config_item_id, cmdb_link=self.link, relation_type=f"DependsOn-{index}", relation_state="Valid")
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "source_state", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {},
        }, user=self.user)
        labels = {row["key"]: row["dimension"] for row in result["detail"]}
        self.assertEqual(labels["em atendimento"], "Em atendimento")
        self.assertEqual(labels["closed successful"], "Encerrado com sucesso")
        self.assertEqual(labels["Encerrado com viabilidade"], "Encerrado com viabilidade")
        self.assertNotIn("estado novo", labels)
        unknown = AnalyticsQueryService().query({"metric":"tickets", "group_by":"source_state", "period_from":"2026-10", "period_to":"2026-10", "filters":{"status":"UNKNOWN"}}, user=self.user)
        self.assertEqual(unknown["value"], 0)
        self.assertEqual(unknown["detail"], [])
        status_options = {row["value"]: row["label"] for row in result["options"]["status"]}
        self.assertEqual(status_options, {"OPEN": "Em atendimento", "PENDING": "Pendente", "SCHEDULED": "Agendado", "RESOLVED": "Resolvido", "CLOSED": "Encerrado"})

    def test_noc_period_and_current_backlog_have_separate_date_semantics(self):
        september = self._ticket(111, "CLOSED", False, 5)
        september.source_created_at = datetime(2026, 9, 5, tzinfo=datetime_timezone.utc)
        september.save()
        october_open = self._ticket(112, "OPEN", True, 5)
        october_pending = self._ticket(113, "PENDING", True, 8)
        october_scheduled = self._ticket(118, "SCHEDULED", True, 8)
        october_resolved = self._ticket(114, "RESOLVED", False, 9)
        other_queue = self._ticket(115, "OPEN", True, 9)
        other_queue.queue_name = "OUTRA FILA"
        other_queue.save()
        october_closed = self._ticket(116, "CLOSED", False, 10)
        october_canceled = self._ticket(117, "CANCELED", False, 11)
        for ticket in (september, october_open, october_pending, october_scheduled, october_resolved, other_queue, october_closed, october_canceled):
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=self.link.source_config_item_id, cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        self.assertFalse(TicketQueryService.noc_tickets().filter(pk=other_queue.pk).exists())
        payload = {"period_from": "2026-09", "period_to": "2026-10", "period": "2026-09", "filters": {}, "group_by": "status"}
        service = AnalyticsQueryService()
        created = service.query({**payload, "metric": "tickets"}, user=self.user)
        backlog = service.query({**payload, "metric": "backlog"}, user=self.user)
        self.assertEqual(created["value"], 6)
        self.assertEqual(backlog["value"], 3)
        scheduled = service.query({**payload, "metric": "scheduled_tickets"}, user=self.user)
        self.assertEqual(scheduled["value"], 1)
        self.assertTrue(backlog["meta"]["backlog_is_current_snapshot"])

    def test_exact_date_bounds_reason_provider_uf_feed_all_analysis_outputs_and_exports(self):
        inside = self._ticket(1701, "CLOSED", False, 5)
        inside.source_created_at = datetime(2026, 10, 20, tzinfo=datetime_timezone.utc)
        inside.terminal_state_at = datetime(2026, 11, 1, 2, 59, 59, 999999, tzinfo=datetime_timezone.utc)
        inside.closing_code_raw = "PendenciaFinanceira"
        inside.save()
        at_end_exclusive = self._ticket(1702, "CLOSED", False, 5)
        at_end_exclusive.terminal_state_at = datetime(2026, 11, 1, 3, 0, tzinfo=datetime_timezone.utc)
        at_end_exclusive.closing_code_raw = "Pendência Financeira"
        at_end_exclusive.save()
        wrong_reason = self._ticket(1703, "CLOSED", False, 5)
        wrong_reason.terminal_state_at = datetime(2026, 10, 31, 20, 0, tzinfo=datetime_timezone.utc)
        wrong_reason.closing_code_raw = "Falha Técnica"
        wrong_reason.save()
        for ticket in (inside, at_end_exclusive, wrong_reason):
            OTRSTicketConfigItem.objects.create(
                ticket=ticket, config_item_id=self.link.source_config_item_id,
                cmdb_link=self.link, relation_type=f"Boundary-{ticket.external_id}", relation_state="Valid")
        payload = {
            "metric": "closed_tickets", "start": "2026-10-01", "end": "2026-10-31",
            "group_by": "provider", "filters": {
                "closing_reason": "Pendência Financeira", "provider": self.operator.name,
                "uf": "sp", "status": "closed", "sla": "Resposta NOC",
                "priority": "3 normal",
            },
        }
        expected = AnalyticsQueryService().query(payload, user=self.user, include_all_rows=True)
        self.assertEqual(expected["value"], 1)
        self.assertEqual(expected["kpis"][0]["value"], 1)
        self.assertEqual(sum(row["value"] for row in expected["series"]), 1)
        self.assertEqual(sum(row["value"] for row in expected["breakdown"]), 1)
        self.assertEqual(sum(row["value"] for row in expected["detail"]), 1)
        self.assertEqual(expected["detail"][0]["dimension"], self.operator.name)
        self.assertEqual(expected["period_date_label"], "data de encerramento")
        self.assertEqual(expected["tickets"][0]["base_date"], "31/10/2026 23:59")

        api = self.client.post(reverse("analytics-query"), data=json.dumps(payload), content_type="application/json")
        self.assertEqual(api.status_code, 200)
        self.assertEqual(api.json()["value"], expected["value"])
        csv_response = self.client.post(reverse("analytics-export"), data=json.dumps({**payload, "format": "csv"}), content_type="application/json")
        self.assertEqual(csv_response.status_code, 200)
        csv_text = csv_response.content.decode("utf-8-sig")
        self.assertEqual(len(csv_text.splitlines()), 2)
        self.assertIn(self.operator.name, csv_text)
        pdf_response = self.client.post(reverse("analytics-export"), data=json.dumps({**payload, "format": "pdf"}), content_type="application/json")
        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_response.content.startswith(b"%PDF-"))
        xlsx_response = self.client.post(reverse("analytics-export"), data=json.dumps({**payload, "format": "xlsx"}), content_type="application/json")
        self.assertEqual(xlsx_response.status_code, 200)
        workbook = __import__("openpyxl").load_workbook(BytesIO(xlsx_response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Resumo", "Distribuição", "Tickets"])
        self.assertEqual(workbook["Distribuição"].max_row, 2)
        self.assertEqual(workbook["Tickets"].max_row, 2)

    def test_provider_ranking_excludes_unresolved_bucket_but_total_keeps_ticket(self):
        ticket = self._ticket(1704, "CLOSED", False, 5)
        ticket.terminal_state_at = datetime(2026, 10, 15, 12, tzinfo=datetime_timezone.utc)
        ticket.closing_code_raw = "PendenciaFinanceira"
        ticket.save()
        result = AnalyticsQueryService().query({
            "metric": "closed_tickets", "start": "2026-10-01", "end": "2026-10-31",
            "group_by": "provider", "filters": {"closing_reason": "Pend\u00eancia Financeira"},
        }, user=self.user, include_all_rows=True)
        self.assertEqual(result["value"], 1)
        self.assertEqual(result["provider_quality"]["unresolved"], 1)
        self.assertEqual(result["provider_quality"]["coverage_percent"], 0)
        self.assertEqual(result["provider_quality"]["unresolved_causes"]["no_config_item"], 1)
        self.assertEqual(result["detail"], [])
        self.assertFalse(any(row["dimension"] == "N\u00e3o classificado" for row in result["breakdown"]))

    def test_provider_quality_reports_unresolved_link_causes(self):
        self._ticket(1710, "OPEN", True, 5)
        missing_mirror = self._ticket(1711, "OPEN", True, 6)
        OTRSTicketConfigItem.objects.create(ticket=missing_mirror, config_item_id="missing-mirror", cmdb_link=None, relation_type="DependsOn", relation_state="Valid")
        invalid_link = CMDBLinkMirror.objects.create(source_config_item_id="invalid-provider-link", is_present_in_source=False)
        invalid_ticket = self._ticket(1712, "OPEN", True, 7)
        OTRSTicketConfigItem.objects.create(ticket=invalid_ticket, config_item_id="invalid-provider-link", cmdb_link=invalid_link, relation_type="DependsOn", relation_state="Valid")
        no_provider = CMDBLinkMirror.objects.create(source_config_item_id="link-without-provider", state="SP")
        no_provider_ticket = self._ticket(1713, "OPEN", True, 8)
        OTRSTicketConfigItem.objects.create(ticket=no_provider_ticket, config_item_id="link-without-provider", cmdb_link=no_provider, relation_type="DependsOn", relation_state="Valid")

        result = AnalyticsQueryService().query({
            "metric": "tickets", "start": "2026-10-01", "end": "2026-10-31",
            "group_by": "provider", "filters": {},
        }, user=self.user)

        self.assertEqual(result["value"], 4)
        self.assertEqual(result["provider_quality"]["unresolved"], 4)
        self.assertEqual(result["provider_quality"]["unresolved_causes"], {
            "no_config_item": 1, "missing_mirror": 1, "invalid_link": 1, "link_without_provider": 1,
        })

    def test_ticket_evidence_uses_the_filtered_analysis_queryset_and_safe_cmdb_fields(self):
        matches = []
        for external_id, day in ((1714, 5), (1715, 6)):
            ticket = self._ticket(external_id, "CLOSED", False, day)
            ticket.closing_code_raw = "PendenciaFinanceira"
            ticket.save()
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=self.link.source_config_item_id,
                                                cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
            matches.append(ticket)
        other = self._ticket(1716, "CLOSED", False, 7)
        other.closing_code_raw = "Falha Tecnica"
        other.save()
        OTRSTicketConfigItem.objects.create(ticket=other, config_item_id=self.other_link.source_config_item_id,
                                            cmdb_link=self.other_link, relation_type="DependsOn", relation_state="Valid")
        result = AnalyticsQueryService().query({
            "metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "group_by": "provider",
            "ticket_page_size": 25, "filters": {"termination_reason": "PendenciaFinanceira", "provider": self.operator.name},
        }, user=self.user)
        self.assertEqual(result["value"], 2)
        self.assertEqual(result["ticket_result_total"], 2)
        self.assertEqual(result["ticket_total"], 2)
        self.assertEqual({row["ticket"] for row in result["tickets"]}, {ticket.ticket_number for ticket in matches})
        self.assertEqual(result["description"]["title"], "An\u00e1lise de Pend\u00eancia Financeira")
        self.assertIn("Criado entre 01/10/2026 e 31/10/2026", result["tickets"][0]["evidence"])
        self.assertEqual(result["tickets"][0]["source"], "OTRS + CMDB")
        self.assertEqual(result["tickets"][0]["links"][0]["config_item_id"], self.link.source_config_item_id)
        self.assertIn("Estado operacional v\u00e1lido", result["tickets"][0]["evidence"])
        self.assertNotIn("customer_id", result["tickets"][0])
        searched = AnalyticsQueryService().query({
            "metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "group_by": "provider",
            "ticket_search": "AN-1", "filters": {"termination_reason": "PendenciaFinanceira"},
        }, user=self.user)
        self.assertEqual(searched["ticket_total"], 2)
        self.assertEqual(len(searched["tickets"]), 2)

    def test_automatic_previous_period_comparison_populates_contextual_kpi_trends(self):
        current = self._ticket(1717, "OPEN", True, 10)
        OTRSTicketConfigItem.objects.create(ticket=current, config_item_id=self.link.source_config_item_id,
                                            cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        previous_identified = self._ticket(1718, "OPEN", True, 8)
        previous_identified.source_created_at = datetime(2026, 9, 8, tzinfo=datetime_timezone.utc)
        previous_identified.save()
        OTRSTicketConfigItem.objects.create(ticket=previous_identified, config_item_id=self.link.source_config_item_id,
                                            cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        previous_unresolved = self._ticket(1719, "OPEN", True, 9)
        previous_unresolved.source_created_at = datetime(2026, 9, 9, tzinfo=datetime_timezone.utc)
        previous_unresolved.save()
        result = AnalyticsQueryService().query({
            "metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "group_by": "provider", "filters": {},
        }, user=self.user)
        self.assertEqual(result["comparison"]["previous_start"], "2026-09-01")
        self.assertEqual(result["comparison"]["previous_end"], "2026-09-30")
        self.assertEqual((result["comparison"]["current_value"], result["comparison"]["previous_value"]), (1, 2))
        self.assertEqual(result["kpis"][0]["trend"]["absolute_change"], -1)
        self.assertEqual(result["kpis"][0]["trend"]["percentage_change"], -50.0)
        coverage = next(item for item in result["kpis"] if item["key"] == "coverage")
        self.assertEqual(coverage["trend"]["absolute_change"], 50.0)
        self.assertEqual(coverage["trend"]["unit"], "p.p.")
        self.assertEqual([spec["type"] for spec in result["chart_specs"] if spec["key"] == "provider-ranking"], ["ranking"])
        ranking = next(spec for spec in result["chart_specs"] if spec["key"] == "provider-ranking")
        self.assertIn("share", ranking["details"][0])
        temporal = next(spec for spec in result["chart_specs"] if spec["key"] == "temporal")
        self.assertEqual(temporal["details"][9]["label"], "10/10/2026")
        self.assertEqual(temporal["details"][7]["previous_value"], 1)

    def test_pdf_report_uses_dynamic_filtered_charts_ranking_and_unicode_ticket_appendix(self):
        first = self._ticket(1723, "CLOSED", False, 5)
        first.closing_code_raw = "PendenciaFinanceira"
        first.save()
        second_operator = CMDBOperatorMirror.objects.create(source_config_item_id="pdf-provider", name="Provedor São José")
        second_link = CMDBLinkMirror.objects.create(source_config_item_id="pdf-link", source_number="PDF-2", operator=second_operator, state="SP")
        second = self._ticket(1724, "CLOSED", False, 6)
        second.closing_code_raw = "Pendência Financeira"
        second.save()
        OTRSTicketConfigItem.objects.create(ticket=first, config_item_id=self.link.source_config_item_id,
                                            cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        OTRSTicketConfigItem.objects.create(ticket=second, config_item_id=second_link.source_config_item_id,
                                            cmdb_link=second_link, relation_type="DependsOn", relation_state="Valid")
        base = {"metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "group_by": "provider",
                "filters": {"termination_reason": "Pendência Financeira"}}
        result = AnalyticsQueryService().query(base, user=self.user, include_all_rows=True)
        self.assertEqual(result["value"], 2)
        self.assertEqual([spec["key"] for spec in result["chart_specs"]], ["temporal", "monthly-calendar", "provider-ranking"])

        from aplicativos.painel.views import _pdf_report
        captured = {}
        class StoryCaptureDocument:
            def __init__(self, buffer, **kwargs):
                self.buffer = buffer
            def build(self, story, **kwargs):
                captured["story"] = story

        def render_story(query_result):
            with patch("reportlab.platypus.SimpleDocTemplate", StoryCaptureDocument):
                _pdf_report("24/10/2026 14:14", self.user.get_username(), query_result,
                            [("Motivo", "Pendência Financeira")])
            texts = []
            def collect(item):
                if isinstance(getattr(item, "text", None), str):
                    texts.append(item.text)
                for row in getattr(item, "_cellvalues", ()):
                    for cell in row:
                        collect(cell)
            for item in captured["story"]:
                collect(item)
            return "\n".join(texts)

        report_text = render_story(result)
        self.assertIn("RELATÓRIO ANALÍTICO DO COCKPIT", report_text)
        self.assertIn("Análise de Pendência Financeira", report_text)
        self.assertIn(self.user.get_username(), report_text)
        self.assertIn("24/10/2026 14:14", report_text)
        self.assertIn("Ranking de provedores", report_text)
        self.assertIn("Detalhe de Outubro/2026", report_text)
        self.assertIn("Seg", report_text)
        self.assertIn("Posição", report_text)
        self.assertIn("Participação", report_text)
        self.assertNotIn("Distribuição completa por provedor", report_text)
        self.assertIn("ANEXO - TICKETS DO RESULTADO", report_text)
        self.assertIn("A-1723", report_text)
        self.assertIn("A-1724", report_text)
        from reportlab.pdfbase import pdfmetrics
        self.assertGreater(pdfmetrics.stringWidth("Período · Pendência Financeira · Criação · Não identificado · Página", "ArionVera", 8), 0)

        filtered = AnalyticsQueryService().query({**base, "filters": {**base["filters"], "provider": self.operator.name}},
                                                user=self.user, include_all_rows=True)
        self.assertEqual([spec["key"] for spec in filtered["chart_specs"]], ["temporal", "monthly-calendar"])
        filtered_text = render_story(filtered)
        self.assertIn("Evolução geral do período", filtered_text)
        self.assertIn("Detalhe de Outubro/2026", filtered_text)
        self.assertNotIn("Ranking de provedores", filtered_text)

    def test_analytics_chart_grain_changes_with_date_span_and_fixed_dimension_is_hidden(self):
        for external_id, day in ((1720, 1), (1721, 3), (1722, 8)):
            ticket = self._ticket(external_id, "OPEN", True, day)
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=f"chart-{external_id}", cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        service = AnalyticsQueryService()
        daily = service.query({
            "metric": "tickets", "start": "2026-10-01", "end": "2026-10-31",
            "group_by": "provider", "filters": {"provider": self.operator.name},
        }, user=self.user)
        weekly = service.query({
            "metric": "tickets", "start": "2026-10-01", "end": "2026-11-30",
            "group_by": "provider", "filters": {},
        }, user=self.user)
        monthly = service.query({
            "metric": "tickets", "start": "2026-01-01", "end": "2026-06-30",
            "group_by": "provider", "filters": {},
        }, user=self.user)
        self.assertEqual(daily["time_grain"], "day")
        self.assertEqual(len(daily["series"]), 31)
        self.assertEqual([spec["key"] for spec in daily["chart_specs"]], ["temporal", "monthly-calendar"])
        self.assertEqual(weekly["time_grain"], "week")
        self.assertEqual(monthly["time_grain"], "month")

    def test_overview_granularity_totals_and_latest_month_selection(self):
        for external_id, day in ((1801, 1), (1802, 3), (1803, 8)):
            self._ticket(external_id, "OPEN", True, day)
        service = AnalyticsQueryService()
        daily = service.query({"metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "filters": {}}, user=self.user)
        weekly = service.query({"metric": "tickets", "start": "2026-09-01", "end": "2026-11-01", "filters": {}}, user=self.user)
        monthly = service.query({"metric": "tickets", "start": "2025-10-01", "end": "2026-10-31", "filters": {}}, user=self.user)

        self.assertEqual((daily["time_grain"], weekly["time_grain"], monthly["time_grain"]), ("day", "week", "month"))
        for result in (daily, weekly, monthly):
            self.assertEqual(sum(item["value"] for item in result["series"]), result["value"])
            temporal = next(spec for spec in result["chart_specs"] if spec["type"] == "temporal_overview")
            self.assertEqual(temporal["grain"], result["time_grain"])
        self.assertEqual(monthly["temporal"]["selected_month"], "2026-10")
        self.assertEqual(monthly["temporal"]["month_summary"]["total"], 3)
        self.assertEqual(sum(day["value"] or 0 for day in monthly["temporal"]["calendar_days"]), 3)
        selected = service.query({"metric": "tickets", "start": "2025-10-01", "end": "2026-10-31", "selected_month": "2026-09", "filters": {}}, user=self.user)
        self.assertEqual(selected["normalized_query"]["selected_month"], "2026-09")
        self.assertEqual(selected["temporal"]["month_summary"]["total"], 0)

    def test_month_calendar_handles_leap_regular_30_and_31_day_months_with_monday_grid(self):
        months = ((1810, date(2024, 2, 29), 29), (1811, date(2025, 2, 28), 28),
                  (1812, date(2025, 4, 30), 30), (1813, date(2026, 10, 31), 31))
        for external_id, created, _days in months:
            ticket = self._ticket(external_id, "OPEN", True, 1)
            ticket.source_created_at = datetime.combine(created, datetime.min.time(), tzinfo=datetime_timezone.utc)
            ticket.save()
        service = AnalyticsQueryService()
        for _external_id, created, expected_days in months:
            result = service.query({"metric": "tickets", "start": created.replace(day=1).isoformat(),
                                    "end": created.isoformat(), "filters": {}}, user=self.user)
            temporal = result["temporal"]
            days = temporal["calendar_days"]
            in_month = [day for day in days if day["inside_month"]]
            self.assertEqual(len(in_month), expected_days)
            self.assertTrue(all(len(week) == 7 for week in temporal["calendar_weeks"]))
            self.assertTrue(all(date.fromisoformat(week[0]["date"]).weekday() == 0 for week in temporal["calendar_weeks"]))
            self.assertTrue(all(date.fromisoformat(week[-1]["date"]).weekday() == 6 for week in temporal["calendar_weeks"]))
            self.assertEqual(sum(day["value"] or 0 for day in days), 1)
            self.assertEqual(temporal["month_summary"]["total"], result["value"])

    def test_partial_month_disables_out_of_range_days_and_day_view_preserves_main_query(self):
        first = self._ticket(1820, "OPEN", True, 15)
        second = self._ticket(1821, "OPEN", True, 20)
        third = self._ticket(1822, "OPEN", True, 24)
        base = {"metric": "tickets", "start": "2026-10-15", "end": "2026-10-24", "filters": {}}
        service = AnalyticsQueryService()
        result = service.query(base, user=self.user)
        days = {item["date"]: item for item in result["temporal"]["calendar_days"]}
        self.assertFalse(days["2026-10-14"]["inside_query_period"])
        self.assertFalse(days["2026-10-25"]["inside_query_period"])
        self.assertTrue(days["2026-10-20"]["inside_query_period"])
        self.assertEqual(sum(item["value"] or 0 for item in days.values()), 3)

        day_view = service.query({**base, "ticket_day": "2026-10-20"}, user=self.user)
        self.assertEqual(day_view["value"], 3)
        self.assertEqual(day_view["ticket_result_total"], 3)
        self.assertEqual(day_view["ticket_day_total"], 1)
        self.assertEqual(day_view["ticket_total"], 1)
        self.assertEqual(day_view["tickets"][0]["ticket"], second.ticket_number)
        self.assertTrue(any("20/10/2026" in item for item in day_view["tickets"][0]["evidence"]))
        self.assertEqual(day_view["normalized_query"]["start"], "2026-10-15")
        self.assertEqual(day_view["normalized_query"]["end"], "2026-10-24")
        with self.assertRaises(ValueError):
            service.query({**base, "ticket_day": "2026-10-30"}, user=self.user)

    def test_calendar_respects_provider_filter_and_data_scope(self):
        ticket = self._ticket(1830, "OPEN", True, 5)
        OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=self.link.source_config_item_id,
                                            cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        hidden_provider = CMDBOperatorMirror.objects.create(source_config_item_id="calendar-hidden-op", name="Fora do escopo")
        hidden_link = CMDBLinkMirror.objects.create(source_config_item_id="calendar-hidden-link", state="AM", operator=hidden_provider, legacy_project="Projeto B")
        hidden = self._ticket(1831, "OPEN", True, 6)
        OTRSTicketConfigItem.objects.create(ticket=hidden, config_item_id=hidden_link.source_config_item_id,
                                            cmdb_link=hidden_link, relation_type="DependsOn", relation_state="Valid")
        scope = self.user.groups.get(name="Analista NOC").data_scope
        scope.mode = GroupDataScope.Mode.CUSTOM
        scope.save()
        allowed_operator = Operadora.objects.create(name=self.operator.name, normalized_name="operadora-analytics")
        allowed_project = Projeto.objects.create(name="Projeto A", normalized_name="projeto-a")
        scope.allowed_operators.set([allowed_operator])
        scope.allowed_projects.set([allowed_project])
        result = AnalyticsQueryService().query({"metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "filters": {}}, user=self.user)
        self.assertEqual(result["value"], 1)
        self.assertEqual(result["temporal"]["month_summary"]["total"], 1)
        self.assertEqual(sum(item["value"] or 0 for item in result["temporal"]["calendar_days"]), 1)
        filtered = AnalyticsQueryService().query({"metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "filters": {"provider": self.operator.name}}, user=self.user)
        self.assertEqual(filtered["value"], 1)
        self.assertEqual(filtered["temporal"]["month_summary"]["total"], 1)
        fixed_dimension = AnalyticsQueryService().query({"metric": "tickets", "start": "2026-10-01", "end": "2026-10-31",
                                                          "group_by": "uf", "filters": {"uf": "SP"}}, user=self.user)
        self.assertFalse(any(item.get("dimension") == "uf" for item in fixed_dimension["chart_specs"]))

    def test_cockpit_calendar_day_navigation_controls_are_present(self):
        response = self.client.get(reverse("cockpit"))
        self.assertEqual(response.status_code, 200)
        for selector in ("data-overview-grid", "data-calendar", "data-month-select", "data-calendar-grid",
                         "data-day-dialog", "data-day-view", "data-day-apply", "data-ticket-day-active"):
            self.assertContains(response, selector)

    def test_cockpit_detail_sections_are_independent_and_temporal_area_stays_visible(self):
        response = self.client.get(reverse("cockpit"))
        self.assertEqual(response.status_code, 200)
        body = response.content.decode("utf-8")
        temporal = body.index('aria-labelledby="cockpit-visual-title"')
        secondary = body.index('id="cockpit-secondary-sections"')
        self.assertLess(temporal, secondary)
        for key in ("dimensions", "distribution", "tickets"):
            self.assertIn(f'data-collapse-toggle="{key}"', body)
            self.assertIn(f'data-collapse-panel="{key}"', body)
        self.assertIn('data-details-global aria-expanded="true" aria-controls="cockpit-secondary-sections"', body)
        self.assertIn('data-analysis-title', body)
        self.assertIn('data-current-period', body)

        static_root = Path(settings.BASE_DIR) / "static"
        script = (Path(settings.BASE_DIR.parent) / "frontend_src" / "js" / "painel" / "cockpit.js").read_text(encoding="utf-8")
        styles = (static_root / "css" / "painel" / "cockpit.css").read_text(encoding="utf-8")
        self.assertIn('DETAIL_COLLAPSE_DEFAULTS = { dimensions: false, distribution: true, tickets: false }', script)
        self.assertIn("sessionStorage", script)
        self.assertIn('setDetailCollapsed("tickets", false)', script)
        self.assertIn('scrollIntoView({ behavior, block: "start" })', script)
        self.assertIn("@media (prefers-reduced-motion:reduce)", styles)

    def test_main_pages_render_without_known_encoding_damage(self):
        routes = ("dashboard", "cockpit", "mapa:index", "base_links:index", "rankings:index")
        damaged = ("per?odo", "cria??o", "Servi?o", "n?o", "reconstr?i", "hist?rico", "m?trica", "integra??o", "\ufffd", "Ãƒ")
        admin = get_user_model().objects.create_superuser(username="encoding-smoke", email="", password="safe-password-123")
        self.client.force_login(admin)
        for route in routes:
            response = self.client.get(reverse(route))
            self.assertEqual(response.status_code, 200, route)
            body = response.content.decode("utf-8")
            for token in damaged:
                self.assertNotIn(token, body, f"{route}: {token}")
        cockpit_script = (Path(settings.BASE_DIR.parent) / "frontend_src" / "js" / "painel" / "cockpit.js").read_text(encoding="utf-8")
        for token in ("servi?o", "per?odo", "Reincid?ncia", "\ufffd", "Ãƒ"):
            self.assertNotIn(token, cockpit_script)

    def test_analytics_uses_real_distinct_tickets_and_supported_dimensions(self):
        ticket = self._ticket(102, "OPEN", True, 5)
        for link in (self.link, self.other_link):
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link, relation_type="DependsOn", relation_state="Valid")
        base = {"period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {}}
        service = AnalyticsQueryService()
        result = service.query({**base, "metric": "tickets", "group_by": "operator"}, user=self.user)
        self.assertEqual(result["source"], "real")
        self.assertEqual(result["value"], 1)
        self.assertEqual(result["meta"]["demo"], False)
        self.assertEqual(result["breakdown"][0]["value"], 1)

    def test_analytics_rejects_unsupported_synthetic_sla_metrics(self):
        service = AnalyticsQueryService()
        with self.assertRaises(ValueError):
            service.query({"metric": "sla_met", "group_by": "operator", "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {}}, user=self.user)

    def test_custom_scope_excludes_unlinked_tickets(self):
        unlinked = self._ticket(103, "OPEN", True, 5)
        group = self.user.groups.get(name="Analista NOC")
        scope = group.data_scope
        scope.mode = GroupDataScope.Mode.CUSTOM
        scope.save()
        allowed_operator = Operadora.objects.create(name=self.operator.name, normalized_name=self.operator.name.casefold())
        allowed_project = Projeto.objects.create(name="Projeto A", normalized_name="projeto a")
        scope.allowed_operators.set([allowed_operator])
        scope.allowed_projects.set([allowed_project])
        from aplicativos.integracoes.services.analytics import TicketQueryService
        self.assertFalse(TicketQueryService.base(self.user).filter(pk=unlinked.pk).exists())

    def test_cockpit_options_and_results_preserve_custom_data_scope(self):
        group = self.user.groups.get(name="Analista NOC")
        scope = group.data_scope
        scope.mode = GroupDataScope.Mode.CUSTOM
        scope.save()
        allowed_operator = Operadora.objects.create(name=self.operator.name, normalized_name=self.operator.name.casefold())
        allowed_project = Projeto.objects.create(name="Projeto A", normalized_name="projeto a")
        scope.allowed_operators.set([allowed_operator])
        scope.allowed_projects.set([allowed_project])
        permitted = self._ticket(501, "OPEN", True, 5)
        hidden = self._ticket(502, "OPEN", True, 5)
        hidden_provider = CMDBOperatorMirror.objects.create(source_config_item_id="out-of-scope-provider", name="Fora do escopo")
        hidden_link = CMDBLinkMirror.objects.create(source_config_item_id="out-of-scope-link", source_number="OUT-1", operator=hidden_provider, legacy_project="Projeto B")
        for ticket, link in ((permitted, self.link), (hidden, hidden_link)):
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link, relation_type="DependsOn", relation_state="Valid")
        result = AnalyticsQueryService().query({
            "metric": "tickets", "group_by": "provider", "period_from": "2026-10",
            "period_to": "2026-10", "period": "2026-10", "filters": {},
        }, user=self.user)
        self.assertEqual(result["value"], 1)
        self.assertEqual(result["ticket_result_total"], 1)
        self.assertEqual([row["ticket"] for row in result["tickets"]], [permitted.ticket_number])
        self.assertEqual(result["options"]["provider"], [self.operator.name])
        self.assertNotIn("Fora do escopo", [row["dimension"] for row in result["detail"]])
        export = self.client.post(reverse("analytics-export"), data=json.dumps({
            "format": "csv", "metric": "tickets", "group_by": "provider",
            "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {},
        }), content_type="application/json")
        self.assertEqual(export.status_code, 200)
        self.assertIn(self.operator.name, export.content.decode("utf-8-sig"))
        self.assertNotIn("Fora do escopo", export.content.decode("utf-8-sig"))

    def test_cockpit_export_uses_the_same_real_mirror_query(self):
        ticket = self._ticket(105, "OPEN", True, 5)
        OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id="analytics-ci-1", cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        payload = {"metric": "tickets", "group_by": "status", "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {"status": "OPEN", "provider": self.operator.name}}
        expected = AnalyticsQueryService().query(payload, user=self.user)
        response = self.client.post(reverse("analytics-export"), data=json.dumps({**payload, "format": "csv"}), content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        csv_text = response.content.decode("utf-8-sig")
        self.assertIn("Em atendimento", csv_text)
        self.assertIn("A-105", csv_text)
        self.assertNotIn("Demo", csv_text)

    def test_exports_include_full_analysis_even_when_ticket_table_search_is_active(self):
        first = self._ticket(1731, "OPEN", True, 9)
        second = self._ticket(1732, "OPEN", True, 10)
        payload = {"metric": "tickets", "start": "2026-10-01", "end": "2026-10-31", "group_by": "status",
                   "ticket_search": first.ticket_number, "filters": {}}
        visible = AnalyticsQueryService().query(payload, user=self.user)
        self.assertEqual((visible["ticket_total"], visible["ticket_result_total"]), (1, 2))
        response = self.client.post(reverse("analytics-export"), data=json.dumps({**payload, "format": "csv"}),
                                    content_type="application/json")
        self.assertEqual(response.status_code, 200)
        csv_text = response.content.decode("utf-8-sig")
        self.assertIn(first.ticket_number, csv_text)
        self.assertIn(second.ticket_number, csv_text)
        self.assertEqual(len(csv_text.splitlines()), 3)

    def test_cockpit_export_csv_pdf_and_xlsx_are_real_downloads_with_consistent_data(self):
        for index in range(27):
            provider = CMDBOperatorMirror.objects.create(source_config_item_id=f"export-provider-{index}", name=f"Provedor São José {index:02d}")
            link = CMDBLinkMirror.objects.create(source_config_item_id=f"export-link-{index}", operator=provider, state="SP")
            ticket = self._ticket(900 + index, "CANCELED" if index == 26 else "OPEN", index != 26, 5)
            OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id=link.source_config_item_id, cmdb_link=link, relation_type=f"Export-{index}", relation_state="Valid")
        base = {"metric": "tickets", "group_by": "provider", "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "detail_page": 2, "filters": {}}
        expected = AnalyticsQueryService().query(base, user=self.user)
        csv_response = self.client.post(reverse("analytics-export"), data=json.dumps({**base, "format": "csv"}), content_type="application/json")
        csv_text = csv_response.content.decode("utf-8-sig")
        self.assertEqual(csv_response.status_code, 200)
        self.assertEqual(len(csv_text.splitlines()), 27)
        self.assertIn("Provedor São José", csv_text)
        self.assertIn("arion-cockpit-tickets-", csv_response["Content-Disposition"])
        self.assertEqual(expected["value"], 26)
        self.assertEqual(expected["ticket_total"], 26)
        self.assertEqual(len(expected["tickets"]), 25)
        pdf_response = self.client.post(reverse("analytics-export"), data=json.dumps({**base, "format": "pdf"}), content_type="application/json")
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response["Content-Type"], "application/pdf")
        self.assertTrue(pdf_response.content.startswith(b"%PDF-"))
        self.assertGreater(len(pdf_response.content), 1000)
        self.assertIn("arion-cockpit-".encode(), pdf_response["Content-Disposition"].encode())
        xlsx_response = self.client.post(reverse("analytics-export"), data=json.dumps({**base, "format": "xlsx"}), content_type="application/json")
        self.assertEqual(xlsx_response.status_code, 200)
        self.assertIn("spreadsheetml", xlsx_response["Content-Type"])
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Resumo", "Distribuição", "Tickets"])
        self.assertEqual(workbook["Distribuição"].max_row, 27)
        self.assertEqual(workbook["Tickets"].max_row, 27)

    def test_cockpit_export_reason_analysis_keeps_terminal_period_and_count(self):
        ticket = self._ticket(950, "CLOSED", False, 3)
        ticket.source_created_at = datetime(2026, 9, 3, tzinfo=datetime_timezone.utc)
        ticket.terminal_state_at = datetime(2026, 10, 17, tzinfo=datetime_timezone.utc)
        ticket.closing_code_raw = "PendenciaFinanceira"
        ticket.save()
        payload = {"metric": "closed_tickets", "group_by": "termination_reason", "period_from": "2026-10", "period_to": "2026-10", "period": "2026-10", "filters": {"termination_reason": "Pendência Financeira"}}
        result = AnalyticsQueryService().query(payload, user=self.user, include_all_rows=True)
        self.assertEqual(result["value"], 1)
        self.assertEqual(result["detail"][0]["dimension"], "Pendência Financeira")
        for export_format in ("csv", "pdf", "xlsx"):
            response = self.client.post(reverse("analytics-export"), data=json.dumps({**payload, "format": export_format}), content_type="application/json")
            self.assertEqual(response.status_code, 200, export_format)
            self.assertGreater(len(response.content), 50, export_format)
            if export_format == "csv":
                self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
                self.assertIn("Pendência Financeira".encode(), response.content)
                self.assertIn(b"closing_reason_raw", response.content)
                self.assertIn(b"PendenciaFinanceira", response.content)
            elif export_format == "xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(response.content), read_only=True)
                self.assertEqual(workbook.sheetnames, ["Resumo", "Distribuição", "Tickets"])
                self.assertEqual(workbook["Distribuição"].max_row, 2)
                self.assertEqual(workbook["Tickets"].max_row, 2)
                headers = [cell.value for cell in workbook["Tickets"][1]]
                raw_column = headers.index("closing_reason_raw") + 1
                self.assertEqual(workbook["Tickets"].cell(row=2, column=raw_column).value, "PendenciaFinanceira")

    def test_cockpit_export_denies_users_without_cockpit_permission(self):
        user = get_user_model().objects.create_user(username="no-cockpit-export", password="safe-password-4")
        self.client.force_login(user)
        response = self.client.post(reverse("analytics-export"), data=json.dumps({"format": "csv"}), content_type="application/json")
        self.assertEqual(response.status_code, 403)

    def test_cockpit_chart_keeps_months_in_one_scrollable_row(self):
        template = Path("arion_geo/static/css/painel/cockpit.css").read_text(encoding="utf-8")
        script = (Path(settings.BASE_DIR.parent) / "frontend_src" / "js" / "painel" / "cockpit.js").read_text(encoding="utf-8")
        self.assertIn("grid-template-columns: repeat(var(--series-count, 12)", template)
        self.assertIn("data-chart-key=", script)
        self.assertIn(".cockpit-chart { overflow-x: auto; }", template)
        self.assertIn("min-width: 760px", template)
        self.assertIn("data.chart_specs", script)
        self.assertIn("--series-count:${values.length}", script)

    def test_cockpit_api_uses_real_query_and_rejects_bad_payload(self):
        ticket = self._ticket(104, "OPEN", True, 5)
        OTRSTicketConfigItem.objects.create(ticket=ticket, config_item_id="analytics-ci-1", cmdb_link=self.link, relation_type="DependsOn", relation_state="Valid")
        response = self.client.post(reverse("analytics-query"), data=json.dumps({"metric":"tickets", "group_by":"status", "period_from":"2026-10", "period_to":"2026-10", "period":"2026-10", "filters":{}}), content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["value"], 1)
        invalid = self.client.post(reverse("analytics-query"), data="[]", content_type="application/json")
        self.assertEqual(invalid.status_code, 400)
