import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from django.utils import timezone
from django.utils.dateparse import parse_datetime

from aplicativos.base_links.classification import detect_operator, detect_project

from .base import CMDBProvider


REQUIRED_HEADERS = (
    "ConfigItem#", "Nome", "Estado de Implantação", "Designação ARION",
    "Endereço", "Cidade", "Estado", "ISP",
)
SENSITIVE_HEADERS = {"Usuário PPOE", "Senha PPOE"}
UF_CODES = set("AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split())
UF_NAMES = {
    "ACRE": "AC", "ALAGOAS": "AL", "AMAPA": "AP", "AMAZONAS": "AM", "BAHIA": "BA",
    "CEARA": "CE", "DISTRITO FEDERAL": "DF", "ESPIRITO SANTO": "ES", "GOIAS": "GO",
    "MARANHAO": "MA", "MATO GROSSO": "MT", "MATO GROSSO DO SUL": "MS", "MINAS GERAIS": "MG",
    "PARA": "PA", "PARAIBA": "PB", "PARANA": "PR", "PERNAMBUCO": "PE", "PIAUI": "PI",
    "RIO DE JANEIRO": "RJ", "RIO GRANDE DO NORTE": "RN", "RIO GRANDE DO SUL": "RS",
    "RONDONIA": "RO", "RORAIMA": "RR", "SANTA CATARINA": "SC", "SAO PAULO": "SP",
    "SERGIPE": "SE", "TOCANTINS": "TO",
}
MAX_ROWS = 25000
MAX_COLUMNS = 80
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
FIELD_MAP = {
    "Classe": "source_class",
    "Estado de Incidente": "incident_state",
    "Nome": "name",
    "ConfigItem#": "config_item_id",
    "Estado de Implantação": "deployment_state",
    "Versão": "version",
    "Hora de Criação": "source_created_at",
    "Designação ARION": "designation_arion",
    "CNPJ ARION": "cnpj_arion",
    "Endereço": "address",
    "Cidade": "city",
    "Estado": "uf_raw",
    "Cliente": "client_raw",
    "Circuito": "circuit",
    "ISP": "isp_raw",
    "Velocidade": "speed_raw",
    "Range IP": "range_ip",
    "Tipo": "link_type",
    "Número de Contrato": "contract_number",
    "Designação da Operadora": "operator_designation",
    "Identificação (BRAD, SIGLA, AG)": "client_identification",
    "Horário de Funcionamento": "business_hours",
    "Observação (Particularidade da ISP)": "notes",
}
HEADER_ALIASES = {
    "Estado de Incidente": ("Estado do Incidente",),
    "Estado de Implantação": ("Estado Implantação",),
    "Versão": ("Version",),
    "Hora de Criação": ("Created", "Creation Time"),
    "Endereço": ("Endereco",),
    "Número de Contrato": ("Numero de Contrato",),
    "Horário de Funcionamento": ("Horario de Funcionamento",),
    "Observação (Particularidade da ISP)": ("Observação", "Observacao (Particularidade da ISP)"),
    "Identificação (BRAD, SIGLA, AG)": ("Identificação", "Identificacao (BRAD, SIGLA, AG)"),
}
AUDITED_FIELDS = (
    "deployment_state", "incident_state", "name", "designation_arion", "address", "city", "uf",
    "provider", "speed_raw", "range_ip", "link_type", "contract_number", "operator_designation",
    "client_identification", "operadora", "projeto",
)


def _header_key(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch)).casefold()
    return re.sub(r"\s+", " ", value.strip())


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _source_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    parsed = parse_datetime(str(value or "").strip())
    if parsed and timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed)
    return parsed


def normalize_provider(raw):
    raw = _text(raw)
    match = re.match(r"^(.*?)\s*\(([^()]*)\)\s*$", raw)
    valid_id = match.group(2).strip() if match and re.fullmatch(r"\d{6,}", match.group(2).strip()) else ""
    name = (match.group(1).strip() if valid_id else raw).strip()
    if _header_key(name) in {
        "", "-", "--", "n/a", "na", "n.a.", "n.d.", "nao informado", "nao disponivel",
        "sem informacao", "unknown", "desconhecido", "null", "none", "undefined", "placeholder",
    }:
        return "", "", ""
    cmdb_id = valid_id
    # Case/spacing normalization only; no business-level provider merging.
    normalized = re.sub(r"\s+", " ", name).casefold()
    return name, normalized, cmdb_id


class ExcelCMDBProvider(CMDBProvider):
    def __init__(self, content, filename="export.xlsx"):
        self.content = content
        self.filename = filename
        self.sha256 = hashlib.sha256(content).hexdigest()
        self.errors = []
        self.headers = []
        self.raw_items = []

    def load_items(self):
        try:
            with ZipFile(io.BytesIO(self.content)) as archive:
                if "[Content_Types].xml" not in archive.namelist() or "xl/workbook.xml" not in archive.namelist():
                    raise ValueError("O arquivo enviado não tem estrutura XLSX válida.")
                if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("O conteúdo XLSX excede o limite de processamento.")
        except (BadZipFile, OSError, InvalidFileException):
            raise ValueError("O conteúdo enviado não é uma planilha XLSX válida.") from None

        try:
            workbook = load_workbook(io.BytesIO(self.content), read_only=True, data_only=True)
        except (InvalidFileException, KeyError, OSError, ValueError):
            raise ValueError("O conteúdo enviado não é uma planilha XLSX válida.") from None
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        self.headers = [_text(value) for value in next(rows, ())]
        if len(self.headers) > MAX_COLUMNS:
            workbook.close()
            raise ValueError(f"A planilha excede o limite de {MAX_COLUMNS} colunas.")
        self.raw_items = []
        for row_number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            if row_number > MAX_ROWS + 1:
                workbook.close()
                raise ValueError(f"A planilha excede o limite de {MAX_ROWS} linhas de dados.")
            self.raw_items.append({header: values[index] if index < len(values) else None for index, header in enumerate(self.headers)})
        workbook.close()
        return self.raw_items

    def validate(self):
        self.errors.clear()
        available = {_header_key(header): header for header in self.headers}
        self.header_map = {}
        for required in REQUIRED_HEADERS:
            candidates = (required, *HEADER_ALIASES.get(required, ()))
            actual = next((available[_header_key(candidate)] for candidate in candidates if _header_key(candidate) in available), None)
            if not actual:
                self.errors.append({"row": 1, "config_item_id": "", "message": f"Coluna obrigatória ausente: {required}"})
            else:
                self.header_map[required] = actual
        for target, aliases in FIELD_MAP.items():
            candidates = (target, *HEADER_ALIASES.get(target, ()))
            actual = next((available[_header_key(candidate)] for candidate in candidates if _header_key(candidate) in available), None)
            if actual:
                self.header_map[target] = actual
        return not any(error["row"] == 1 for error in self.errors)

    def normalize(self):
        items, seen = [], set()
        if not self.validate():
            return items
        for row_number, raw in enumerate(self.raw_items, start=2):
            item = {}
            for source, target in FIELD_MAP.items():
                header = self.header_map.get(source)
                if header is None:
                    continue
                value = raw.get(header)
                item[target] = _source_datetime(value) if target == "source_created_at" else _text(value)
            source_class = _header_key(item.get("source_class", ""))
            if source_class and source_class != _header_key("Links de internet"):
                self.errors.append({"row": row_number, "config_item_id": item.get("config_item_id", ""), "message": f"Classe não importada para a Base de Links: {item['source_class']}"})
                continue
            config_id = item.get("config_item_id", "").strip()
            if not config_id:
                self.errors.append({"row": row_number, "config_item_id": "", "message": "ConfigItem# vazio."})
                continue
            if config_id in seen:
                self.errors.append({"row": row_number, "config_item_id": config_id, "message": "ConfigItem# duplicado no arquivo; segunda ocorrência ignorada."})
                continue
            seen.add(config_id)
            item["config_item_id"] = config_id
            item["version"] = str(item.get("version", ""))
            uf_raw = item.get("uf_raw", "")
            uf_candidate = _header_key(uf_raw).upper()
            item["uf"] = uf_candidate if uf_candidate in UF_CODES else UF_NAMES.get(uf_candidate, "")
            provider_name, provider_normalized, provider_id = normalize_provider(item.get("isp_raw", ""))
            item["provider_name"] = provider_name
            item["provider_normalized_name"] = provider_normalized
            item["provider_cmdb_config_item_id"] = provider_id
            item["operadora_name"] = detect_operator(item)
            item["projeto_name"] = detect_project(item)
            item["classification_status"] = "classified" if item["operadora_name"] and item["projeto_name"] else "unclassified"
            # Preserve raw data for audit while dropping both credential fields by policy.
            item["source_payload"] = {
                str(key): _text(value) for key, value in raw.items()
                if key and _header_key(key) not in {_header_key(field) for field in SENSITIVE_HEADERS}
            }
            items.append(item)
        return items
